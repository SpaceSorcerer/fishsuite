from pathlib import Path
from types import SimpleNamespace
import importlib.util

import matplotlib.pyplot as plt
import pandas as pd
import pytest

from fishsuite.report import figures, slides


def test_explicit_significance_bracket_has_two_vertical_ticks():
    ctx = SimpleNamespace(
        group_order=["Control", "Treatment"],
        colors=figures.group_colors(["Control", "Treatment"]),
        technical_layer="none",
        axis_groups={},
        axis_windows={},
        reference="Control",
        channel_labels={},
    )
    wells = pd.DataFrame({
        "endpoint": ["rna1_spots_per_nucleus"] * 4,
        "group": ["Control", "Control", "Treatment", "Treatment"],
        "well_id": ["C1", "C2", "T1", "T2"],
        "well_mean_of_field_values": [2.0, 3.0, 5.0, 7.0],
    })
    canvas, ax = plt.subplots()

    figures.draw_replicate_simple(
        ax, ctx, "rna1_spots_per_nucleus", wells, pd.DataFrame(),
        pd.DataFrame(), pd.DataFrame(), "Puncta per nucleus", None,
        bracket={"x0": 0, "x1": 1, "label": "adjusted p = 0.04"},
    )
    ax._replicate_simple_axis("focus")

    artists = {artist.get_gid(): artist for artist in [*ax.lines, *ax.texts]}
    assert artists["bracket:bar"].get_xdata().tolist() == [0, 1]
    assert artists["bracket:bar"].get_ydata()[0] > 7.0
    for key in ("bracket:tick0", "bracket:tick1"):
        tick = artists[key]
        assert tick.get_xdata()[0] == tick.get_xdata()[1]
        assert tick.get_ydata()[0] < tick.get_ydata()[1]
    assert artists["bracket:label"].get_text() == "adjusted p = 0.04"
    assert artists["bracket:label"].get_position()[1] < ax.get_ylim()[1]
    plt.close(canvas)


def test_footer_lines_wrap_at_readable_size_and_raise_axes():
    canvas, ax = plt.subplots(figsize=(4.8, 3.6))
    ax.set_xticks([0, 1], ["Control", "Treatment"])
    footer_lines = [
        "Detection floor: recorded from the source run configuration",
        "Gate: all declared exclusions applied before aggregation",
        "Replicates: wells; test: two-sided comparison of well means",
    ]

    figures.layout_replicate_simple(
        canvas, ax, SimpleNamespace(run_name="synthetic-run"), "Endpoint", "Source",
        footer_lines=footer_lines,
    )
    canvas.canvas.draw()

    assert len(canvas.texts) == 3
    block = canvas.texts[2]
    flat = " ".join(block.get_text().split())
    assert all(" ".join(line.split()) in flat for line in footer_lines)
    assert block.get_fontsize() >= 6
    assert ax.get_position().y0 > block.get_window_extent().y1 / canvas.bbox.height
    plt.close(canvas)


def test_speaker_note_trailer_accepts_generic_overrides_and_omissions():
    definition = {
        "values": [{"label": "readout", "value": "Measured endpoint."}],
    }

    notes = slides.speaker_notes(
        Path("report.xlsx"), definition,
        notes_trailer={
            "replicate_unit": "Images are technical units and cultures are biological units",
            "headline_statistic": "the headline is the culture-level comparison",
            "decision_line": "",
        },
    )

    assert "Images are technical units" in notes
    assert "culture-level comparison" in notes
    assert "Mixed-model headline decision" not in notes


def test_speaker_note_trailer_rejects_unknown_keys():
    with pytest.raises(RuntimeError, match="unknown notes_trailer key"):
        slides.speaker_notes(Path("report.xlsx"), {"values": []}, {"headline": "x"})


@pytest.mark.skipif(importlib.util.find_spec("pptx") is None, reason="python-pptx missing")
def test_build_deck_uses_the_spec_note_trailer(tmp_path):
    from fishsuite.report.provenance import sha256
    from fishsuite.report.workbook import write
    from pptx import Presentation

    workbook = write(
        tmp_path / "REPORT.xlsx",
        {"Values": pd.DataFrame({"readout": ["Measured endpoint."]})},
        order=["Values"],
    )
    spec = {
        "workbook_sha256": sha256(workbook),
        "notes_trailer": {
            "replicate_unit": "Cultures are biological units",
            "headline_statistic": "the headline is the culture-level comparison",
            "decision_line": "",
        },
        "slides": [{
            "title": "Measured endpoint",
            "values": [{"sheet": "Values", "cell": "A3", "label": "readout"}],
            "figures": [],
        }],
    }

    deck = slides.build_deck(workbook, spec, tmp_path / "deck.pptx")
    notes = Presentation(deck).slides[0].notes_slide.notes_text_frame.text

    assert "Cultures are biological units" in notes
    assert "culture-level comparison" in notes
    assert "Mixed-model headline decision" not in notes
