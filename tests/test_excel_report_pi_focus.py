"""Regression tests for excel_report PI_Focus / Comparison_Table helpers."""
import pandas as pd
from openpyxl import Workbook

from fishsuite.core import excel_report as er


def _frames():
    per_image = pd.DataFrame(dict(
        image=["a.nd2", "b.nd2", "s.nd2"],
        condition=["WT", "KO", "WT"],
        secondary_only=[False, False, True],
        nuclei_analyzed=[3, 3, 2],
    ))
    nuclei = pd.DataFrame(dict(
        image=["a.nd2"] * 3 + ["b.nd2"] * 3 + ["s.nd2"] * 2,
        nucleus_id=[1, 2, 3, 1, 2, 3, 1, 2],
        condition=["WT"] * 3 + ["KO"] * 3 + ["WT"] * 2,
        secondary_only=[False] * 6 + [True] * 2,
        n_spots_rna1=[4, 2, 0, 1, 1, 1, 0, 0],
    ))
    # WT nucleus 1: two nucleolar spots; nucleus 2: one non-nucleolar
    # nuclear spot; nucleus 3: no spots. KO nucleus 2: one nucleolar spot.
    spots = pd.DataFrame(dict(
        image=["a.nd2", "a.nd2", "a.nd2", "b.nd2"],
        condition=["WT", "WT", "WT", "KO"],
        secondary_only=[False] * 4,
        channel=["rna1"] * 4,
        nucleus_id=[1, 1, 2, 2],
        in_nucleus=[1, 1, 1, 1],
        in_nucleolus=[1, 1, 0, 1],
        spot_peak_intensity=[10.0, 12.0, 8.0, 9.0],
    ))
    return per_image, nuclei, spots


def _row_for(ws, label):
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == label:
            return row
    raise AssertionError(f"row {label!r} not found")


def test_pi_focus_nucleolar_counts_include_cells_with_zero_spots():
    per_image, nuclei, spots = _frames()
    ws = Workbook().active
    er._build_pi_focus(ws, per_image_df=per_image, nuclei_df=nuclei, spots_df=spots,
                       comparison_unit="nucleus")
    row = _row_for(ws, "Nucleolar RNA1 spots per cell")
    # WT: [2, 0, 0] over all three real nuclei (sec-only nuclei excluded).
    assert row[1].startswith("0.667") and row[1].endswith("(n=3)")
    # KO: [0, 1, 0].
    assert row[2].startswith("0.333") and row[2].endswith("(n=3)")
    # Default per-image unit: the image mean still counts the zero cells.
    ws = Workbook().active
    er._build_pi_focus(ws, per_image_df=per_image, nuclei_df=nuclei, spots_df=spots)
    row = _row_for(ws, "Nucleolar RNA1 spots per cell")
    assert row[1] == "0.667 (n=1)"


def test_pi_focus_labels_use_generic_channel_tokens():
    per_image, nuclei, spots = _frames()
    ws = Workbook().active
    er._build_pi_focus(ws, per_image_df=per_image, nuclei_df=nuclei, spots_df=spots)
    labels = [r[0] for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str)]
    assert not any("Intron" in s or "Exon" in s for s in labels)
    assert "Total RNA1 spots per cell" in labels
    er._relabel_worksheet(ws, {"rna_label": "GeneA", "rna2_label": "GeneB", "dapi_label": "DAPI"})
    labels = [r[0] for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str)]
    assert "Total GeneA spots per cell" in labels


def test_comparison_sec_only_uses_same_source_and_nuclei_row_label():
    per_image, nuclei, _ = _frames()
    per_image["mean_spots_per_nucleus_rna1"] = [2.0, 1.0, 99.0]
    nuclei.loc[nuclei["secondary_only"], "n_spots_rna1"] = [1, 3]
    ws = Workbook().active
    er._build_comparison_table(ws, per_image_df=per_image, nuclei_df=nuclei)
    hdr = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "Metric")
    row = _row_for(ws, "Mean RNA1 spots per nucleus")
    # WT/KO use per-nucleus values (averaged per image), so sec-only must
    # too: nuclei [1, 3] in one image -> 2.0 (n=1), not the per-image 99.
    assert row[hdr.index("sec-only_mean")] == 2.0
    assert row[hdr.index("sec-only_n")] == 1
    labels = [r[0] for r in ws.iter_rows(values_only=True) if r]
    assert "Nuclei analyzed (sum)" not in labels
    assert "Nuclei analyzed per image (mean)" in labels


def test_secondary_only_coercion_and_sheet_only_relabel():
    df = pd.DataFrame(dict(condition=["WT", "KO", "KO"],
                           secondary_only=["False", "True", None]))
    out = er._coerce_secondary_only(df)
    assert out["secondary_only"].dtype == bool
    assert out["secondary_only"].tolist() == [False, True, False]
    # Sorting keeps the parent condition; only the data-sheet view relabels.
    sorted_df = er._sort_df_by_condition(out)
    assert set(sorted_df["condition"]) == {"WT", "KO"}
    sheet = er._relabel_sec_only_for_sheet(sorted_df)
    assert sheet.loc[sheet["secondary_only"], "condition"].tolist() == ["sec-only"]
    assert sheet["secondary_only"].tolist() == sorted_df["secondary_only"].tolist()


def test_fractional_columns_not_integer_formatted():
    for col in ("spot_fwhm_px", "voxel_xy_nm", "voxel_z_nm", "z_position_um"):
        assert col not in er.INT_COLUMNS
    assert er.PER_SPOT_GLOSSARY["z_position_um"][0] == "float"
    assert er.PER_SPOT_GLOSSARY["spot_fwhm_px"][0] == "float"
    assert er.PER_IMAGE_GLOSSARY["voxel_xy_nm"][0] == "float"


def _two_images_per_arm(n_nuc=12, with_well=True):
    rows = []
    for cond, imgs, base in (("WT", ["w1.nd2", "w2.nd2"], 5.0), ("KO", ["k1.nd2", "k2.nd2"], 1.0)):
        for i, img in enumerate(imgs):
            for nid in range(1, n_nuc + 1):
                rows.append(dict(image=img, nucleus_id=nid, condition=cond, secondary_only=False,
                                 well_id=f"{cond}_well", n_spots_rna1=base + i + nid % 3))
    nuclei = pd.DataFrame(rows)
    per_image = nuclei.groupby(["image", "condition", "well_id"], as_index=False).agg(
        nuclei_analyzed=("nucleus_id", "size"), mean_spots_per_nucleus_rna1=("n_spots_rna1", "mean"))
    per_image["secondary_only"] = False
    if not with_well:
        nuclei = nuclei.drop(columns="well_id")
        per_image = per_image.drop(columns="well_id")
    return per_image, nuclei


def _cmp_row(unit, **kw):
    per_image, nuclei = _two_images_per_arm(**kw)
    ws = Workbook().active
    er._build_comparison_table(ws, per_image_df=per_image, nuclei_df=nuclei, comparison_unit=unit)
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(r for r in rows if r and r[0] == "Metric")
    row = _row_for(ws, "Mean RNA1 spots per nucleus")
    return dict(zip(hdr, row)), rows[1][0]


def test_comparison_unit_sets_test_n():
    image, note = _cmp_row(None)                       # default = image
    assert image["WT_n"] == 2 and image["KO_n"] == 2
    assert image["MWU p-value"] is not None            # 2 vs 2 images
    assert "per-image means" in note and "n = images" in note
    nucleus, note = _cmp_row("nucleus")
    assert nucleus["WT_n"] == 24 and nucleus["KO_n"] == 24
    assert "pseudoreplicates" in note
    assert nucleus["MWU p-value"] < image["MWU p-value"]
    well, note = _cmp_row("well")
    assert well["WT_n"] == 1 and well["KO_n"] == 1 and well["MWU p-value"] is None
    assert "per-well means" in note
    fallback, note = _cmp_row("well", with_well=False)
    assert fallback["WT_n"] == 2 and "fell back to per-image" in note


def test_pi_focus_uses_the_comparison_unit():
    per_image, nuclei = _two_images_per_arm()
    for unit, n in ((None, 2), ("nucleus", 24)):
        ws = Workbook().active
        er._build_pi_focus(ws, per_image_df=per_image, nuclei_df=nuclei, spots_df=pd.DataFrame(),
                           comparison_unit=unit)
        assert _row_for(ws, "Total RNA1 spots per cell")[1].endswith(f"(n={n})")
