from pathlib import Path
import importlib.util
import pytest
import pandas as pd


@pytest.mark.parametrize('rna,partner',[('BIN1 intron','RNASEH2B'),('MIAT','QKI')])
def test_simple_coloc_slide_identities(tmp_path,rna,partner):
    from fishsuite.report.slides import simple_coloc_slides
    assets=[]
    for name in ['FIG_SIMPLE_COLOC_focus','FIG_CYTOFLUOROGRAM']:
        p=tmp_path/(name+'.png'); p.write_bytes(b'asset'); assets.append(p)
    sheets={'Simple coloc metrics':pd.DataFrame({'endpoint':['Pearson r'],'p_mixed':[.1]})}
    slides=simple_coloc_slides(sheets,assets,rna,partner,'retained')
    assert [s['identity'] for s in slides]==['simple_coloc','cytofluorogram']
    assert slides[0]['title']==f'Pixel colocalization, {rna} × {partner}'
    assert all(s['values'] for s in slides)


def test_recorded_deck_semantic_assets_and_localization(tmp_path, monkeypatch):
    import yaml
    from fishsuite.report.build import build_report
    from pptx import Presentation
    from fishsuite.report import figures
    original = figures.draw_replicate_simple
    checked = []
    def check_panel(ax, ctx, endpoint, *args, **kwargs):
        result = original(ax, ctx, endpoint, *args, **kwargs)
        texts = [t.get_text() for t in ax.texts]
        assert any('p =' in t or 'descriptive, no test' in t for t in texts), (endpoint, texts)
        checked.append(endpoint)
        return result
    monkeypatch.setattr(figures, 'draw_replicate_simple', check_panel)
    root = Path(__file__).resolve().parents[1]
    spec = root/'_closeout_evidence/A/deck_spec.yaml'
    run = Path('F:/Image Analysis Work/RNASEH2B_BIN1introns_2026_08_25/13b_FULL_HARMONIZED_T36_FIXEDNUCLEAR_2026-09-05/RUN_T36_fixed_2026-09-05_0915')
    result = build_report(run, tmp_path/'report', groups=['WT=WT_1,WT_2,WT_3', 'QKI-KO=KO_1,KO_2,KO_3'],
                          reference='WT', make_figures=False, deck_spec=spec, deck=True)
    resolved = yaml.safe_load((result['out_dir']/'deck_spec.resolved.yaml').read_text())
    expected = {'overview':'A01_overview', 'q1_count_size':'A04_question1',
                'q1_localization':'A05_BIN1_localization', 'q1_total_if':'A04_RNASEH2B_total_IF',
                'q2':'A06_question2', 'q3':'A07_question3', 'q4_reverse_anchor':'A08_question4',
                'nucleus_selection':'A03_nucleus_selection', 'methods_provenance':'A02_methods',
                'micrographs':'A10_micrographs', 'standard_coloc':'A09_standard_coloc'}
    assert [s['identity'] for s in resolved['slides']] == list(expected)
    for slide in resolved['slides']:
        for entry in slide['figures']:
            path=Path(entry['path'])
            assert 'composites' not in path.parts
            if slide['identity'] != 'micrographs':
                assert path.with_suffix('.svg').is_file()
        if slide['identity'] in {'q1_count_size','q1_total_if','q2','q3','q4_reverse_anchor'}:
            assert len(slide['figures']) == len(next(s for s in __import__('yaml').safe_load(spec.read_text())['slides'] if s['identity']==slide['identity'])['endpoints'])
    index = pd.read_excel(result['xlsx'], sheet_name='FIGURE_INDEX', header=1)
    assert index.path.fillna('').str.endswith('_focus.png').sum() > 7
    sources = pd.read_csv(result['out_dir']/'slide_sources.csv')
    paired = sources.loc[sources.figure.fillna('').str.endswith('_focus.png')]
    assert len(paired) > 7
    assert paired.full_figure.str.endswith('_full.png').all()
    assert paired.full_figure.map(lambda p: Path(p).is_file()).all()
    crops = pd.read_excel(result['xlsx'], sheet_name='Localization crops', header=1)
    assert crops.crop_status.eq('available').all()
    assert crops.territory_boundary.eq('missing').all()
    svg = (result['out_dir']/'localization/rna1_cyto_spots_per_nucleus_localization_focus.svg').read_text()
    import re
    import xml.etree.ElementTree as ET
    svg_text=' '.join(''.join(node.itertext()) for node in ET.fromstring(svg).iter()
                      if node.tag.endswith('}text'))
    assert 'per nucleus and assigned cell territory' in svg_text.lower()
    assert 'Assigned-cytoplasmic puncta count' in svg_text or 'cytoplasmic puncta' in svg_text
    ppt = Presentation(result['out_dir']/'Sam_RNASEH2B_BIN1.pptx')
    panels = pd.read_excel(result['xlsx'], sheet_name='Micrograph panels', header=1)
    assert panels.groupby('group').size().eq(4).all()
    assert not panels.panel.str.contains('DAPI-only',case=False).any()
    assert 'FOV outlier sensitivity' in pd.ExcelFile(result['xlsx']).sheet_names
    for slide in ppt.slides:
        notes = slide.notes_slide.notes_text_frame.text
        assert notes.startswith('Workbook: '+str(result['xlsx'].resolve()))
        assert '!' not in notes
        assert 'Levels of comparison' in notes
    assert [s.shapes[0].text for s in ppt.slides] == [s['title'] for s in resolved['slides']]
    assert resolved['slides'][1]['title'] == '1. Are BIN1 intron puncta larger, more numerous and more nuclear in KO than WT? (sanity check)'
    assert resolved['slides'][3]['title'] == '1b. RNASEH2B total signal and puncta, WT vs KO'
    for i in (4, 5, 6):
        assert resolved['slides'][i]['title'].startswith(str(i-2)+'. ')
    assert {'frac_called_coloc_shuffle_runthr', 'frac_called_coloc_minus_shuffle_runthr'} <= set(checked)
    composite=(result['out_dir']/'localization/composites/FIG_LOCALIZATION_focus.svg').read_text()
    assert 'cytoplasmic = outside the 2D nuclear mask within the assigned territory; single plane' in composite
    for slide, definition in zip(ppt.slides, resolved['slides']):
        if definition.get('layout')!='micrographs':
            assert any(s.has_text_frame and s.top > 5000000 for s in slide.shapes)
    for slide, definition in zip(ppt.slides, resolved['slides']):
        assert len([s for s in slide.shapes if s.shape_type == 13]) == len(definition['figures'])


@pytest.mark.parametrize('claim', ['999 nuclei', '1e9 nuclei', '25um scale bar'])
@pytest.mark.parametrize('field', ['title', 'body', 'label'])
def test_export_rejects_literal_numeric_text(tmp_path, field, claim):
    from fishsuite.report.slides import build_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    book = write(tmp_path/'REPORT.xlsx', {'Values': pd.DataFrame({'value': [1]})}, order=['Values'])
    slide = dict(title='Measured value', values=[dict(sheet='Values', cell='A3', label='nuclei')])
    (slide['values'][0] if field == 'label' else slide)[field] = claim
    with pytest.raises(RuntimeError, match='literal numeric'):
        build_deck(book, dict(workbook_sha256=sha256(book), slides=[slide]), tmp_path/'bad.pptx')
    assert not (tmp_path/'bad.pptx').exists()


@pytest.mark.parametrize('via_junction', [False, True])
@pytest.mark.parametrize('writer', ['save', 'localization', 'workbook', 'deck'])
def test_frozen_writers_before_side_effect(tmp_path, monkeypatch, writer, via_junction):
    from unittest.mock import Mock
    from fishsuite.report import figures, workbook, slides
    target = tmp_path/'DELIVERY_frozen'/'child'
    (tmp_path/'DELIVERY_frozen').mkdir()
    (tmp_path/'DELIVERY_frozen'/'MANIFEST_SHA256.tsv').write_text('')  # release manifest = frozen
    if via_junction:
        import subprocess
        frozen = Path('F:/Image Analysis Work/RNASEH2B_BIN1introns_2026_08_25/DELIVERY_RNASEH2B_BIN1intron_2026-09-05_v3')
        link = tmp_path/'alias'
        subprocess.run(['powershell', '-NoProfile', '-Command',
                        f"New-Item -ItemType Junction -Path '{link}' -Target '{frozen}' | Out-Null"], check=True)
        assert link.resolve() == frozen.resolve()
        target = link/'child'
    mkdir = Mock(side_effect=AssertionError('mkdir side effect'))
    opened = Mock(side_effect=AssertionError('open side effect'))
    monkeypatch.setattr(Path, 'mkdir', mkdir)
    monkeypatch.setattr('builtins.open', opened)
    with pytest.raises(RuntimeError, match='frozen output'):
        if writer == 'save':
            figures.save(Mock(), target, 'test', [], '', '')
        elif writer == 'localization':
            figures.render_localization(None, None, None, None, None, None, target)
        elif writer == 'workbook':
            workbook.write(target/'test.xlsx', {})
        else:
            slides.build_deck(tmp_path/'absent.xlsx', {}, target/'test.pptx')
    mkdir.assert_not_called()
    opened.assert_not_called()


def test_save_nonfrozen_and_descendant_guard(tmp_path, monkeypatch):
    from unittest.mock import Mock
    from fishsuite.report import figures
    figure, ax = figures.plt.subplots()
    ax.plot([0, 1])
    manifest = []
    figures.save(figure, tmp_path, 'allowed', manifest, '', '')
    assert (tmp_path/'allowed.png').is_file() and (tmp_path/'allowed.svg').is_file()
    (tmp_path/'DELIVERY_frozen').mkdir(exist_ok=True)
    (tmp_path/'DELIVERY_frozen'/'MANIFEST_SHA256.tsv').write_text('')
    mkdir = Mock(side_effect=AssertionError('mkdir side effect'))
    monkeypatch.setattr(Path, 'mkdir', mkdir)
    with pytest.raises(RuntimeError, match='frozen output'):
        figures.save(Mock(), tmp_path, 'DELIVERY_frozen/nested', [], '', '')
    mkdir.assert_not_called()


def test_slide_validation_missing_and_stale(tmp_path):
    from fishsuite.report.slides import validate_deck
    from fishsuite.report.provenance import sha256
    from fishsuite.report.workbook import write
    book = write(tmp_path/'REPORT.xlsx', {'Slide values':pd.DataFrame({'value':[1]})}, order=['Slide values'])
    spec = {'workbook_sha256':sha256(book), 'slides':[{'title':'Measured value', 'values':[{'sheet':'Slide values','cell':'A3'}], 'figures':[]}]}
    validate_deck(book, spec)
    spec['slides'][0]['values'][0]['cell'] = 'A300'
    with pytest.raises(RuntimeError, match='missing'):
        validate_deck(book, spec)
    spec['slides'][0]['values'][0]['allow_missing'] = True
    with pytest.raises(RuntimeError, match='missing workbook cell'):
        validate_deck(book, spec)
    spec['workbook_sha256'] = '0'*64
    with pytest.raises(RuntimeError, match='stale'):
        validate_deck(book, spec)


def test_slide_validation_rejects_untraced_numbers(tmp_path):
    from fishsuite.report.slides import validate_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    book = write(tmp_path/'REPORT.xlsx', {'Slide values':pd.DataFrame({'value':[1]})}, order=['Slide values'])
    with pytest.raises(RuntimeError, match='literal numeric'):
        validate_deck(book, {'workbook_sha256':sha256(book), 'slides':[{'title':'393 nuclei', 'values':[], 'figures':[]}]})


def test_missing_stale_figure_and_cohort(tmp_path):
    from fishsuite.report.slides import validate_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    book=write(tmp_path/'REPORT.xlsx',{'Slide values':pd.DataFrame({'value':[1]})},order=['Slide values'])
    asset=tmp_path/'figure.png'
    spec={'workbook_sha256':sha256(book),'cohort':'retained',
          'slides':[{'title':'Measured endpoint','values':[],
                     'figures':[{'path':str(asset),'sha256':'0'*64,'cohort':'retained'}]}]}
    with pytest.raises(RuntimeError,match='missing figure'):
        validate_deck(book,spec)
    asset.write_bytes(b'persisted test asset')
    with pytest.raises(RuntimeError,match='stale figure'):
        validate_deck(book,spec)
    spec['slides'][0]['figures'][0]['sha256']=sha256(asset)
    spec['slides'][0]['figures'][0]['cohort']='filtered'
    with pytest.raises(RuntimeError,match='cohort mismatch'):
        validate_deck(book,spec)


def test_stale_slide_value_lineage_rejected(tmp_path):
    from fishsuite.report.slides import validate_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    sheets={'Contrasts':pd.DataFrame({'diff':[-2.]}),
            'Slide values':pd.DataFrame([dict(label='MEASURED direction',value='higher',source_sheet='Contrasts',source_cell='A3')])}
    book=write(tmp_path/'REPORT.xlsx',sheets,order=list(sheets))
    with pytest.raises(RuntimeError,match='stale Slide values lineage'):
        validate_deck(book,{'schema':'resolved-deck-1','workbook_sha256':sha256(book),
                            'slides':[{'title':'Measured direction','values':[],'figures':[]}]})


@pytest.mark.skipif(importlib.util.find_spec('pptx') is None,
                    reason='python-pptx missing from fishproc_dml; dispatch prohibits installation')
def test_pptx_notes_media_and_count(tmp_path):
    from fishsuite.report.slides import build_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    from pptx import Presentation
    book = write(tmp_path/'REPORT.xlsx', {'Slide values':pd.DataFrame({'value':[1]})}, order=['Slide values'])
    spec = {'workbook_sha256':sha256(book), 'slides':[{'title':'Measured value','values':[{'sheet':'Slide values','cell':'A3'}], 'figures':[]}]}
    path = build_deck(book, spec, tmp_path/'deck.pptx')
    ppt = Presentation(path)
    assert len(ppt.slides) == 1
    assert 'Workbook: '+str(book.resolve()) in ppt.slides[0].notes_slide.notes_text_frame.text
    assert '!' not in ppt.slides[0].notes_slide.notes_text_frame.text


def test_export_allows_explicit_source_identifiers(tmp_path):
    from fishsuite.report.slides import build_deck
    from fishsuite.report.workbook import write
    from fishsuite.report.provenance import sha256
    from pptx import Presentation
    table = pd.DataFrame(dict(stem=['2026_09_field'], group=['25KO']))
    book = write(tmp_path/'REPORT.xlsx', {'Micrographs': table, '25 assays': table}, order=['Micrographs', '25 assays'])
    slide = dict(title='2026_09_field / 25KO / 25 assays',
                 values=[dict(sheet='Micrographs', cell='A3', label='file stem')])
    path = build_deck(book, dict(workbook_sha256=sha256(book), slides=[slide]), tmp_path/'ids.pptx')
    assert Presentation(path).slides[0].shapes[0].text == slide['title']


def test_micrograph_labels_are_slide_readable_and_inside_panel(tmp_path):
    import numpy as np
    import matplotlib.pyplot as plt
    from fishsuite.report.slides import draw_micrograph_panel
    path = tmp_path / "micrograph.png"
    plt.imsave(path, np.zeros((100, 150, 3)))
    canvas, ax = plt.subplots(figsize=(6.2, 4.8))
    row = dict(path=str(path), group="WT", bar_um=5., voxel_xy_um=.1,
               native_bar_px=50, width_px=150, height_px=100)
    draw_micrograph_panel(ax, row)
    canvas.canvas.draw()
    assert all(t.get_fontsize() >= 14 for t in [ax.title, *ax.texts])
    for text in ax.texts:
        box = text.get_window_extent()
        panel = ax.get_window_extent()
        assert panel.contains(box.x0, box.y0) and panel.contains(box.x1, box.y1)
    bar = ax.lines[-1]
    assert abs(np.diff(bar.get_xdata())[0]) == 50
    assert min(bar.get_xdata()) > ax.get_xlim()[0]
    assert max(bar.get_xdata()) < ax.get_xlim()[1]
    plt.close(canvas)


def test_cytoplasmic_call_csv_and_missing_overlay(tmp_path):
    import numpy as np
    import tifffile
    from PIL import Image
    from fishsuite.report.cyto_calls import build_cyto_calls
    run = tmp_path/'run'
    (run/'publication_images').mkdir(parents=True)
    (run/'masks').mkdir()
    labels = np.zeros((80,80), dtype=np.uint16)
    labels[20:40,20:40] = 1
    tifffile.imwrite(run/'masks/WT_1__field__nuclei_label_mask.tif', labels)
    Image.new('RGB',(80,80)).save(run/'publication_images/WT_1__field__merge_all.png')
    pd.DataFrame([dict(image=image, condition='WT_1', channel='rna1', spot_id=1,
                      nucleus_id=1, in_cytoplasm=1, x_px=41, y_px=30)
                  for image in ['field.vsi','absent.vsi']]).to_csv(run/'spot_metrics.csv', index=False)
    nuclei = pd.DataFrame([dict(image=image, nucleus_id=1, group='WT', voxel_xy_um=1.)
                           for image in ['field.vsi','absent.vsi']])
    out=tmp_path/'qc'
    build_cyto_calls(run,out,nuclei)
    calls=pd.read_csv(out/'cytoplasmic_calls.csv')
    assert len(calls)==2
    assert calls.set_index('image').loc['field.vsi','distance_to_mask_edge_um']==pytest.approx(1.5)
    assert calls.set_index('image').loc['absent.vsi','status']=='missing overlay'
    assert calls.set_index('image').loc['field.vsi','crop_px']==6
    assert len(pd.read_csv(out/'missing.csv'))==1
    assert (out/'WT_sheet_001.png').is_file()
