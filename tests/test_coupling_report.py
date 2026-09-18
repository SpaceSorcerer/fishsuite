from pathlib import Path

import numpy as np
import pandas as pd
import pytest


def toy_table():
    rows = []
    for condition, well, n, c, intensity in (
        ('control', 'C1', 2, 1, 10), ('control', 'C2', 6, 3, 30),
        ('treated', 'T1', 8, 6, 40), ('treated', 'T2', 16, 12, 80),
    ):
        for nucleus in range(5):
            rows.append(dict(condition=condition, well=well, image=well + '.tif',
                nucleus_id=nucleus + 1, threshold_multiplier=1.0,
                miat_min_used=10, qki_min_used=20, measurement_plane='single-plane',
                n_miat_spots=n, n_miat_spots_qki_pos=c, integrated_nuclear_miat=intensity,
                nuclear_area_um2=100, mean_nuclear_qki=10, mean_qki_at_miat_spots=12,
                qki_at_spots_minus_nuclear=2, frac_miat_spots_qki_pos=c / n,
                obs_minus_null_frac_miat_spots_qki_pos=c / n - .25,
                frac_miat_footprint_area_qki_pos=c / n,
                obs_minus_null_frac_miat_footprint_area_qki_pos=c / n - .25,
                sat_frac_miat=0, sat_frac_qki=0))
    return pd.DataFrame(rows)


def summary(df=None, **kwargs):
    from fishsuite.report.coupling_stats import summarize
    return summarize(toy_table() if df is None else df, 'treated', 'control', n_boot=kwargs.pop('n_boot', 30), **kwargs)


def three_well_table(inner_variation=False):
    rows = []
    template = toy_table().iloc[0].to_dict()
    for condition in ('control', 'treated'):
        for index, (constant, count) in enumerate(zip((3, 33, 63), (5, 5, 10)), 1):
            well = f'{condition}_{index}'
            values = (3, 9) if inner_variation else [constant] * count
            for nucleus, count_value in enumerate(values, 1):
                rows.append(dict(template, condition=condition, well=well, image=well + '.tif',
                                 nucleus_id=nucleus, n_miat_spots=count_value,
                                 integrated_nuclear_miat=count_value * 5))
    return pd.DataFrame(rows)


def test_literal_summary_and_secondary_ratios():
    result = summary()
    wells = result['per_well'].set_index('well')
    assert wells.loc['C1', 'mean_n_miat_spots'] == 2
    assert wells.loc['C2', 'median_n_miat_spots'] == 6
    assert wells.loc['T1', 'n_nuclei'] == 5
    assert wells.loc['T1', 'n_nuclei_with_spots'] == 5
    row = result['contrast'].set_index('metric').loc['n_miat_spots']
    assert row.treated_mean_of_well_means == 12
    assert row.control_mean_of_well_means == 4
    assert row.difference == 8
    assert row.ratio == 3
    assert (row.control_well_1_mean, row.control_well_2_mean) == (2, 6)
    assert (row.treated_well_1_mean, row.treated_well_2_mean) == (8, 16)
    ratios = result['ratio_of_ratios'].set_index('metric')
    assert ratios.loc['ror_count', 'value'] == 1.5
    assert ratios.loc['ror_intensity', 'value'] == 1.5
    assert ratios.loc['ror_corrected', 'value'] == .25
    assert 'difference' in ratios.loc['ror_corrected', 'definition']


def test_unequal_nuclei_keep_equal_well_weights():
    df = toy_table()
    df = pd.concat([df[df.well == w].iloc[[0] * size] for w, size in
                    [('C1', 4), ('C2', 12), ('T1', 4), ('T2', 12)]], ignore_index=True)
    df['nucleus_id'] = df.groupby('well').cumcount() + 1
    row = summary(df)['contrast'].set_index('metric').loc['n_miat_spots']
    assert row.control_mean_of_well_means == 4
    assert row.treated_mean_of_well_means == 12
    assert row.ratio == 3


def test_zero_control_abundance_is_nan_with_reason():
    df = toy_table()
    df.loc[df.condition == 'control', 'n_miat_spots'] = 0
    ratio = summary(df)['ratio_of_ratios'].set_index('metric').loc['ror_count']
    assert pd.isna(ratio.value)
    assert ratio.na_reason == 'CONTROL_NONPOSITIVE'


@pytest.mark.parametrize('count,distinct,expected', [(9, 9, None), (12, 12, 1.0), (12, 4, None)])
def test_correlation_gate(count, distinct, expected):
    df = pd.concat([toy_table().iloc[[0] * count], toy_table().iloc[[10] * count]], ignore_index=True)
    df['nucleus_id'] = df.groupby('well').cumcount() + 1
    vals = np.arange(count) % distinct + 1
    df['n_miat_spots'] = np.tile(vals, 2)
    df['integrated_nuclear_miat'] = df.n_miat_spots * 10
    df['obs_minus_null_frac_miat_spots_qki_pos'] = df.n_miat_spots / 100
    df['qki_at_spots_minus_nuclear'] = df.n_miat_spots * 2
    rows = summary(df)['within_well_correlation']
    if expected is None:
        assert rows.spearman_r.isna().all()
        assert rows.na_reason.str.len().gt(0).all()
    else:
        assert np.allclose(rows.spearman_r, expected)
        assert np.allclose(rows.pearson_r, 1.0)
        assert rows.na_reason.eq('').all()


def test_bootstrap_seed_and_exploratory_columns():
    first = summary(three_well_table(), seed=42)['contrast']
    second = summary(three_well_table(), seed=42)['contrast']
    pd.testing.assert_frame_equal(first, second)
    assert {'exploratory_boot_lo', 'exploratory_boot_hi'} <= set(first)
    assert 'boot_lo' not in first and 'boot_hi' not in first
    row = first.set_index('metric').loc['n_miat_spots']
    assert np.isfinite(row.exploratory_boot_lo)
    assert row.exploratory_boot_na_reason == ''


def test_bootstrap_resamples_wells_with_hand_derived_support():
    result = summary(three_well_table(), seed=7, n_boot=5000)['contrast'].set_index('metric')
    row = result.loc['n_miat_spots']
    # Six independent ternary draws give 729 allocations. CDF(-50)=7/729;
    # CDF(-40)=28/729, so the 2.5/97.5 percentile endpoints are -40 and 40.
    assert (row.exploratory_boot_lo, row.exploratory_boot_hi) == (-40, 40)
    hierarchical_support = set(range(-60, 61, 10))
    # A pooled nucleus-only resample of 20 nuclei per arm has a 30/20=1.5 lattice.
    pooled_nucleus_only_support = {1.5 * k for k in range(-40, 41)}
    fixed_well_nucleus_only_support = {0}
    for bound in (row.exploratory_boot_lo, row.exploratory_boot_hi):
        assert bound in hierarchical_support
        assert bound not in pooled_nucleus_only_support
        assert bound not in fixed_well_nucleus_only_support
    intensity = result.loc['integrated_nuclear_miat']
    assert (intensity.exploratory_boot_lo, intensity.exploratory_boot_hi) == (-200, 200)


def test_bootstrap_resamples_nuclei_inside_selected_wells():
    result = summary(three_well_table(inner_variation=True), seed=7, n_boot=5000)['contrast'].set_index('metric')
    row = result.loc['n_miat_spots']
    # Every fixed well mean is 6; wells-only support is {0}. Twelve Bernoulli
    # draws give CDF(-4)=79/4096 and CDF(-3)=299/4096, hence bounds -3 and 3.
    assert (row.exploratory_boot_lo, row.exploratory_boot_hi) == (-3, 3)


def test_two_wells_skip_bootstrap_and_explain_why(monkeypatch):
    from fishsuite.report import coupling_stats
    def forbidden(*args, **kwargs):
        raise AssertionError('bootstrap must not run with fewer than three wells per arm')
    monkeypatch.setattr(coupling_stats, '_bootstrap', forbidden)
    result = summary()
    contrast = result['contrast']
    assert contrast.exploratory_boot_lo.isna().all()
    assert contrast.exploratory_boot_hi.isna().all()
    assert contrast.exploratory_boot_valid_draws.eq(0).all()
    assert contrast.exploratory_boot_na_reason.eq('TOO_FEW_WELLS').all()
    assert 'with fewer than 3 wells per arm no interval is reported; read the individual well means' in '\n'.join(result['README'].description)


@pytest.mark.parametrize('metric', [
    'obs_minus_null_frac_miat_spots_qki_pos', 'obs_minus_null_frac_miat_footprint_area_qki_pos',
    'qki_at_spots_minus_nuclear', 'obs_minus_null_extra_statistic'])
def test_difference_scale_ratios_are_suppressed(metric):
    df = toy_table()
    df[metric] = np.where(df.condition.eq('control'), -3.5e-18, .2)
    level = df.copy()
    level['threshold_multiplier'] = .8
    result = summary(pd.concat([df, level], ignore_index=True))
    row = result['contrast'].set_index('metric').loc[metric]
    assert row.difference == pytest.approx(.2)
    assert pd.isna(row.ratio)
    assert row.ratio_na_reason == 'DIFFERENCE_SCALE'
    sensitivity = result['sensitivity'].set_index('metric').loc[metric]
    for level in ('0.8', '1'):
        assert pd.isna(sensitivity[f'ratio_multiplier_{level}'])
        assert sensitivity[f'ratio_na_reason_multiplier_{level}'] == 'DIFFERENCE_SCALE'


@pytest.mark.parametrize('control_mean', [0, -1, 1e-13, -1e-13])
def test_ratio_scale_control_nonpositive_or_near_zero(control_mean):
    df = toy_table()
    df['mean_nuclear_qki'] = df.mean_nuclear_qki.astype(float)
    df.loc[df.condition.eq('control'), 'mean_nuclear_qki'] = control_mean
    row = summary(df)['contrast'].set_index('metric').loc['mean_nuclear_qki']
    assert pd.isna(row.ratio)
    assert row.ratio_na_reason == 'CONTROL_NONPOSITIVE'


def test_zero_spot_nucleus_retained_for_counts_excluded_for_fraction():
    df = toy_table()
    rows = df.index[df.well.eq('C1')]
    df.loc[rows, 'n_miat_spots'] = [0, 4, 6, 10, 10]
    df.loc[rows, 'n_miat_spots_qki_pos'] = [0, 1, 2, 3, 3]
    fractions = np.array([np.nan, .25, 1/3, .3, .3])
    for metric in ('frac_miat_spots_qki_pos', 'frac_miat_footprint_area_qki_pos'):
        df.loc[rows, metric] = fractions
        df.loc[rows, 'obs_minus_null_' + metric] = fractions - .25
    row = summary(df)['per_well'].set_index('well').loc['C1']
    assert row.n_nuclei == 5
    assert row.n_nuclei_with_spots == 4
    assert row.n_finite_n_miat_spots == 5
    assert row.n_excluded_n_miat_spots == 0
    assert row.mean_n_miat_spots == 6
    assert row.mean_n_miat_spots_qki_pos == 1.8
    for metric, mean in (
        ('frac_miat_spots_qki_pos', 71/240), ('frac_miat_footprint_area_qki_pos', 71/240),
        ('obs_minus_null_frac_miat_spots_qki_pos', 11/240),
        ('obs_minus_null_frac_miat_footprint_area_qki_pos', 11/240),
    ):
        assert row[f'n_finite_{metric}'] == 4
        assert row[f'n_excluded_{metric}'] == 1
        assert row[f'mean_{metric}'] == pytest.approx(mean)
    assert row.n_finite_sat_frac_miat == 5
    assert row.n_excluded_sat_frac_miat == 0


def test_ratio_of_ratios_allows_tiny_positive_abundance_ratio():
    df = toy_table()
    df['n_miat_spots'] = np.where(df.condition.eq('control'), 1e14, 1)
    df['n_miat_spots_qki_pos'] = 1
    row = summary(df)['ratio_of_ratios'].set_index('metric').loc['ror_count']
    assert row.value == 1e14
    assert row.na_reason == ''


def test_ratio_of_ratios_zero_treated_abundance_has_denominator_reason():
    df = toy_table()
    df.loc[df.condition.eq('treated'), ['n_miat_spots', 'n_miat_spots_qki_pos']] = 0
    row = summary(df)['ratio_of_ratios'].set_index('metric').loc['ror_count']
    assert pd.isna(row.value)
    assert 'zero n_miat_spots treated/control ratio' in row.na_reason
    assert 'CONTROL_NONPOSITIVE' not in row.na_reason


def test_empty_well_falls_back_to_image_and_readme_records_it():
    df = toy_table()
    df['well'] = ''
    result = summary(df)
    assert set(result['per_well'].well) == {'C1.tif', 'C2.tif', 'T1.tif', 'T2.tif'}
    text = result['README'].to_string(index=False).lower()
    assert 'fallback' in text and 'image' in text and '20' in text
    assert 'single-plane' in text


def test_sensitivity_all_arms_and_missing_metric_values():
    df = toy_table()
    extra = df[df.well == 'C1'].copy()
    extra['condition'], extra['well'], extra['image'] = 'other', 'O1', 'O1.tif'
    df = pd.concat([df, extra], ignore_index=True)
    level = df.copy()
    level['threshold_multiplier'] = .8
    df = pd.concat([df, level], ignore_index=True)
    df.loc[df.well == 'C1', 'mean_nuclear_qki'] = np.nan
    result = summary(df)
    assert set(result['all_arms'].condition) == {'other'}
    assert len(result['all_arms']) == 2
    assert {'difference_multiplier_0.8', 'ratio_multiplier_0.8',
            'difference_multiplier_1', 'ratio_multiplier_1'} <= set(result['sensitivity'])
    row = result['contrast'].set_index('metric').loc['mean_nuclear_qki']
    assert pd.isna(row.control_mean_of_well_means)
    assert 'missing' in row.na_reason.lower()


def test_duplicate_nuclei_rejected():
    df = toy_table()
    with pytest.raises(ValueError, match='duplicate'):
        summary(pd.concat([df, df.iloc[[0]]], ignore_index=True))


def test_missing_well_at_sensitivity_threshold_is_not_silently_dropped():
    baseline = toy_table()
    level = baseline[baseline.well.ne('C2')].copy()
    level['threshold_multiplier'] = .8
    result = summary(pd.concat([baseline, level], ignore_index=True))
    row = result['sensitivity'].set_index('metric').loc['n_miat_spots']
    assert pd.isna(row['difference_multiplier_0.8'])
    assert 'missing' in row['na_reason_multiplier_0.8']
    absent = result['per_well'].query("well == 'C2' and threshold_multiplier == 0.8")
    assert len(absent) == 1
    assert absent.iloc[0].n_nuclei == 0
    assert pd.isna(absent.iloc[0].mean_n_miat_spots)


def test_report_artifacts(tmp_path):
    from click.testing import CliRunner
    from fishsuite.cli import cli
    from fishsuite.report.coupling_stats import METRICS
    from openpyxl import load_workbook
    source = tmp_path / 'association.csv'
    baseline = toy_table()
    sensitivity = baseline.copy()
    sensitivity['threshold_multiplier'] = .8
    sensitivity['miat_min_used'], sensitivity['qki_min_used'] = 8, 16
    pd.concat([baseline[baseline.well.ne('C2')], sensitivity], ignore_index=True).to_csv(source, index=False)
    out = tmp_path / 'report'
    config = tmp_path / 'colors.yaml'
    config.write_text('condition_order: [treated, control]\ngroup_colors:\n  treated: "#0072B2"\n  control: "#7F7F7F"\n', encoding='utf-8')
    result = CliRunner().invoke(cli, ['coupling', '--assoc-csv', str(source),
        '--treated', 'treated', '--control', 'control', '--out', str(out),
        '--config', str(config), '--seed', '4', '--n-boot', '10'])
    assert result.exit_code == 0, result.output + repr(result.exception)
    for name in ('coupling_summary.xlsx', 'command.log', 'versions.txt', 'figure_well_key.csv'):
        assert (out / name).stat().st_size > 0
    book = load_workbook(out / 'coupling_summary.xlsx', read_only=False)
    assert book.sheetnames == ['README', 'per_well', 'contrast', 'ratio_of_ratios',
                              'within_well_correlation', 'sensitivity', 'all_arms']
    assert book['per_well'].auto_filter.ref
    assert book['per_well'].freeze_panes == 'A3'
    readme_text = '\n'.join(str(cell.value) for row in book['README'] for cell in row if cell.value is not None)
    assert 'with fewer than 3 wells per arm no interval is reported; read the individual well means' in readme_text
    figure_files = list(out.rglob('*.png')) + list(out.rglob('*.svg'))
    assert figure_files
    for path in figure_files:
        assert path.stat().st_size > 0
        if path.suffix == '.svg':
            svg = path.read_text(encoding='utf-8')
            assert 'single-plane' in svg
            assert 'Filter:' in svg
            assert 'threshold_multiplier=' in svg
            assert '/'.join(source.parts[-2:]) in svg
            assert str(source) in (out / 'command.log').read_text(encoding='utf-8')
            import re
            for color in re.findall(r'#[0-9a-fA-F]{6}\b', svg):
                color = color.lower()
                assert color in {'#0072b2', '#7f7f7f'} or color[1:3] == color[3:5] == color[5:7]
    distribution_svg = next(p for p in figure_files if p.name == 'distribution_n_miat_spots_full.svg').read_text(encoding='utf-8')
    assert '#0072b2' in distribution_svg.lower()
    assert '#7f7f7f' in distribution_svg.lower()
    import xml.etree.ElementTree as ET
    label_x = {node.text: float(node.attrib['x']) for node in
               ET.fromstring(distribution_svg).iter('{http://www.w3.org/2000/svg}text')
               if node.text in ('treated', 'control')}
    assert label_x['treated'] < label_x['control']
    assert 'well 2: NA' in distribution_svg
    assert 'figure well key: control; well 2' in readme_text
    key = pd.read_csv(out / 'figure_well_key.csv')
    assert set(key.image) == {'C1.tif', 'C2.tif', 'T1.tif', 'T2.tif'}
    sensitivity_svg = next(p for p in figure_files if p.name == 'sensitivity_n_miat_spots_full.svg').read_text(encoding='utf-8')
    sensitivity_root = ET.fromstring(sensitivity_svg)
    x_axis = next(node for node in sensitivity_root.iter() if node.attrib.get('id') == 'matplotlib.axis_1')
    x_labels = [''.join(node.itertext()).strip() for node in x_axis.iter('{http://www.w3.org/2000/svg}text')]
    assert '0.8' in x_labels and '1' in x_labels
    sensitivity_text = ' '.join(' '.join(node.itertext()) for node in sensitivity_root.iter('{http://www.w3.org/2000/svg}text'))
    assert 'NA' in sensitivity_text and 'missing well' in sensitivity_text
    from PIL import Image
    with Image.open(next(p for p in figure_files if p.suffix == '.png')) as image:
        assert image.info['dpi'][0] == pytest.approx(600, abs=1)
    for kind in ('distribution', 'well_means', 'sensitivity'):
        for metric in METRICS:
            for view in ('full', 'focus'):
                for ext in ('png', 'svg'):
                    assert any(p.name == f'{kind}_{metric}_{view}.{ext}' for p in figure_files)
    for stem in ('precision_funnel',):
        for view in ('full', 'focus'):
            for ext in ('png', 'svg'):
                assert any(p.name == f'{stem}_{view}.{ext}' for p in figure_files)
    for abundance in ('n_miat_spots', 'integrated_nuclear_miat'):
        for association in ('obs_minus_null_frac_miat_spots_qki_pos', 'qki_at_spots_minus_nuclear'):
            for view in ('full', 'focus'):
                for ext in ('png', 'svg'):
                    assert any(p.name == f'scatter_{abundance}_vs_{association}_{view}.{ext}' for p in figure_files)
    book.close()


def test_cli_registration():
    from click.testing import CliRunner
    from fishsuite.cli import cli
    result = CliRunner().invoke(cli, ['coupling', '--help'])
    assert result.exit_code == 0, result.output
    assert '--assoc-csv' in result.output
    assert '--treated' in result.output
    assert '--control' in result.output


def test_well_markers_are_ordered_within_each_arm():
    from fishsuite.report.coupling import _well_markers
    mapping = _well_markers(toy_table(), ['control', 'treated'])
    assert set(mapping) == {('control', 'C1'), ('control', 'C2'), ('treated', 'T1'), ('treated', 'T2')}
    assert mapping == {('control', 'C1'): 'o', ('control', 'C2'): '^',
                       ('treated', 'T1'): 'o', ('treated', 'T2'): '^'}
    assert mapping == _well_markers(toy_table().sample(frac=1, random_state=3), ['control', 'treated'])


def test_red_green_guard_and_resolved_colors(tmp_path):
    from fishsuite.report.coupling import _configuration
    config = tmp_path / 'forbidden.yaml'
    config.write_text('group_colors:\n  treated: "#D55E00"\n  control: "#009E73"\n', encoding='utf-8')
    with pytest.raises(ValueError, match='red and green'):
        _configuration(config, ['control', 'treated'], 'treated')
    order, colors = _configuration(None, ['control', 'treated', 'other'], 'treated')
    assert order == ['control', 'treated', 'other']
    assert {key: colors[key].upper() for key in order} == {
        'control': '#7F7F7F', 'treated': '#56B4E9', 'other': '#7F7F7F'}


def test_six_arm_distribution_layout_markers_and_bars():
    import matplotlib.pyplot as plt
    from matplotlib.markers import MarkerStyle
    from fishsuite.report import coupling
    template = toy_table().iloc[0].to_dict()
    arms = ['g2 noDox (control)', 'g2 +Dox (MIAT OE)', 'g3 noDox (control)',
            'g3 +Dox (MIAT OE)', 'WT noDox (control)', 'WT +Dox (control)']
    rows = [dict(template, condition=arm, well=f'{well}.vsi', image=f'{well}.vsi', nucleus_id=n,
                 n_miat_spots=well + n) for arm in arms for well in range(1, 6) for n in (1, 2)]
    df = pd.DataFrame(rows)
    coupling.figlib.set_style()
    fig, ax = plt.subplots(figsize=(1.1 * len(arms) + 1.5, 4.8))
    coupling._distribution(ax, df, df, 'n_miat_spots', arms,
                           {arm: '#7F7F7F' for arm in arms}, np.random.default_rng(0))
    fig.subplots_adjust(left=1.1 / fig.get_figwidth(), right=.97, bottom=.32, top=.88)
    fig.canvas.draw()
    boxes = [label.get_window_extent(fig.canvas.get_renderer()) for label in ax.get_xticklabels()]
    assert all(not a.overlaps(b) for i, a in enumerate(boxes) for b in boxes[i+1:])
    assert len(ax.patches) == len(arms)
    assert all(p.get_facecolor()[-1] == pytest.approx(.35) for p in ax.patches)
    assert all(p.get_edgecolor()[-1] == 1 for p in ax.patches)
    assert all(p.get_height() == pytest.approx(4.5) for p in ax.patches)
    assert not any(type(c).__name__ == 'ErrorbarContainer' for c in ax.containers)
    markers = coupling._well_markers(df, arms)
    assert [markers[arms[0], f'{i}.vsi'] for i in range(1, 6)] == ['o', '^', 's', 'D', 'o']
    assert all(markers[arm, '1.vsi'] == 'o' for arm in arms)
    legend = [text.get_text() for text in fig.legends[0].get_texts()]
    assert legend == ['well 1', 'well 2', 'well 3', 'well 4', 'well 5',
                      'well mean', 'arm mean of well means']
    assert not any('.vsi' in text for text in legend)
    # First two collections are nuclei and well mean; the same circle path is used.
    circle = MarkerStyle('o')
    expected = circle.get_path().transformed(circle.get_transform()).vertices
    assert np.array_equal(ax.collections[0].get_paths()[0].vertices, expected)
    assert np.array_equal(ax.collections[1].get_paths()[0].vertices, expected)
    assert ax.collections[0].get_alpha() == .45
    assert ax.collections[1].get_zorder() > ax.collections[0].get_zorder()
    plt.close(fig)


@pytest.mark.parametrize('metric,expected', [
    ('n_miat_spots', 'MIAT puncta per nucleus'),
    ('n_miat_spots_per_100um2', 'MIAT puncta per 100 µm² nuclear area'),
    ('integrated_nuclear_miat', 'Total nuclear MIAT intensity (a.u.)'),
    ('frac_miat_spots_qki_pos', 'Fraction of MIAT puncta with QKI ≥ threshold (single plane)'),
    ('obs_minus_null_frac_miat_spots_qki_pos', 'QKI-positive MIAT puncta: observed − chance (single plane)'),
    ('frac_miat_footprint_area_qki_pos', 'Fraction of MIAT punctum area that is QKI-positive (single plane)'),
    ('obs_minus_null_frac_miat_footprint_area_qki_pos', 'QKI-positive MIAT area: observed − chance (single plane)'),
    ('mean_qki_at_miat_spots', 'Mean QKI intensity at MIAT puncta (a.u.)'),
    ('qki_at_spots_minus_nuclear', 'QKI at MIAT puncta − nuclear mean QKI (a.u.)'),
    ('mean_nuclear_qki', 'Mean nuclear QKI intensity (a.u.)'),
    ('nuclear_area_um2', 'Nuclear area (µm²)'),
    ('unknown_metric', 'unknown_metric'),
])
def test_plain_language_metric_labels(metric, expected):
    from fishsuite.report.coupling import _label
    assert _label(metric) == expected


@pytest.mark.parametrize('metric', ['obs_minus_null_frac_miat_spots_qki_pos',
                                    'obs_minus_null_frac_miat_footprint_area_qki_pos'])
def test_observed_minus_chance_reference_and_shared_threshold_bounds(metric):
    import matplotlib.pyplot as plt
    from fishsuite.report import coupling
    base = toy_table()
    level = base.copy()
    level['threshold_multiplier'] = .8
    level[metric] = 3
    df = pd.concat([base, level])
    fig, ax = plt.subplots()
    coupling._distribution(ax, df, base, metric, ['control', 'treated'],
                           {'control': '#7F7F7F', 'treated': '#0072B2'}, np.random.default_rng(0))
    chance = [line for line in ax.lines if line.get_label() == 'chance']
    assert len(chance) == 1
    assert list(chance[0].get_ydata()) == [0, 0]
    assert chance[0].get_linestyle() == '--'
    ax._replicate_simple_axis('focus')
    first = ax.get_ylim()
    assert first[1] > 3
    assert first[0] <= 0 <= first[1]
    fig2, ax2 = plt.subplots()
    coupling._distribution(ax2, df, level, metric, ['control', 'treated'],
                           {'control': '#7F7F7F', 'treated': '#0072B2'}, np.random.default_rng(0))
    ax2._replicate_simple_axis('focus')
    assert ax2.get_ylim() == first
    plt.close(fig)
    plt.close(fig2)


def test_missing_well_suppresses_arm_bar():
    import matplotlib.pyplot as plt
    from fishsuite.report import coupling
    df = toy_table()
    base = df.loc[df.well.ne('C2')]
    fig, ax = plt.subplots()
    coupling._distribution(ax, df, base, 'n_miat_spots', ['control', 'treated'],
                           {'control': '#7F7F7F', 'treated': '#0072B2'}, np.random.default_rng(0))
    assert len(ax.patches) == 1
    assert ax.patches[0].get_x() + ax.patches[0].get_width() / 2 == 1
    assert ax.patches[0].get_height() == 12
    plt.close(fig)


def test_nested_preset_condition_colors_and_order(tmp_path):
    from fishsuite.report.coupling import _configuration
    config = tmp_path / 'preset.yaml'
    config.write_text('conditions:\n  condition_order: [treated, control]\n  group_colors:\n    treated: "#0072B2"\n    control: "#7F7F7F"\n', encoding='utf-8')
    order, colors = _configuration(config, ['control', 'treated'], 'treated')
    assert order == ['treated', 'control']
    assert colors['treated'] == '#0072B2'


def test_footer_is_two_lines_and_title_one_line(monkeypatch, tmp_path):
    import matplotlib.pyplot as plt
    from fishsuite.report import coupling
    coupling.figlib.set_style()
    fig, ax = plt.subplots(figsize=(8.1, 4.8))
    ax.set_ylabel(coupling._label('n_miat_spots'))
    coupling._axis_variants(ax, [1, 2])
    monkeypatch.setattr(coupling.figlib, 'save', lambda *args: None)
    coupling._finish(fig, ax, tmp_path, 'test', 'MIAT puncta per nucleus',
                     'Filter: threshold_multiplier=1; finite metric; N=0 retained when defined',
                     'Input: SMOKE2_PROVISIONAL_LEVELS_NOT_RESULTS_20260917-175544/qki_association_per_nucleus.csv; single-plane', [])
    fig.canvas.draw()
    assert fig._suptitle.get_text() == 'MIAT puncta per nucleus'
    footer = [t for t in fig.texts if t.get_text().startswith('Descriptive;')]
    assert len(footer) == 1
    assert footer[0].get_text().count('\n') <= 1
    assert footer[0].get_fontsize() >= 7
    assert footer[0].get_window_extent(fig.canvas.get_renderer()).x1 <= fig.bbox.x1
    plt.close(fig)


def allcols_table():
    data = toy_table()
    for metric in ('mean_nuclear_miat', 'n_miat_spots_per_100um2', 'qki_pos_area_frac',
                   'miat_pos_area_frac', 'frac_qki_area_on_miat_area',
                   'frac_qki_area_on_miat_footprints', 'miat_footprint_area_frac',
                   'median_footprint_area_px', 'null_valid_center_count',
                   'null_mean_frac_miat_spots_qki_pos', 'null_sd_frac_miat_spots_qki_pos',
                   'exp_frac_qki_area_on_miat_area', 'obs_minus_exp_frac_qki_area_on_miat_area'):
        data[metric] = data.n_miat_spots / 100
    data['arbitrary numeric measurement (a.u.)'] = data.well.map({'C1': 10., 'C2': 20., 'T1': 30., 'T2': 40.})
    data['all_missing_numeric'] = np.nan
    data['nullable_integer_measurement'] = pd.Series([1, 2, pd.NA] + list(range(4, 21)), dtype='Int64')
    data['obs_minus_null_extra'] = np.where(data.condition.eq('control'), -.01, .02)
    data['sample_id'] = 55
    data['extra_min_used'] = 60
    data['n_null_effective'] = 200
    data['na_reason_numeric'] = 7
    data['well_from_image'] = False
    data['unrelated_boolean_flag'] = True
    data['text_measurement'] = 'not numeric'
    data['nuclear_area_um2'] = data.nuclear_area_um2.astype(str)
    return data


def test_allcols_numeric_discovery_and_literal_contrasts():
    from fishsuite.report.coupling_stats import METRICS
    from fishsuite.core.qki_association import COLUMN_DEFINITIONS
    data = allcols_table()
    level = data.copy()
    level['threshold_multiplier'] = .8
    result = summary(pd.concat([data, level], ignore_index=True))
    numeric_extras = {
        'mean_nuclear_miat', 'n_miat_spots_per_100um2', 'qki_pos_area_frac', 'miat_pos_area_frac',
        'frac_qki_area_on_miat_area', 'frac_qki_area_on_miat_footprints', 'miat_footprint_area_frac',
        'median_footprint_area_px', 'null_valid_center_count', 'null_mean_frac_miat_spots_qki_pos',
        'null_sd_frac_miat_spots_qki_pos', 'exp_frac_qki_area_on_miat_area',
        'obs_minus_exp_frac_qki_area_on_miat_area', 'arbitrary numeric measurement (a.u.)', 'all_missing_numeric',
        'nullable_integer_measurement', 'obs_minus_null_extra',
    }
    assert set(result['contrast'].metric) == set(METRICS) | numeric_extras
    assert set(result['sensitivity'].metric) == set(METRICS) | numeric_extras
    per_well = result['per_well'].set_index(['well', 'threshold_multiplier'])
    arbitrary = 'arbitrary numeric measurement (a.u.)'
    assert per_well.loc[('C1', 1), 'mean_' + arbitrary] == 10
    row = result['contrast'].set_index('metric').loc[arbitrary]
    assert row.control_well_1_mean == 10
    assert row.control_well_2_mean == 20
    assert row.control_mean_of_well_means == 15
    assert row.treated_mean_of_well_means == 35
    assert row.difference == 20
    assert row.ratio == pytest.approx(7 / 3)
    assert result['sensitivity'].set_index('metric').loc[arbitrary, 'difference_multiplier_0.8'] == 20
    assert per_well.loc[('C1', 1), 'n_excluded_all_missing_numeric'] == 5
    assert per_well.loc[('C1', 1), 'mean_nullable_integer_measurement'] == 3
    assert per_well.loc[('C1', 1), 'n_excluded_nullable_integer_measurement'] == 1
    assert pd.isna(result['contrast'].set_index('metric').loc['all_missing_numeric', 'difference'])
    readme = result['README'].set_index('topic').description
    assert readme[arbitrary] == arbitrary
    assert readme['frac_qki_area_on_miat_footprints'] == COLUMN_DEFINITIONS['frac_qki_area_on_miat_footprints']
    for metric in ('obs_minus_null_frac_miat_spots_qki_pos', 'qki_at_spots_minus_nuclear', 'obs_minus_null_extra'):
        assert result['contrast'].set_index('metric').loc[metric, 'ratio_na_reason'] == 'DIFFERENCE_SCALE'


def test_allcols_plot_selection_is_bounded():
    from fishsuite.report.coupling import _plot_metrics
    from fishsuite.report.coupling_stats import METRICS
    extras = ['mean_nuclear_miat', 'n_miat_spots_per_100um2']
    available = list(allcols_table().columns) + ['obs_minus_null_extra', 'sat_frac_extra']
    assert _plot_metrics(available) == METRICS + extras
    assert _plot_metrics(pd.Series(available)) == METRICS + extras
    assert _plot_metrics(['unknown_numeric', 'mean_nuclear_miat']) == ['mean_nuclear_miat']


@pytest.mark.parametrize('metric,expected', [
    ('mean_nuclear_miat', 'Mean nuclear MIAT intensity (a.u.)'),
    ('n_miat_spots_per_100um2', 'MIAT puncta per 100 µm² nuclear area'),
    ('qki_pos_area_frac', 'Fraction of nuclear area QKI-positive'),
    ('miat_pos_area_frac', 'Fraction of nuclear area MIAT-positive'),
    ('frac_qki_area_on_miat_area', 'Fraction of QKI-positive area on MIAT-positive pixels (raw)'),
    ('frac_qki_area_on_miat_footprints', 'Fraction of QKI-positive area on MIAT puncta (raw, uncorrected)'),
])
def test_allcols_additional_plain_language_labels(metric, expected):
    from fishsuite.report.coupling import _label
    assert _label(metric) == expected


def test_allcols_workbook_expands_without_expanding_every_numeric_plot(tmp_path, monkeypatch):
    import matplotlib.pyplot as plt
    from fishsuite.report import coupling
    from fishsuite.report.coupling_stats import METRICS
    source = tmp_path / 'all_columns.csv'
    allcols_table().to_csv(source, index=False)
    stems = []
    def capture_save(fig, destination, stem, *args):
        Path(destination).mkdir(parents=True, exist_ok=True)
        stems.append(stem)
        plt.close(fig)
    monkeypatch.setattr(coupling.figlib, 'save', capture_save)
    output = coupling.build_coupling(source, 'treated', 'control', tmp_path / 'report', n_boot=10)
    contrast = pd.read_excel(output / 'coupling_summary.xlsx', sheet_name='contrast', header=1)
    assert 'arbitrary numeric measurement (a.u.)' in set(contrast.metric)
    expected = {f'{kind}_{metric}' for kind in ('distribution', 'well_means', 'sensitivity')
                for metric in METRICS + ['mean_nuclear_miat', 'n_miat_spots_per_100um2']}
    expected |= {f'scatter_{abundance}_vs_{association}' for abundance in coupling.ABUNDANCE
                 for association in coupling.ASSOCIATION}
    expected.add('precision_funnel')
    assert set(stems) == expected
    assert len(stems) == len(expected)
