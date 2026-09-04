"""``fishsuite report`` end to end, on a synthetic run.

What this holds: the workbook opens with the plain sheet names Brian asked for,
every sheet starts with a description row rather than a bare header, no sheet
name or column carries a ``Q1``-style code, the readout names the producing run
on its first line, and the run directory is not modified.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from fishsuite.report import workbook as wb
from fishsuite.report.build import build_report


def _synthetic_run(tmp_path: Path) -> Path:
    rng = np.random.default_rng(0)
    run = tmp_path / "RUN_SYNTH_2026-09-04"
    run.mkdir()
    img, nuc, thr = [], [], []
    means = {"WT_1": 4.0, "WT_2": 5.0, "WT_3": 6.0,
             "KO_1": 10.0, "KO_2": 11.0, "KO_3": 12.0}
    for well, mu in means.items():
        for f in (1, 2, 3):
            image = f"{well}_{f:02d}.vsi"
            img.append({"image": image, "condition": well, "secondary_only": False,
                        "total_spots_rna1": mu * 6,
                        "protein_pooled_rotation_enrichment_at_rna1_spots": 1 + mu / 100,
                        "rna_bigfish_log_threshold": 36.0,
                        "protein_bigfish_log_threshold": 15.0})
            thr.append({"image": image, "rna_bigfish_log_threshold": 36.0,
                        "protein_bigfish_log_threshold": 15.0,
                        "protein_threshold_value": 900.0})
            for k in range(6):
                nuc.append({"image": image, "nucleus_id": k,
                            "n_spots_rna1": float(mu + rng.normal(0, 0.5)),
                            "nuclear_spot_count": float(mu - 1),
                            "nuclear_spot_fraction": 0.9,
                            "protein_nuclear_mean": 800.0 + mu,
                            "protein_nc_ratio": 1.4,
                            "nucleus_area_px": 16000.0, "voxel_xy_um": 0.065,
                            "rotation_null_usable": True,
                            "protein_rotation_enrichment_at_rna1_spots": 1 + mu / 100})
    for f in (8, 9):
        image = f"Sec-Only_{f:02d}.vsi"
        img.append({"image": image, "condition": "Sec-Only", "secondary_only": True,
                    "total_spots_rna1": 2.0,
                    "rna_bigfish_log_threshold": 36.0,
                    "protein_bigfish_log_threshold": 15.0})
        thr.append({"image": image, "rna_bigfish_log_threshold": 36.0,
                    "protein_bigfish_log_threshold": 15.0,
                    "protein_threshold_value": 900.0})
        n_nuc = 12 if f == 8 else 3       # the second control field trips the floor rule
        for k in range(n_nuc):
            nuc.append({"image": image, "nucleus_id": k, "n_spots_rna1": 0.2,
                        "nuclear_spot_count": 0.0, "nuclear_spot_fraction": 0.0,
                        "protein_nuclear_mean": 100.0, "protein_nc_ratio": 1.0,
                        "nucleus_area_px": 16000.0, "voxel_xy_um": 0.065,
                        "rotation_null_usable": False,
                        "protein_rotation_enrichment_at_rna1_spots": np.nan})
    pd.DataFrame(img).to_csv(run / "per_image_summary.csv", index=False)
    pd.DataFrame(nuc).to_csv(run / "nuclei_metrics.csv", index=False)
    pd.DataFrame(thr).to_csv(run / "thresholds.csv", index=False)
    (run / "run_config.json").write_text(json.dumps({
        "SEGMENTATION_BACKEND": "cellpose",
        "config_resolved": {
            "channels": {"analysis_mode": "rna_protein", "rna_label": "BIN1 intron",
                         "antibody_label": "RNASEH2B", "dapi_label": "DAPI"},
            "nuclei": {"backend": "cellpose", "cellpose_model_type": "cpsam_v2",
                       "min_area_px": 16000}}}), encoding="utf-8")
    (run / "versions.txt").write_text("cellpose: 4.2.1.1\n", encoding="utf-8")
    return run


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("report")
    run = _synthetic_run(tmp)
    before = {p.name: p.stat().st_mtime_ns for p in run.iterdir() if p.is_file()}
    r = build_report(run_dir=run, groups=["WT=WT_1,WT_2,WT_3", "QKI-KO=KO_1,KO_2,KO_3"],
                     reference="WT", stamp="test", make_figures=False,
                     coloc_panel=False, argv=["fishsuite", "report"])
    after = {p.name: p.stat().st_mtime_ns for p in run.iterdir() if p.is_file()}
    return r, run, before, after


def test_the_run_directory_is_not_modified(built):
    _, _, before, after = built
    assert before == after


def test_the_expected_products_are_written(built):
    r, _, _, _ = built
    out = r["out_dir"]
    for name in ("REPORT.xlsx", "READOUT.md", "per_well.csv", "contrasts.csv",
                 "versions.txt", "command.log"):
        p = out / name
        assert p.is_file() and p.stat().st_size > 0, f"{name} missing or empty"


def test_sheet_names_are_plain_language_and_in_the_expected_order(built):
    r, _, _, _ = built
    import openpyxl
    book = openpyxl.load_workbook(r["xlsx"], read_only=False)
    assert book.sheetnames == wb.SHEET_ORDER


def test_no_sheet_name_or_column_carries_a_q_code(built):
    r, _, _, _ = built
    import openpyxl
    book = openpyxl.load_workbook(r["xlsx"])
    q = re.compile(r"^Q[1-9]\b|_Q[1-9]_|\bQ[1-9]$")
    for name in book.sheetnames:
        assert not q.search(name), f"sheet name {name} carries a Q code"
        ws = book[name]
        for cell in next(ws.iter_rows(min_row=2, max_row=2)):
            if cell.value:
                assert not q.search(str(cell.value)), f"{name}: column {cell.value}"


def test_every_sheet_opens_with_a_description_row(built):
    r, _, _, _ = built
    import openpyxl
    book = openpyxl.load_workbook(r["xlsx"])
    for name in book.sheetnames:
        desc = book[name].cell(row=1, column=1).value
        assert desc and len(str(desc)) > 80, f"{name} has no description row"
        # Two to three sentences, not a label.
        assert str(desc).count(".") >= 2, f"{name} description is not prose"


def test_the_readout_names_the_run_on_its_first_line(built):
    r, run, _, _ = built
    first = (r["out_dir"] / "READOUT.md").read_text(encoding="utf-8").splitlines()[0]
    assert str(run) in first


def test_contrasts_carry_the_raw_p_the_adjusted_p_and_the_mde(built):
    r, _, _, _ = built
    c = r["contrasts"]
    for col in ("p_welch", "significant_raw_0p05", "p_welch_holm_within_family",
                "mde_hedges_g_at_family_alpha", "hedges_g", "ci_low", "ci_high"):
        assert col in c.columns, f"contrasts is missing {col}"
    gated = c[c["in_holm_family"]].dropna(subset=["p_welch"])
    assert len(gated) > 0
    assert gated["mde_hedges_g_at_family_alpha"].notna().any()


def test_the_control_field_below_the_nucleus_floor_is_excluded_by_rule(built):
    r, _, _, _ = built
    sec = r["sec"]
    row = sec[sec["field"] == "Sec-Only_09.vsi"].iloc[0]
    assert bool(row["excluded"])
    assert "fewer than" in row["exclusion_reason"]
    kept = sec[sec["field"] == "Sec-Only_08.vsi"].iloc[0]
    assert not bool(kept["excluded"])


def test_an_operator_exclusion_records_its_reason(tmp_path):
    run = _synthetic_run(tmp_path)
    r = build_report(run_dir=run, groups=["WT=WT_1,WT_2,WT_3", "QKI-KO=KO_1,KO_2,KO_3"],
                     reference="WT", stamp="excl", make_figures=False, coloc_panel=False,
                     exclude_fields={"WT_1_01.vsi": "the field was out of focus"},
                     argv=["fishsuite", "report"])
    prov = pd.read_excel(r["xlsx"], sheet_name="Run provenance", header=1)
    line = prov[prov["Item"] == "fields excluded"]["Value"].iloc[0]
    assert "WT_1_01.vsi" in line and "out of focus" in line
    # The excluded field's nuclei are gone from the tested well.
    n = r["well"]
    wt1 = n[(n["endpoint"] == "rna1_spots_per_nucleus") & (n["well_id"] == "WT_1")].iloc[0]
    assert int(wt1["n_fields"]) == 2


def test_the_reference_group_must_exist(tmp_path):
    run = _synthetic_run(tmp_path)
    from fishsuite.report.aggregate import ReportInputError
    with pytest.raises(ReportInputError, match="reference group"):
        build_report(run_dir=run, groups=["WT=WT_1,WT_2,WT_3"], reference="Nope",
                     stamp="bad", make_figures=False, coloc_panel=False)
