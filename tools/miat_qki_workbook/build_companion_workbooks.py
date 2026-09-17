#!/usr/bin/env python3
"""Build memory-safe MIAT×QKI companion Excel workbooks from staged CSV tables."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import sys
import warnings
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.reference import Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo


ANALYSIS_KEYS = [
    "summary",
    "fov",
    "set",
    "endpoint_inference",
    "cartesian",
    "ratio_of_ratios",
    "fov_correlations",
    "set_correlations",
    "correlation_inference",
    "representatives",
    "controls",
    "sampling_audit",
    "image_qc",
    "nucleus_qc",
    "threshold_sensitivity",
]

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "666666"
THIN_GRAY = Side(style="thin", color="D9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color=WHITE)
BODY_FONT = Font(name="Arial", size=10, color="222222")
TITLE_FONT = Font(name="Arial", size=18, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Arial", size=11, italic=True, color=GRAY)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILL = PatternFill("solid", fgColor=BLUE)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=False)
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)


class CompanionWorkbookError(RuntimeError):
    """Base error for clear CLI failures."""


class OutputExistsError(CompanionWorkbookError):
    """Raised when a requested output directory already exists."""


class MissingInputError(CompanionWorkbookError):
    """Raised when required staged inputs or columns are missing."""


DIGIT_INTEGER_RE = re.compile(r"^[+-]?\d+$")
NUMBER_RE = re.compile(
    r"^[+-]?(?:(?:\d+\.\d*)|(?:\.\d+)|(?:\d+))(?:[eE][+-]?\d+)?$"
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_cell_value(value: str) -> Any:
    """Convert safe short numerics/bools while preserving loss-prone text exactly."""
    if value == "":
        return None
    if value != value.strip():
        return value
    lower = value.lower()
    if lower == "true":
        return True
    if lower == "false":
        return False
    if DIGIT_INTEGER_RE.fullmatch(value):
        unsigned = value.lstrip("+-")
        if len(unsigned) > 1 and unsigned.startswith("0"):
            return value
        if len(unsigned) <= 15:
            return int(value)
        return value
    if NUMBER_RE.fullmatch(value):
        try:
            number = Decimal(value)
        except InvalidOperation:
            return value
        significant_digits = len(number.as_tuple().digits)
        if significant_digits <= 15 and number.is_finite():
            converted = float(value)
            if math.isfinite(converted):
                return converted
    return value


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def merge_companion_config(
    schema: dict[str, Any], companion_config_path: Path | None
) -> dict[str, Any]:
    if companion_config_path is None:
        return schema
    companion = load_json(companion_config_path)
    merged = dict(schema)
    merged["companionWorkbooks"] = companion.get("companionWorkbooks", companion)
    return merged


def table_map(schema: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {table["key"]: table for table in schema.get("tables", [])}


def validate_inputs(
    input_dir: Path, schema: dict[str, Any], required_keys: Iterable[str]
) -> dict[str, dict[str, Any]]:
    tables = table_map(schema)
    missing_keys = [key for key in required_keys if key not in tables]
    if missing_keys:
        raise MissingInputError(f"Schema is missing required tables: {missing_keys}")
    validated: dict[str, dict[str, Any]] = {}
    for key in required_keys:
        spec = tables[key]
        path = input_dir / spec["file"]
        if not path.is_file():
            raise MissingInputError(f"Required staged CSV is missing: {path}")
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle)
            try:
                headers = next(reader)
            except StopIteration as exc:
                raise MissingInputError(f"Required staged CSV is empty: {path}") from exc
        if len(headers) != len(set(headers)):
            raise MissingInputError(f"Duplicate headers in {path}")
        required_columns = spec.get("requiredColumns", [])
        missing_columns = [column for column in required_columns if column not in headers]
        if missing_columns:
            raise MissingInputError(
                f"Required columns missing from {path}: {missing_columns}"
            )
        if len(headers) > 16384:
            raise MissingInputError(f"Excel column limit exceeded by {path}")
        validated[key] = {"spec": spec, "path": path, "headers": headers}
    return validated


def configure_workbook_defaults(workbook: Workbook) -> None:
    normal = workbook._named_styles[0]
    normal.font = BODY_FONT
    normal.alignment = BODY_ALIGNMENT


def number_format_for_header(header: str) -> str:
    lower = header.lower()
    if "percent_change" in lower:
        return '0.00"%"'
    if lower.startswith("p_") or lower.endswith("_p") or "p_value" in lower or "p_two_sided" in lower or "permutation_p" in lower:
        return "0.000E+00"
    if any(token in lower for token in ("fraction", "coverage", "ratio", "correlation", "enrichment", "coefficient", "mean", "median", "sd", "ci95", "difference", "intensity", "mass")):
        return "0.0000"
    if any(token in lower for token in ("count", "n_", "_n", "_id", "z_", "_px", "bytes", "seed", "replicate", "fov")):
        return "0"
    return "0.###############"


def inferred_type(header: str, sample_values: list[Any] | None = None) -> str:
    lower = header.lower()
    if "p_value" in lower or "p_two_sided" in lower or "permutation_p" in lower:
        return "p-value"
    if "percent_change" in lower:
        return "percentage-point numeric"
    if any(token in lower for token in ("path", "source", "note", "status", "reason", "uid", "key", "method", "endpoint", "population", "cohort", "arm", "condition", "slide")):
        return "text"
    if sample_values:
        nonempty = [value for value in sample_values if value is not None]
        if nonempty and all(isinstance(value, bool) for value in nonempty):
            return "boolean"
        if nonempty and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in nonempty):
            return "numeric"
    return "mixed / source-preserved"


def description_for_column(spec: dict[str, Any], header: str) -> str:
    configured = spec.get("columns", {}).get(header, {})
    if configured.get("definition"):
        return configured["definition"]
    return header.replace("_", " ").strip().capitalize() + "."


def sensible_width(header: str, values: Iterable[Any] = ()) -> float:
    lower = header.lower()
    if any(token in lower for token in ("path", "source", "note", "reason", "definition")):
        cap = 52
    elif any(token in lower for token in ("endpoint", "measurement", "method", "biological_set", "population")):
        cap = 36
    else:
        cap = 26
    width = len(header) + 2
    for value in values:
        if value is not None:
            width = max(width, min(len(str(value)) + 2, cap))
    return max(9, min(width, cap))


def set_sheet_page_defaults(worksheet) -> None:
    worksheet.freeze_panes = None
    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:1"


def safe_text_cell(cell, value: Any) -> None:
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def style_header_cell(cell) -> None:
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGNMENT
    cell.border = Border(bottom=THIN_GRAY)


def apply_body_cell_style(cell, header: str, value: Any) -> None:
    cell.font = BODY_FONT
    cell.alignment = WRAP_ALIGNMENT if any(
        token in header.lower() for token in ("note", "reason", "path", "source", "definition")
    ) else BODY_ALIGNMENT
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = number_format_for_header(header)


def add_excel_table(worksheet, table_name: str, headers: list[str], data_rows: int) -> str:
    if not data_rows:
        return ""
    last_column = get_column_letter(len(headers))
    reference = f"A1:{last_column}{data_rows + 1}"
    table = Table(displayName=table_name, ref=reference)
    table.tableColumns = [
        TableColumn(id=index, name=str(name))
        for index, name in enumerate(headers, start=1)
    ]
    table.autoFilter = AutoFilter(ref=reference)
    table.tableStyleInfo = TABLE_STYLE
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="In write-only mode you must add table columns manually",
            category=UserWarning,
        )
        worksheet.add_table(table)
    # The formal table's own AutoFilter is the single authoritative filter.
    # Duplicating the same range as a worksheet-level filter creates an
    # overlapping _FilterDatabase name that native Excel rejects.
    return reference


def write_readme_sheet(
    worksheet,
    workbook_meta: dict[str, Any],
    role: str,
    generated_utc: str,
    extra_rows: list[tuple[str, str]],
) -> int:
    set_sheet_page_defaults(worksheet)
    rows = [
        (workbook_meta.get("title", "MIAT×QKI exact-footprint analysis"), ""),
        (workbook_meta.get("subtitle", ""), ""),
        ("Workbook role", role),
        ("Purpose", workbook_meta.get("purpose", "")),
        ("Analysis convention", workbook_meta.get("analysisConvention", "")),
        ("Interpretation guard", workbook_meta.get("interpretationGuard", "")),
        ("Provenance", workbook_meta.get("provenance", "")),
        ("Generated UTC", generated_utc),
    ] + extra_rows
    notes = workbook_meta.get("readmeNotes", [])
    if notes:
        rows.append(("Review notes", ""))
        rows.extend((f"Note {index}", note) for index, note in enumerate(notes, start=1))

    for row_index, (label, value) in enumerate(rows, start=1):
        worksheet.cell(row=row_index, column=1, value=label)
        worksheet.cell(row=row_index, column=2, value=value)
        worksheet.cell(row=row_index, column=1).font = BODY_FONT
        worksheet.cell(row=row_index, column=2).font = BODY_FONT
        worksheet.cell(row=row_index, column=1).alignment = WRAP_ALIGNMENT
        worksheet.cell(row=row_index, column=2).alignment = WRAP_ALIGNMENT
        if row_index == 1:
            worksheet.cell(row=row_index, column=1).font = TITLE_FONT
        elif row_index == 2:
            worksheet.cell(row=row_index, column=1).font = SUBTITLE_FONT
        elif label == "Review notes":
            worksheet.cell(row=row_index, column=1).font = Font(name="Arial", size=11, bold=True, color=NAVY)
            worksheet.cell(row=row_index, column=1).fill = SECTION_FILL
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 110
    for row in range(1, len(rows) + 1):
        worksheet.row_dimensions[row].height = 30 if row > 2 else 28
    return len(rows)


def append_write_only_readme(
    worksheet,
    workbook_meta: dict[str, Any],
    role: str,
    generated_utc: str,
    extra_rows: list[tuple[str, str]],
) -> int:
    set_sheet_page_defaults(worksheet)
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 110
    rows = [
        (workbook_meta.get("title", "MIAT×QKI exact-footprint analysis"), ""),
        (workbook_meta.get("subtitle", ""), ""),
        ("Workbook role", role),
        ("Purpose", workbook_meta.get("purpose", "")),
        ("Analysis convention", workbook_meta.get("analysisConvention", "")),
        ("Interpretation guard", workbook_meta.get("interpretationGuard", "")),
        ("Provenance", workbook_meta.get("provenance", "")),
        ("Generated UTC", generated_utc),
    ] + extra_rows
    notes = workbook_meta.get("readmeNotes", [])
    if notes:
        rows.append(("Review notes", ""))
        rows.extend((f"Note {index}", note) for index, note in enumerate(notes, start=1))
    for row_index, (label, value) in enumerate(rows, start=1):
        first = WriteOnlyCell(worksheet, value=label)
        second = WriteOnlyCell(worksheet, value=value)
        first.font = BODY_FONT
        second.font = BODY_FONT
        first.alignment = WRAP_ALIGNMENT
        second.alignment = WRAP_ALIGNMENT
        if row_index == 1:
            first.font = TITLE_FONT
        elif row_index == 2:
            first.font = SUBTITLE_FONT
        elif label == "Review notes":
            first.font = Font(name="Arial", size=11, bold=True, color=NAVY)
            first.fill = SECTION_FILL
        worksheet.append([first, second])
    return len(rows)


def write_data_dictionary(
    worksheet,
    table_entries: list[dict[str, Any]],
    streaming: bool = False,
) -> tuple[int, int, str]:
    set_sheet_page_defaults(worksheet)
    headers = ["Workbook", "Sheet", "Field", "Inferred type", "Number format", "Definition"]
    widths = [20, 24, 48, 24, 20, 70]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    if streaming:
        header_cells = []
        for value in headers:
            cell = WriteOnlyCell(worksheet, value=value)
            style_header_cell(cell)
            header_cells.append(cell)
        worksheet.append(header_cells)
    else:
        worksheet.append(headers)
        for cell in worksheet[1]:
            style_header_cell(cell)
        worksheet.row_dimensions[1].height = 36

    row_count = 0
    for entry in table_entries:
        spec = entry["spec"]
        sheet_name = spec["sheet"]
        for header in entry["headers"]:
            values = entry.get("samples", {}).get(header, [])
            row_values = [
                entry["workbookLabel"],
                sheet_name,
                header,
                inferred_type(header, values),
                number_format_for_header(header),
                description_for_column(spec, header),
            ]
            if streaming:
                cells = []
                for column_index, value in enumerate(row_values, start=1):
                    cell = WriteOnlyCell(worksheet, value=value)
                    cell.font = BODY_FONT
                    cell.alignment = WRAP_ALIGNMENT if column_index == 6 else BODY_ALIGNMENT
                    cells.append(cell)
                worksheet.append(cells)
            else:
                worksheet.append(row_values)
                for column_index, value in enumerate(row_values, start=1):
                    apply_body_cell_style(worksheet.cell(worksheet.max_row, column_index), headers[column_index - 1], value)
            row_count += 1

    reference = add_excel_table(worksheet, "DataDictionaryTable", headers, row_count)
    return row_count, len(headers), reference


def read_csv_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        rows = [[parse_cell_value(value) for value in row] for row in reader]
    for row_index, row in enumerate(rows, start=2):
        if len(row) != len(headers):
            raise MissingInputError(
                f"CSV row {row_index} in {path} has {len(row)} values; expected {len(headers)}"
            )
    return headers, rows


def write_analysis_table(
    workbook: Workbook,
    entry: dict[str, Any],
) -> dict[str, Any]:
    spec = entry["spec"]
    headers, rows = read_csv_rows(entry["path"])
    worksheet = workbook.create_sheet(spec["sheet"])
    set_sheet_page_defaults(worksheet)
    worksheet.append(headers)
    worksheet.row_dimensions[1].height = 42
    for cell in worksheet[1]:
        style_header_cell(cell)

    sample_limit = min(100, len(rows))
    for column_index, header in enumerate(headers, start=1):
        samples = [rows[row_index][column_index - 1] for row_index in range(sample_limit)]
        worksheet.column_dimensions[get_column_letter(column_index)].width = sensible_width(header, samples)

    status_columns = {index + 1 for index, header in enumerate(headers) if "status" in header.lower()}
    p_columns = {
        index + 1
        for index, header in enumerate(headers)
        if "p_value" in header.lower() or "p_two_sided" in header.lower() or "permutation_p" in header.lower()
    }
    primary_column = next((index + 1 for index, header in enumerate(headers) if header == "is_primary"), None)

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column_index)
            safe_text_cell(cell, value)
            apply_body_cell_style(cell, headers[column_index - 1], value)
            if column_index in status_columns and isinstance(value, str):
                upper = value.upper()
                if upper == "PASS":
                    cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
                elif upper == "REVIEW":
                    cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
                elif upper in {"FAIL", "ERROR"}:
                    cell.fill = PatternFill("solid", fgColor=PALE_RED)
            if column_index in p_columns and isinstance(value, (int, float)) and not isinstance(value, bool):
                if value < 0.05:
                    cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
                elif value < 0.1:
                    cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        if primary_column and row[primary_column - 1] is True:
            for cell in worksheet[row_index]:
                cell.font = Font(name="Arial", size=10, bold=True, color="222222")
                if cell.fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)

    table_ref = add_excel_table(worksheet, spec["tableName"], headers, len(rows))
    return {
        "key": spec["key"],
        "sheet": spec["sheet"],
        "sourceCsv": str(entry["path"]),
        "sourceSha256": sha256_file(entry["path"]),
        "dataRows": len(rows),
        "dataColumns": len(headers),
        "headers": headers,
        "tableName": spec["tableName"],
        "tableRef": table_ref,
        "samples": {
            header: [rows[index][column_index] for index in range(sample_limit)]
            for column_index, header in enumerate(headers)
        },
        "rows": rows,
    }


def filtered_superplot_rows(
    headers: list[str], rows: list[list[Any]], chart_spec: dict[str, Any]
) -> dict[str, list[tuple[str, float]]]:
    index = {header: position for position, header in enumerate(headers)}
    required = [
        chart_spec["conditionColumn"],
        chart_spec["valueColumn"],
        chart_spec["labelColumn"],
        *chart_spec.get("filters", {}).keys(),
    ]
    missing = [column for column in required if column not in index]
    if missing:
        if chart_spec.get("required"):
            raise MissingInputError(f"Superplot {chart_spec['title']} is missing columns: {missing}")
        return {}
    grouped = {group["value"]: [] for group in chart_spec["groups"]}
    for row in rows:
        if any(row[index[column]] != expected for column, expected in chart_spec.get("filters", {}).items()):
            continue
        group_value = row[index[chart_spec["conditionColumn"]]]
        value = row[index[chart_spec["valueColumn"]]]
        label = row[index[chart_spec["labelColumn"]]]
        chart_value = None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            chart_value = float(value)
        elif isinstance(value, str) and NUMBER_RE.fullmatch(value):
            try:
                chart_value = float(Decimal(value))
            except (InvalidOperation, OverflowError, ValueError):
                chart_value = None
        if group_value in grouped and chart_value is not None and math.isfinite(chart_value):
            # Chart helper values are numeric display copies. The source worksheet
            # cell remains exact text when Excel cannot preserve all source digits.
            grouped[group_value].append((str(label), chart_value))
    if chart_spec.get("required") and any(not grouped[group["value"]] for group in chart_spec["groups"]):
        raise MissingInputError(f"Superplot {chart_spec['title']} has an empty required group")
    return grouped


def add_superplots(
    workbook: Workbook,
    schema: dict[str, Any],
    analysis_data: dict[str, dict[str, Any]],
) -> int:
    chart_specs = schema.get("superplots", [])
    if not chart_specs:
        return 0
    summary_sheet = workbook[table_map(schema)["summary"]["sheet"]]
    dashboard = workbook.create_sheet(
        "Summary Charts", index=workbook.index(summary_sheet) + 1
    )
    set_sheet_page_defaults(dashboard)
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    chart_data = workbook.create_sheet("_Chart Data")
    chart_data.sheet_state = "hidden"
    chart_data.sheet_view.showGridLines = False
    chart_count = 0
    row_cursor = 1
    anchors = ["A1", "N1", "A18", "N18", "A35", "N35"]

    for chart_index, chart_spec in enumerate(chart_specs):
        source = analysis_data.get(chart_spec["sourceTable"])
        if source is None:
            if chart_spec.get("required"):
                raise MissingInputError(f"Superplot source table missing: {chart_spec['sourceTable']}")
            continue
        grouped = filtered_superplot_rows(source["headers"], source["rows"], chart_spec)
        chart_data.cell(row=row_cursor, column=1, value=chart_spec["title"])
        chart_data.cell(row=row_cursor, column=1).font = Font(name="Arial", bold=True)
        start_row = row_cursor + 2
        chart = ScatterChart()
        chart.title = chart_spec["title"]
        chart.x_axis.title = chart_spec.get("xAxisTitle", "Condition")
        chart.y_axis.title = chart_spec.get("yAxisTitle", chart_spec["valueColumn"])
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"
        chart.x_axis.scaling.min = 0.5
        chart.x_axis.scaling.max = 2.5
        chart.x_axis.majorUnit = 1
        chart.height = 7.2
        chart.width = 13.0
        chart.legend.position = "b"
        max_group_rows = 0

        for group_index, group in enumerate(chart_spec["groups"]):
            values = grouped.get(group["value"], [])
            max_group_rows = max(max_group_rows, len(values))
            x_column = 1 + group_index * 3
            y_column = x_column + 1
            label_column = x_column + 2
            chart_data.cell(row=start_row - 1, column=x_column, value=f"{group['label']} x")
            chart_data.cell(row=start_row - 1, column=y_column, value=f"{group['label']} value")
            chart_data.cell(row=start_row - 1, column=label_column, value=f"{group['label']} set")
            offset_center = (len(values) - 1) / 2
            for value_index, (label, value) in enumerate(values):
                chart_data.cell(row=start_row + value_index, column=x_column, value=group["x"] + (value_index - offset_center) * 0.045)
                chart_data.cell(row=start_row + value_index, column=y_column, value=value)
                chart_data.cell(row=start_row + value_index, column=label_column, value=label)
            if values:
                x_values = Reference(chart_data, min_col=x_column, min_row=start_row, max_row=start_row + len(values) - 1)
                y_values = Reference(chart_data, min_col=y_column, min_row=start_row, max_row=start_row + len(values) - 1)
                series = Series(y_values, x_values, title=group["label"])
                series.marker.symbol = "circle"
                series.marker.size = 7
                series.graphicalProperties.line.noFill = True
                series.marker.graphicalProperties.solidFill = group["color"].lstrip("#")
                series.marker.graphicalProperties.line.solidFill = group["color"].lstrip("#")
                chart.series.append(series)

                mean_row = start_row + len(values) + 1
                mean_value = sum(value for _label, value in values) / len(values)
                chart_data.cell(row=mean_row, column=x_column, value=group["x"])
                chart_data.cell(row=mean_row, column=y_column, value=mean_value)
                mean_x = Reference(chart_data, min_col=x_column, min_row=mean_row, max_row=mean_row)
                mean_y = Reference(chart_data, min_col=y_column, min_row=mean_row, max_row=mean_row)
                mean_series = Series(mean_y, mean_x, title=f"{group['label']} mean")
                mean_series.marker.symbol = "diamond"
                mean_series.marker.size = 10
                mean_series.graphicalProperties.line.noFill = True
                mean_series.marker.graphicalProperties.solidFill = group["color"].lstrip("#")
                mean_series.marker.graphicalProperties.line.solidFill = "000000"
                chart.series.append(mean_series)

        anchor = anchors[chart_index] if chart_index < len(anchors) else f"Q{2 + chart_index * 16}"
        dashboard.add_chart(chart, anchor)
        row_cursor = start_row + max_group_rows + 5
        chart_count += 1
    if chart_count == 0:
        workbook.remove(dashboard)
        workbook.remove(chart_data)
    return chart_count


def build_analysis_workbook(
    output_path: Path,
    schema: dict[str, Any],
    entries: dict[str, dict[str, Any]],
    analysis_keys: list[str],
    generated_utc: str,
) -> dict[str, Any]:
    workbook = Workbook()
    configure_workbook_defaults(workbook)
    readme = workbook.active
    readme.title = "README"
    readme_rows = write_readme_sheet(
        readme,
        schema.get("workbook", {}),
        "Polished analysis, inference, representative-selection, control, and QC workbook.",
        generated_utc,
        [
            ("Raw spot companion", "MIAT_QKI_spots_exact_complete.xlsx"),
            ("Raw nucleus companion", "MIAT_QKI_nucleus_endpoints_complete.xlsx"),
            ("Formula policy", "No worksheet formulas; plotted helper values are literal and internal to this file."),
        ],
    )
    data_dictionary = workbook.create_sheet("Data Dictionary")
    analysis_data: dict[str, dict[str, Any]] = {}
    sheet_manifest: dict[str, Any] = {}
    dictionary_entries = []

    for key in analysis_keys:
        result = write_analysis_table(workbook, entries[key])
        analysis_data[key] = result
        sheet_manifest[key] = {k: v for k, v in result.items() if k not in {"rows", "samples"}}
        dictionary_entries.append(
            {
                "workbookLabel": "Analysis/QC",
                "spec": entries[key]["spec"],
                "headers": result["headers"],
                "samples": result["samples"],
            }
        )

    dictionary_rows, dictionary_columns, dictionary_ref = write_data_dictionary(
        data_dictionary, dictionary_entries, streaming=False
    )
    chart_count = add_superplots(workbook, schema, analysis_data)
    workbook.save(output_path)
    expected_sheets = ["README", "Data Dictionary"] + [
        entries[key]["spec"]["sheet"] for key in analysis_keys
    ]
    if chart_count:
        summary_index = expected_sheets.index(entries["summary"]["spec"]["sheet"])
        expected_sheets.insert(summary_index + 1, "Summary Charts")
        expected_sheets.append("_Chart Data")
    return {
        "path": str(output_path),
        "sha256": sha256_file(output_path),
        "bytes": output_path.stat().st_size,
        "role": "analysis",
        "expectedSheets": expected_sheets,
        "readmeRows": readme_rows,
        "dictionaryRows": dictionary_rows,
        "dictionaryColumns": dictionary_columns,
        "dictionaryTableRef": dictionary_ref,
        "chartCount": chart_count,
        "tables": sheet_manifest,
    }


def write_streaming_raw_workbook(
    output_path: Path,
    schema: dict[str, Any],
    entry: dict[str, Any],
    generated_utc: str,
    workbook_label: str,
) -> dict[str, Any]:
    spec = entry["spec"]
    headers = entry["headers"]
    workbook = Workbook(write_only=True)
    configure_workbook_defaults(workbook)

    readme = workbook.create_sheet("README")
    source_hash = sha256_file(entry["path"])
    readme_rows = append_write_only_readme(
        readme,
        schema.get("workbook", {}),
        workbook_label,
        generated_utc,
        [
            ("Data sheet", spec["sheet"]),
            ("Source staged CSV", str(entry["path"])),
            ("Source SHA256", source_hash),
            ("Value policy", "Safe short numerics and booleans are typed; leading-zero and >15-significant-digit values remain exact text."),
            ("Formula policy", "No worksheet formulas; formula-like source text is stored as text."),
        ],
    )

    dictionary = workbook.create_sheet("Data Dictionary")
    dictionary_rows, dictionary_columns, dictionary_ref = write_data_dictionary(
        dictionary,
        [
            {
                "workbookLabel": workbook_label,
                "spec": spec,
                "headers": headers,
                "samples": {},
            }
        ],
        streaming=True,
    )

    worksheet = workbook.create_sheet(spec["sheet"])
    set_sheet_page_defaults(worksheet)
    for column_index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = sensible_width(header)

    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=header)
        style_header_cell(cell)
        header_cells.append(cell)
    worksheet.append(header_cells)

    row_count = 0
    with entry["path"].open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        source_headers = next(reader)
        if source_headers != headers:
            raise MissingInputError(f"Header changed during build: {entry['path']}")
        for source_row_index, source_row in enumerate(reader, start=2):
            if len(source_row) != len(headers):
                raise MissingInputError(
                    f"CSV row {source_row_index} in {entry['path']} has {len(source_row)} values; expected {len(headers)}"
                )
            output_row = []
            for header, raw_value in zip(headers, source_row):
                value = parse_cell_value(raw_value)
                cell = WriteOnlyCell(worksheet, value=value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"
                apply_body_cell_style(cell, header, value)
                output_row.append(cell)
            worksheet.append(output_row)
            row_count += 1
            if row_count >= 1_048_576:
                raise MissingInputError(f"Excel row limit exceeded by {entry['path']}")

    table_ref = add_excel_table(worksheet, spec["tableName"], headers, row_count)
    workbook.save(output_path)
    return {
        "path": str(output_path),
        "sha256": sha256_file(output_path),
        "bytes": output_path.stat().st_size,
        "role": spec["key"],
        "expectedSheets": ["README", "Data Dictionary", spec["sheet"]],
        "readmeRows": readme_rows,
        "dictionaryRows": dictionary_rows,
        "dictionaryColumns": dictionary_columns,
        "dictionaryTableRef": dictionary_ref,
        "chartCount": 0,
        "tables": {
            spec["key"]: {
                "key": spec["key"],
                "sheet": spec["sheet"],
                "sourceCsv": str(entry["path"]),
                "sourceSha256": source_hash,
                "dataRows": row_count,
                "dataColumns": len(headers),
                "headers": headers,
                "tableName": spec["tableName"],
                "tableRef": table_ref,
            }
        },
    }


def build_companion_package(
    input_dir: Path | str,
    config_path: Path | str,
    output_dir: Path | str,
    companion_config_path: Path | str | None = None,
) -> dict[str, Any]:
    input_dir = Path(input_dir).resolve()
    config_path = Path(config_path).resolve()
    output_dir = Path(output_dir).resolve()
    companion_config_path = Path(companion_config_path).resolve() if companion_config_path else None
    if output_dir.exists():
        raise OutputExistsError(f"Output directory already exists; refusing overwrite: {output_dir}")
    if not input_dir.is_dir():
        raise MissingInputError(f"Input directory does not exist: {input_dir}")

    schema = merge_companion_config(load_json(config_path), companion_config_path)
    companion = schema.get("companionWorkbooks", {})
    analysis_keys = companion.get("analysisKeys", ANALYSIS_KEYS)
    spot_key = companion.get("spotKey", "spots")
    nucleus_key = companion.get("nucleusKey", "nucleus")
    required_keys = list(dict.fromkeys([*analysis_keys, spot_key, nucleus_key]))
    entries = validate_inputs(input_dir, schema, required_keys)

    output_dir.mkdir(parents=True, exist_ok=False)
    generated_utc = utc_now()
    try:
        analysis_path = output_dir / companion.get("analysisFilename", "MIAT_QKI_analysis_QC.xlsx")
        spots_path = output_dir / companion.get("spotsFilename", "MIAT_QKI_spots_exact_complete.xlsx")
        nucleus_path = output_dir / companion.get("nucleusFilename", "MIAT_QKI_nucleus_endpoints_complete.xlsx")

        analysis_manifest = build_analysis_workbook(
            analysis_path, schema, entries, analysis_keys, generated_utc
        )
        spots_manifest = write_streaming_raw_workbook(
            spots_path,
            schema,
            entries[spot_key],
            generated_utc,
            "Complete corrected per-spot exact-footprint measurements.",
        )
        nucleus_manifest = write_streaming_raw_workbook(
            nucleus_path,
            schema,
            entries[nucleus_key],
            generated_utc,
            "Complete per-nucleus endpoint measurements.",
        )

        adapter_manifest_path = input_dir / "adapter_manifest.json"
        adapter_manifest = load_json(adapter_manifest_path) if adapter_manifest_path.is_file() else {}
        manifest = {
            "status": "success",
            "architecture": "three_workbook_streaming_fallback",
            "generatedUtc": generated_utc,
            "inputDir": str(input_dir),
            "schemaPath": str(config_path),
            "schemaSha256": sha256_file(config_path),
            "companionConfigPath": str(companion_config_path) if companion_config_path else None,
            "companionConfigSha256": sha256_file(companion_config_path) if companion_config_path else None,
            "outputDir": str(output_dir),
            "v2SourcesOnly": adapter_manifest.get("v2SourcesOnly"),
            "valuesPolicy": "All CSV rows and columns are retained. Safe short numerics and booleans are typed; loss-prone values remain exact text.",
            "noCrossFileFormulas": True,
            "freezePanes": False,
            "omittedFromWorkbooks": adapter_manifest.get("omittedFromWorkbook"),
            "workbooks": {
                "analysis": analysis_manifest,
                "spots": spots_manifest,
                "nucleus": nucleus_manifest,
            },
        }
        manifest_path = output_dir / "companion_build_manifest.json"
        with manifest_path.open("x", encoding="utf-8", newline="\n") as handle:
            json.dump(manifest, handle, indent=2, ensure_ascii=False)
            handle.write("\n")
        manifest["manifestPath"] = str(manifest_path)
        return manifest
    except Exception:
        # Preserve partial outputs for diagnosis; never delete or overwrite them silently.
        raise


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--schema", required=True, type=Path)
    parser.add_argument("--companion-config", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        manifest = build_companion_package(
            args.input_dir, args.schema, args.output_dir, args.companion_config
        )
        print(json.dumps(manifest, ensure_ascii=False))
        return 0
    except OutputExistsError as exc:
        print(json.dumps({"status": "error", "code": "OUTPUT_EXISTS", "message": str(exc)}))
        return 2
    except MissingInputError as exc:
        print(json.dumps({"status": "error", "code": "MISSING_INPUT", "message": str(exc)}))
        return 3
    except Exception as exc:
        print(json.dumps({"status": "error", "code": "UNEXPECTED_ERROR", "message": str(exc)}))
        return 1


if __name__ == "__main__":
    sys.exit(main())
