import csv
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


TOOL_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_DIR))

from build_companion_workbooks import (  # noqa: E402
    OutputExistsError,
    build_companion_package,
    parse_cell_value,
)
from verify_companion_workbooks import verify_package  # noqa: E402
from render_raw_first_pages import render_raw_first_pages  # noqa: E402


def write_csv(path: Path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


class CompanionWorkbookTests(unittest.TestCase):
    def test_production_global_chart_uses_all_detected_endpoint(self):
        """Catches a floor-only field or stale label in the reader-facing global chart."""
        schema = json.loads((TOOL_DIR / "schema.production.json").read_text(encoding="utf-8"))
        global_chart = schema["superplots"][0]
        self.assertEqual(global_chart["title"], "Global MIAT spots per nucleus")
        self.assertEqual(global_chart["valueColumn"], "n_spots_all")
        self.assertEqual(global_chart["yAxisTitle"], "All-detected MIAT spots per nucleus")

    def test_loss_aware_cell_parsing(self):
        self.assertIsNone(parse_cell_value(""))
        self.assertIs(parse_cell_value("true"), True)
        self.assertEqual(parse_cell_value("42"), 42)
        self.assertEqual(parse_cell_value("0.125"), 0.125)
        self.assertEqual(parse_cell_value("00042"), "00042")
        self.assertEqual(parse_cell_value("1234567890123456"), "1234567890123456")
        self.assertEqual(parse_cell_value("0.12345678901234567"), "0.12345678901234567")

    def test_build_verify_and_refuse_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / "inputs"
            output = root / "COMPANION_WORKBOOKS_20990101-000000"
            config_path = root / "schema.json"

            tables = [
                ("summary", "summary.csv", "Summary", ["section", "result", "mean_nt", "mean_kd"]),
                ("set", "set.csv", "Set Wide", ["cohort", "arm", "biological_set", "metric"]),
                ("controls", "controls.csv", "Controls", ["image_key", "status", "notes"]),
                ("spots", "spots.csv", "Spots Exact", ["spot_uid", "value", "long_id"]),
                ("nucleus", "nucleus.csv", "Nucleus Endpoints", ["nucleus_uid", "value", "flag"]),
            ]
            for key, filename, _sheet, headers in tables:
                rows = {
                    "summary": [["Primary", "Metric", "1.0", "0.5"]],
                    "set": [
                        ["sampled_primary", "NT", "Slide_1_NT_1", "0.12345678901234567"],
                        ["sampled_primary", "KD", "Slide_1_KD_1", "0.06798619990978495"],
                    ],
                    "controls": [["KD_sec_only_42", "PASS", "Expected-negative structural zero"]],
                    "spots": [["S1", "2.5", "1234567890123456"], ["S2", "3.5", "00042"]],
                    "nucleus": [["N1", "7", "true"], ["N2", "8", "false"]],
                }[key]
                write_csv(inputs / filename, headers, rows)

            config = {
                "workbook": {
                    "title": "Synthetic MIAT×QKI workbook",
                    "subtitle": "Streaming companion test",
                    "purpose": "Exercise the companion architecture.",
                    "interpretationGuard": "Spatial association is not a direct-binding test.",
                    "analysisConvention": "Single reviewed z plane.",
                    "provenance": "Synthetic fixture values.",
                    "freezePanes": False,
                    "readmeNotes": ["Synthetic test only."],
                },
                "tables": [
                    {
                        "key": key,
                        "file": filename,
                        "sheet": sheet,
                        "tableName": f"{key.title()}Table".replace("_", ""),
                        "description": f"Synthetic {key} table.",
                        "required": True,
                        "requiredColumns": headers,
                    }
                    for key, filename, sheet, headers in tables
                ],
                "companionWorkbooks": {
                    "analysisKeys": ["summary", "set", "controls"],
                    "spotKey": "spots",
                    "nucleusKey": "nucleus",
                },
                "superplots": [
                    {
                        "title": "Synthetic metric",
                        "sourceTable": "set",
                        "conditionColumn": "arm",
                        "valueColumn": "metric",
                        "labelColumn": "biological_set",
                        "filters": {"cohort": "sampled_primary"},
                        "xAxisTitle": "Condition",
                        "yAxisTitle": "Metric",
                        "required": True,
                        "groups": [
                            {"value": "NT", "label": "NT sets", "x": 1, "color": "#0072B2"},
                            {"value": "KD", "label": "MIAT-KD sets", "x": 2, "color": "#D55E00"},
                        ],
                    }
                ],
            }
            config_path.write_text(json.dumps(config), encoding="utf-8")

            manifest = build_companion_package(inputs, config_path, output)
            self.assertEqual(manifest["status"], "success")
            self.assertEqual(len(manifest["workbooks"]), 3)

            verification = verify_package(output / "companion_build_manifest.json")
            self.assertEqual(verification["status"], "passed")
            self.assertEqual(verification["totals"]["formulaCells"], 0)
            self.assertEqual(verification["totals"]["errorCells"], 0)
            self.assertEqual(verification["workbooks"]["spots"]["dataRows"], 2)
            self.assertEqual(verification["workbooks"]["nucleus"]["dataColumns"], 3)
            self.assertEqual(verification["workbooks"]["analysis"]["chartCount"], 1)
            self.assertIn(
                "Summary Charts",
                verification["workbooks"]["analysis"]["sheetNames"],
            )

            analysis_path = Path(manifest["workbooks"]["analysis"]["path"])
            reloaded = load_workbook(analysis_path, read_only=True, data_only=False)
            try:
                self.assertEqual(
                    reloaded["Set Wide"].cell(row=2, column=4).value,
                    "0.12345678901234567",
                )
            finally:
                reloaded.close()

            chart_namespace = {
                "c": "http://schemas.openxmlformats.org/drawingml/2006/chart"
            }
            with zipfile.ZipFile(analysis_path) as archive:
                chart_root = ET.fromstring(archive.read("xl/charts/chart1.xml"))
            axes = {
                int(axis.find("c:axId", chart_namespace).attrib["val"]): axis
                for axis in chart_root.findall(".//c:valAx", chart_namespace)
            }
            self.assertEqual(
                axes[10].find("c:axPos", chart_namespace).attrib["val"],
                "b",
                "Condition axis must serialize at the bottom so outcome values are not clipped.",
            )
            self.assertEqual(
                axes[20].find("c:axPos", chart_namespace).attrib["val"],
                "l",
                "Outcome axis must serialize at the left.",
            )

            for workbook in manifest["workbooks"].values():
                with zipfile.ZipFile(workbook["path"]) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertNotIn(
                        b"_xlnm._FilterDatabase",
                        archive.read("xl/workbook.xml"),
                        "Table filtering must not be duplicated as a sheet-level filter database.",
                    )
                    for member in archive.namelist():
                        if member.startswith("xl/worksheets/sheet") and member.endswith(".xml"):
                            self.assertNotIn(b"<autoFilter", archive.read(member))

            raw_preview = render_raw_first_pages(
                manifest["workbooks"]["spots"]["path"],
                manifest["workbooks"]["nucleus"]["path"],
                root / "raw_previews",
            )
            self.assertEqual(raw_preview["status"], "success")
            self.assertEqual(raw_preview["previewCount"], 2)
            self.assertTrue(
                all(Path(item["png"]).stat().st_size > 0 for item in raw_preview["previews"])
            )

            with self.assertRaises(OutputExistsError):
                build_companion_package(inputs, config_path, output)


if __name__ == "__main__":
    unittest.main(verbosity=2)
