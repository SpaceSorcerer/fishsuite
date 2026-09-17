#!/usr/bin/env python3
"""Render bounded first-page previews from verified raw XLSX workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


NAVY = "17365D"
STRIPE = "EAF3F8"
GRID = "B4C6E7"


class RawPreviewError(RuntimeError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def font(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    candidate = Path("C:/Windows/Fonts") / name
    if candidate.is_file():
        return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return f"{value:.8g}"
    return str(value)


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, active_font, max_width: int, max_lines: int) -> list[str]:
    if not text:
        return [""]
    words = text.replace("_", "_ ").split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textlength(candidate, font=active_font) <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
        if len(lines) >= max_lines - 1:
            break
    if current and len(lines) < max_lines:
        lines.append(current)
    if not lines:
        lines = [text]
    if len(lines) == max_lines and " ".join(lines).replace("_ ", "_") != text:
        while lines[-1] and draw.textlength(lines[-1] + "…", font=active_font) > max_width:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "…"
    return lines


def normalized_rgb(color, fallback: str) -> str:
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and len(rgb) in {6, 8}:
        return rgb[-6:].upper()
    return fallback


def render_raw_preview(
    workbook_path: Path | str,
    sheet_name: str,
    output_path: Path | str,
    maximum_rows: int = 25,
    maximum_columns: int = 12,
) -> dict[str, Any]:
    workbook_path = Path(workbook_path).resolve()
    output_path = Path(output_path).resolve()
    if output_path.exists():
        raise RawPreviewError(f"Preview already exists; refusing overwrite: {output_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise RawPreviewError(f"Sheet {sheet_name!r} is missing from {workbook_path}")
        worksheet = workbook[sheet_name]
        rows = list(
            islice(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=maximum_rows,
                    min_col=1,
                    max_col=maximum_columns,
                ),
                maximum_rows,
            )
        )
        if not rows:
            raise RawPreviewError(f"Sheet {sheet_name!r} is empty in {workbook_path}")
        headers = [display_value(cell.value) for cell in rows[0]]
        if rows[0][0].font.name != "Arial" or not rows[0][0].font.bold:
            raise RawPreviewError(f"Expected Arial bold header style in {workbook_path} / {sheet_name}")
        header_fill = normalized_rgb(rows[0][0].fill.fgColor, NAVY)
        values = [[display_value(cell.value) for cell in row] for row in rows]
    finally:
        workbook.close()

    body_font = font(16)
    header_font = font(16, bold=True)
    measure_image = Image.new("RGB", (32, 32), "white")
    measure = ImageDraw.Draw(measure_image)
    column_widths = []
    for column in range(len(headers)):
        candidates = [row[column] for row in values[:25]]
        widest = max(
            [measure.textlength(text, font=body_font) for text in candidates]
            + [measure.textlength(headers[column], font=header_font)]
        )
        column_widths.append(int(max(120, min(250, widest + 28))))

    margin = 24
    header_height = 88
    body_height = 38
    canvas_width = margin * 2 + sum(column_widths)
    canvas_height = margin * 2 + header_height + body_height * (len(values) - 1)
    image = Image.new("RGB", (canvas_width, canvas_height), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    x = margin
    for column, width in enumerate(column_widths):
        draw.rectangle(
            (x, margin, x + width, margin + header_height),
            fill=f"#{header_fill}",
            outline=f"#{GRID}",
            width=1,
        )
        lines = wrap_lines(draw, headers[column], header_font, width - 14, 4)
        line_height = 19
        y = margin + max(6, (header_height - line_height * len(lines)) // 2)
        for line in lines:
            draw.text((x + 7, y), line, fill="#FFFFFF", font=header_font)
            y += line_height
        x += width

    for row_index, row in enumerate(values[1:], start=1):
        y0 = margin + header_height + (row_index - 1) * body_height
        fill = f"#{STRIPE}" if row_index % 2 == 0 else "#FFFFFF"
        x = margin
        for column, width in enumerate(column_widths):
            draw.rectangle(
                (x, y0, x + width, y0 + body_height),
                fill=fill,
                outline=f"#{GRID}",
                width=1,
            )
            text = row[column]
            if measure.textlength(text, font=body_font) > width - 14:
                while text and measure.textlength(text + "…", font=body_font) > width - 14:
                    text = text[:-1]
                text += "…"
            draw.text((x + 7, y0 + 9), text, fill="#222222", font=body_font)
            x += width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, format="PNG", optimize=True)
    return {
        "workbook": str(workbook_path),
        "workbookSha256": sha256_file(workbook_path),
        "sheet": sheet_name,
        "rowsRendered": len(values),
        "columnsRendered": len(headers),
        "headers": headers,
        "headerFont": "Arial bold",
        "headerFill": f"#{header_fill}",
        "png": str(output_path),
        "pngBytes": output_path.stat().st_size,
        "pngSha256": sha256_file(output_path),
    }


def render_raw_first_pages(
    spots_workbook: Path | str,
    nucleus_workbook: Path | str,
    output_directory: Path | str,
) -> dict[str, Any]:
    output_directory = Path(output_directory).resolve()
    if output_directory.exists():
        raise RawPreviewError(
            f"Raw preview output directory already exists; refusing overwrite: {output_directory}"
        )
    output_directory.mkdir(parents=True, exist_ok=False)
    previews = [
        render_raw_preview(
            spots_workbook,
            "Spots Exact",
            output_directory / "spots_exact_first_page.png",
        ),
        render_raw_preview(
            nucleus_workbook,
            "Nucleus Endpoints",
            output_directory / "nucleus_endpoints_first_page.png",
        ),
    ]
    report = {
        "status": "success",
        "renderer": "openpyxl read-only first rows + Pillow",
        "previewCount": len(previews),
        "previews": previews,
    }
    manifest_path = output_directory / "raw_preview_manifest.json"
    with manifest_path.open("x", encoding="utf-8", newline="\n") as handle:
        json.dump(report, handle, indent=2, ensure_ascii=False)
        handle.write("\n")
    report["manifestPath"] = str(manifest_path)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spots-workbook", required=True, type=Path)
    parser.add_argument("--nucleus-workbook", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        result = render_raw_first_pages(
            args.spots_workbook, args.nucleus_workbook, args.output_dir
        )
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except RawPreviewError as exc:
        print(json.dumps({"status": "error", "code": "RAW_PREVIEW_FAILED", "message": str(exc)}))
        return 5
    except Exception as exc:
        print(json.dumps({"status": "error", "code": "UNEXPECTED_ERROR", "message": str(exc)}))
        return 1


if __name__ == "__main__":
    sys.exit(main())
