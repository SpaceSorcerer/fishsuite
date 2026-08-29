#!/usr/bin/env python3
"""Reload and independently verify MIAT×QKI companion workbooks."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from build_companion_workbooks import parse_cell_value


class VerificationError(RuntimeError):
    """Raised when an independent workbook integrity check fails."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def values_equal(expected: Any, observed: Any) -> bool:
    if isinstance(expected, float) and isinstance(observed, (int, float)) and not isinstance(observed, bool):
        return math.isclose(expected, float(observed), rel_tol=1e-12, abs_tol=1e-12)
    return expected == observed


def source_sentinels(path: Path, row_count: int, column_count: int) -> dict[str, Any]:
    row_indices = sorted({1, max(1, (row_count + 1) // 2), row_count})
    column_indices = sorted({0, max(0, (column_count - 1) // 2), column_count - 1})
    selected: dict[int, list[Any]] = {}
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        if len(headers) != column_count:
            raise VerificationError(
                f"Source column count changed for {path}: {len(headers)} != {column_count}"
            )
        seen = 0
        for seen, row in enumerate(reader, start=1):
            if len(row) != column_count:
                raise VerificationError(
                    f"Source row {seen + 1} in {path} has {len(row)} columns; expected {column_count}"
                )
            if seen in row_indices:
                selected[seen] = [parse_cell_value(row[index]) for index in column_indices]
        if seen != row_count:
            raise VerificationError(
                f"Source row count changed for {path}: {seen} != {row_count}"
            )
    return {
        "rowIndices": row_indices,
        "columnIndices": column_indices,
        "values": selected,
    }


def stream_contains(archive: zipfile.ZipFile, member: str, needle: bytes) -> bool:
    overlap = max(64, len(needle) * 2)
    tail = b""
    with archive.open(member) as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                return needle in tail
            combined = tail + chunk
            if needle in combined:
                return True
            tail = combined[-overlap:]


def collect_table_parts(archive: zipfile.ZipFile) -> dict[str, str]:
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    result = {}
    for member in archive.namelist():
        if member.startswith("xl/tables/table") and member.endswith(".xml"):
            root = ET.fromstring(archive.read(member))
            name = root.attrib.get("displayName") or root.attrib.get("name")
            reference = root.attrib.get("ref")
            if not name or not reference:
                raise VerificationError(f"Malformed table part: {member}")
            auto_filter = root.find(f"{namespace}autoFilter")
            if auto_filter is None or auto_filter.attrib.get("ref") != reference:
                raise VerificationError(f"Table autofilter mismatch in {member}")
            result[name] = reference
    return result


def scan_workbook_cells(
    workbook_path: Path,
    workbook_manifest: dict[str, Any],
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=False)
    expected_sheets = workbook_manifest["expectedSheets"]
    if workbook.sheetnames != expected_sheets:
        raise VerificationError(
            f"Sheet order mismatch in {workbook_path}: {workbook.sheetnames} != {expected_sheets}"
        )

    formula_cells = 0
    error_cells = 0
    sheet_results: dict[str, Any] = {}
    table_by_sheet = {
        table["sheet"]: table for table in workbook_manifest.get("tables", {}).values()
    }
    for worksheet in workbook.worksheets:
        if getattr(worksheet, "freeze_panes", None) is not None:
            raise VerificationError(
                f"Frozen panes found in {workbook_path} / {worksheet.title}: {getattr(worksheet, 'freeze_panes', None)}"
            )
        expected_table = table_by_sheet.get(worksheet.title)
        row_count = 0
        header = None
        sentinel_rows: dict[int, list[Any]] = {}
        if expected_table:
            expected_data_rows = expected_table["dataRows"]
            expected_columns = expected_table["dataColumns"]
            sentinel_spec = source_sentinels(
                Path(expected_table["sourceCsv"]), expected_data_rows, expected_columns
            )
            target_excel_rows = {index + 1 for index in sentinel_spec["rowIndices"]}
        else:
            expected_data_rows = None
            expected_columns = None
            sentinel_spec = None
            target_excel_rows = set()

        for excel_row_index, row in enumerate(worksheet.iter_rows(), start=1):
            row_count += 1
            if excel_row_index == 1:
                header = [cell.value for cell in row]
            if excel_row_index in target_excel_rows:
                sentinel_rows[excel_row_index - 1] = [
                    row[index].value for index in sentinel_spec["columnIndices"]
                ]
            for cell in row:
                if cell.data_type == "f":
                    formula_cells += 1
                elif cell.data_type == "e":
                    error_cells += 1

        if expected_table:
            if header != expected_table["headers"]:
                raise VerificationError(f"Header mismatch in {workbook_path} / {worksheet.title}")
            if row_count - 1 != expected_data_rows:
                raise VerificationError(
                    f"Row count mismatch in {workbook_path} / {worksheet.title}: {row_count - 1} != {expected_data_rows}"
                )
            if len(header) != expected_columns:
                raise VerificationError(
                    f"Column count mismatch in {workbook_path} / {worksheet.title}: {len(header)} != {expected_columns}"
                )
            for row_index, expected_values in sentinel_spec["values"].items():
                observed_values = sentinel_rows.get(row_index)
                if observed_values is None:
                    raise VerificationError(
                        f"Missing sentinel row {row_index} in {workbook_path} / {worksheet.title}"
                    )
                for column_index, (expected, observed) in enumerate(
                    zip(expected_values, observed_values)
                ):
                    if not values_equal(expected, observed):
                        source_column = sentinel_spec["columnIndices"][column_index] + 1
                        raise VerificationError(
                            f"Sentinel mismatch in {worksheet.title} row {row_index} column {source_column}: {observed!r} != {expected!r}"
                        )
            first_cell = worksheet.cell(row=1, column=1)
            if first_cell.font.name != "Arial" or not first_cell.font.bold:
                raise VerificationError(
                    f"Header Arial/bold style missing in {workbook_path} / {worksheet.title}"
                )
            sheet_results[worksheet.title] = {
                "dataRows": expected_data_rows,
                "dataColumns": expected_columns,
                "headerMatched": True,
                "sentinelRows": sentinel_spec["rowIndices"],
                "sentinelColumns1Based": [index + 1 for index in sentinel_spec["columnIndices"]],
            }
        else:
            sheet_results[worksheet.title] = {
                "usedRows": row_count,
                "usedColumns": len(header or []),
            }
    workbook.close()
    return {
        "formulaCells": formula_cells,
        "errorCells": error_cells,
        "sheets": sheet_results,
    }


def verify_one_workbook(
    label: str, workbook_manifest: dict[str, Any]
) -> dict[str, Any]:
    path = Path(workbook_manifest["path"])
    if not path.is_file():
        raise VerificationError(f"Workbook is missing: {path}")
    observed_hash = sha256_file(path)
    if observed_hash != workbook_manifest["sha256"]:
        raise VerificationError(
            f"Workbook SHA256 mismatch for {path}: {observed_hash} != {workbook_manifest['sha256']}"
        )
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise VerificationError(f"ZIP integrity failure in {path}: {bad_member}")
        sheet_xml = [
            member
            for member in archive.namelist()
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", member)
        ]
        pane_members = [member for member in sheet_xml if stream_contains(archive, member, b"<pane")]
        if pane_members:
            raise VerificationError(f"Frozen pane XML found in {path}: {pane_members}")
        table_parts = collect_table_parts(archive)
        expected_tables = {
            table["tableName"]: table["tableRef"]
            for table in workbook_manifest.get("tables", {}).values()
        }
        expected_tables["DataDictionaryTable"] = workbook_manifest["dictionaryTableRef"]
        if table_parts != expected_tables:
            raise VerificationError(
                f"Excel table mismatch in {path}: {table_parts} != {expected_tables}"
            )
        width_sheet_count = sum(
            1 for member in sheet_xml if stream_contains(archive, member, b"<cols>")
        )
        if width_sheet_count < len(workbook_manifest.get("tables", {})) + 1:
            raise VerificationError(
                f"Expected column-width definitions are missing in {path}: {width_sheet_count}"
            )
        chart_count = len(
            [member for member in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", member)]
        )
        if chart_count != workbook_manifest.get("chartCount", 0):
            raise VerificationError(
                f"Chart count mismatch in {path}: {chart_count} != {workbook_manifest.get('chartCount', 0)}"
            )
        styles_xml = archive.read("xl/styles.xml")
        if b'val="Arial"' not in styles_xml and b'name="Arial"' not in styles_xml:
            raise VerificationError(f"Arial font is not declared in {path}")

    cell_scan = scan_workbook_cells(path, workbook_manifest)
    return {
        "path": str(path),
        "sha256": observed_hash,
        "bytes": path.stat().st_size,
        "zipIntegrity": "passed",
        "sheetNames": workbook_manifest["expectedSheets"],
        "tableCount": len(table_parts),
        "chartCount": chart_count,
        "columnWidthSheets": width_sheet_count,
        "noFrozenPanes": True,
        "formulaCells": cell_scan["formulaCells"],
        "errorCells": cell_scan["errorCells"],
        "sheets": cell_scan["sheets"],
        "dataRows": sum(table["dataRows"] for table in workbook_manifest.get("tables", {}).values()),
        "dataColumns": sum(table["dataColumns"] for table in workbook_manifest.get("tables", {}).values()),
    }


def verify_package(manifest_path: Path | str) -> dict[str, Any]:
    manifest_path = Path(manifest_path).resolve()
    with manifest_path.open("r", encoding="utf-8") as handle:
        manifest = json.load(handle)
    workbook_results = {
        label: verify_one_workbook(label, workbook_manifest)
        for label, workbook_manifest in manifest["workbooks"].items()
    }
    formula_cells = sum(result["formulaCells"] for result in workbook_results.values())
    error_cells = sum(result["errorCells"] for result in workbook_results.values())
    if formula_cells:
        raise VerificationError(f"Formula scan found {formula_cells} worksheet formula cells")
    if error_cells:
        raise VerificationError(f"Formula/error scan found {error_cells} worksheet error cells")
    return {
        "status": "passed",
        "manifestPath": str(manifest_path),
        "architecture": manifest.get("architecture"),
        "workbooks": workbook_results,
        "totals": {
            "formulaCells": formula_cells,
            "errorCells": error_cells,
            "workbookCount": len(workbook_results),
        },
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    try:
        result = verify_package(args.manifest)
        if args.output:
            output_path = args.output.resolve()
            if output_path.exists():
                raise VerificationError(f"Verification output already exists: {output_path}")
            with output_path.open("x", encoding="utf-8", newline="\n") as handle:
                json.dump(result, handle, indent=2, ensure_ascii=False)
                handle.write("\n")
            result["verificationPath"] = str(output_path)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except VerificationError as exc:
        print(json.dumps({"status": "error", "code": "VERIFICATION_FAILED", "message": str(exc)}))
        return 4
    except Exception as exc:
        print(json.dumps({"status": "error", "code": "UNEXPECTED_ERROR", "message": str(exc)}))
        return 1


if __name__ == "__main__":
    sys.exit(main())
