"""Workbook-sourced slide preparation and optional python-pptx export."""
from __future__ import annotations

from pathlib import Path
import json
import re
from functools import lru_cache

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, coordinate_to_tuple

from .aggregate import ReportInputError
from . import figures
from .provenance import guard_output, sha256


def _literal_number(text, identifiers=()):
    return _literal_number_cached(str(text), tuple(sorted(set(identifiers))))


@lru_cache(maxsize=2048)
def _literal_number_cached(text, identifiers):
    # Gene labels such as BIN1 and RNASEH2B are identities, not numeric claims.
    text = str(text)
    # Sam's question identifiers are ordering labels, not measured quantities.
    text = re.sub(r'^(?:1b|[1-4])\.\s+', '', text)
    for identifier in sorted(set(identifiers), key=len, reverse=True):
        text = re.sub(r'(?<!\w)' + re.escape(identifier) + r'(?!\w)', '', text)
    return re.search(r'(?<![\w])[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?(?:[a-zA-Zµμ%]+)?', text)


def validate_deck(workbook: Path, spec: dict) -> list:
    """Fail closed on untraced literal numbers, stale assets and absent cells."""
    workbook = Path(workbook)
    if not workbook.is_file():
        raise ReportInputError(f'missing workbook: {workbook}')
    if spec.get('workbook_sha256') != sha256(workbook):
        raise ReportInputError('stale workbook hash')
    book = load_workbook(workbook, data_only=True, read_only=True)
    cache={}
    def cell_value(sheet,cell):
        if sheet not in book:
            raise ReportInputError(f'missing workbook sheet: {sheet}')
        if sheet not in cache:
            cache[sheet]=list(book[sheet].values)
        row,col=coordinate_to_tuple(cell)
        values=cache[sheet]
        if row>len(values) or col>len(values[row-1]):
            raise ReportInputError(f'missing workbook cell: {sheet}!{cell}')
        return values[row-1][col-1]
    identifiers = set(book.sheetnames)
    for sheet in ('Per nucleus', 'Per well', 'Micrographs'):
        if sheet in book:
            rows = list(book[sheet].values)
            if len(rows) > 1:
                for index, column in enumerate(rows[1]):
                    if column in ('image', 'stem', 'group', 'condition', 'well_id'):
                        identifiers.update(str(row[index]) for row in rows[2:] if row[index] is not None)
    resolved = []
    try:
        if 'Slide values' in book and spec.get('schema') == 'resolved-deck-1':
            table=list(book['Slide values'].values)
            columns=list(table[1])
            for row in table[2:]:
                entry=dict(zip(columns,row))
                source_sheet,source_cell=entry['source_sheet'],entry['source_cell']
                if source_sheet not in book or not source_cell:
                    raise ReportInputError('missing Slide values source')
                source_value=cell_value(source_sheet,source_cell)
                expected=source_value
                if entry['label']=='MEASURED direction':
                    expected=('higher' if source_value>0 else 'lower' if source_value<0 else 'unchanged') if source_value is not None else 'missing'
                elif entry['label']=='display multiplier':
                    expected=figures.fraction_scale(source_value)
                if entry['value'] != expected:
                    raise ReportInputError(f'stale Slide values lineage: {source_sheet}!{source_cell}')
        for slide in spec.get('slides', []):
            for text in [slide.get('title',''), slide.get('body',''), slide.get('qualification','')]:
                if _literal_number(text, identifiers):
                    raise ReportInputError(f'literal numeric claim must use workbook cells: {text}')
            values = []
            for item in slide.get('values', []):
                if _literal_number(item.get('label', ''), identifiers):
                    raise ReportInputError('literal numeric claim must use workbook cells: ' + item['label'])
                sheet, cell = item['sheet'], item['cell']
                if sheet not in book.sheetnames or not re.fullmatch(r'[A-Z]+[1-9][0-9]*',cell):
                    raise ReportInputError(f'missing workbook cell: {sheet}!{cell}')
                value = cell_value(sheet,cell)
                if value is None and not item.get('allow_missing',False):
                    raise ReportInputError(f'missing workbook value: {sheet}!{cell}')
                if isinstance(value, str) and _literal_number(value, identifiers) and item.get('display', True):
                    raise ReportInputError(f'literal numeric claim must use typed workbook cells: {sheet}!{cell}')
                values.append(dict(item, value=value))
            if slide.get('title_cell'):
                title_ref=slide['title_cell']
                if cell_value(title_ref['sheet'],title_ref['cell']) != slide['title']:
                    raise ReportInputError('stale title claim')
            resolved_assets = []
            for asset in slide.get('figures', []):
                path = Path(asset['path'])
                if not path.is_file():
                    raise ReportInputError(f'missing figure: {path}')
                if sha256(path) != asset.get('sha256'):
                    raise ReportInputError(f'stale figure: {path}')
                if asset.get('cohort') != spec.get('cohort'):
                    raise ReportInputError(f'figure cohort mismatch: {path}')
                caption = ''
                if asset.get('caption_ref'):
                    ref = asset['caption_ref']
                    caption = cell_value(ref['sheet'], ref['cell'])
                    if not isinstance(caption, str) or _literal_number(caption, identifiers):
                        raise ReportInputError('literal numeric or missing micrograph caption; use typed workbook cells')
                resolved_assets.append(dict(asset, caption_value=caption))
            resolved.append(dict(slide, values=values, figures=resolved_assets))
    finally:
        book.close()
    if not resolved:
        raise ReportInputError('missing slide definitions')
    return resolved


def speaker_notes(workbook, definition):
    """Short spoken summary from already validated workbook cells, never cell refs."""
    if definition.get('identity') in ('simple_coloc','cytofluorogram'):
        groups={}
        for item in definition.get('values',[]):
            if item['sheet']=='Simple coloc metrics':
                row,_=coordinate_to_tuple(item['cell'])
                groups.setdefault(row,{})[item.get('label','')]=item['value']
        lines=[f'Workbook: {Path(workbook).resolve()}', 'Levels of comparison: per nucleus → per FOV → per well.']
        for row in groups.values():
            lines.append(f"{row.get('endpoint')}: {row.get('reference_group')} well mean {row.get('mean_ref')}; {row.get('test_group')} well mean {row.get('mean_test')}; mixed p {row.get('p_mixed')}; Welch (wells) p {row.get('p_welch')}. {row.get('sensitivity_mixed_status')}.")
            for arm in ('reference','test'):
                lines.append(f"{arm}: {row.get('sensitivity_n_nuclei_'+arm)} defined nuclei, {row.get('sensitivity_n_fovs_'+arm)} FOVs, {row.get('sensitivity_n_wells_'+arm)} wells; nucleus median {row.get('sensitivity_nucleus_median_'+arm)}.")
        lines += ['Nucleus-level model: fixed arm, random well and FOV within well. Two-sided asymptotic Wald p; exploratory, unadjusted.',
                  'Manders uses Costes-converged nuclei only; failed thresholds remain missing. Pearson and ICQ use all defined biological nuclei.']
        if definition['identity']=='cytofluorogram':
            lines.append('One nucleus per arm, nearest its median Pearson r; descriptive pixel density on identical raw-intensity axes. No pixel-level test.')
            lines.extend(f"{v.get('label')}: {v['value']}" for v in definition.get('values',[]) if v['sheet']=='Simple coloc representatives')
        return '\n'.join(lines)
    endpoints = {}
    for item in definition.get('values', []):
        label = item.get('label', '')
        if item.get('sheet') == 'Contrasts' and ': ' in label:
            endpoint, column = label.split(': ', 1)
            endpoints.setdefault(endpoint, {})[column] = item['value']
    def number(value):
        value = pd.to_numeric(value, errors='coerce')
        return f'{value:.4g}' if np.isfinite(value) else 'missing'
    notes = [f'Workbook: {Path(workbook).resolve()}']
    if definition.get('question'):
        notes.append(definition['question'])
    levels, means, tests, adjusted = [], [], [], []
    for endpoint, row in endpoints.items():
        name = row.get('endpoint_plain', endpoint)
        test, ref = row.get('test_group', 'test'), row.get('reference_group', 'reference')
        n = lambda key: number(row.get(key))
        means.append(f"{name}: {test} well mean {n('mean_test')}, {ref} well mean {n('mean_ref')}, difference {n('diff')} ({test} minus {ref})")
        tests.append(f"{name}: mixed model p {n('p_mixed')}; Welch (wells) p {n('p_welch')}, Hedges g {n('hedges_g')}")
        adjusted.append(f"{name}: Holm p {n('p_headline_holm')}, MDE g {n('mde_hedges_g_at_family_alpha')} at family alpha and 80% power")
        levels.extend([str(name),
            f"Per-nucleus median: {test} {n('sensitivity_nucleus_median_test')} (n={n('sensitivity_n_nuclei_test')} nuclei), {ref} {n('sensitivity_nucleus_median_reference')} (n={n('sensitivity_n_nuclei_reference')} nuclei); descriptive median, no pooled nucleus t test.",
            f"Per-FOV mean: {test} n={n('sensitivity_n_fovs_test')} FOVs, {ref} n={n('sensitivity_n_fovs_reference')} FOVs; technical-level Welch p {n('sensitivity_welch_fov_p')}.",
            f"Per-well mean: {test} n={n('n_wells_test')} wells, {ref} n={n('n_wells_reference')} wells; two-sided Welch comparison; pooled-variance Student sensitivity p {n('sensitivity_student_well_p')}.",
            f"Nucleus-level mixed model: fixed arm, random well and FOV within well; Wald p {n('sensitivity_mixed_p')}; {row.get('sensitivity_mixed_status', 'missing')}."])
    if endpoints:
        notes.extend(['; '.join(means)+'.', '; '.join(tests)+'.', '; '.join(adjusted)+'. MDE is not an exclusion bound.'])
    else:
        readouts = [str(v['value']) for v in definition.get('values',[]) if v.get('label') == 'readout']
        notes.extend(readouts or ['This slide describes the recorded measurements and study design.'])
        notes.append('No endpoint contrast is displayed on this slide.')
        levels.append('Nuclei are measurement units, FOVs are technical replicates, and well means are biological replicates; the headline is the nucleus-level mixed model.')
    return '\n'.join(notes + ['', 'Levels of comparison'] + levels + ['Mixed-model headline decision 2026-09-07 after data inspection; Welch on well means remains reported. Student and FOV Welch are sensitivities.'])


def build_deck(workbook: Path, spec: dict, destination: Path) -> Path:
    destination = guard_output(destination)
    sources_path = guard_output(destination.with_name('slide_sources.csv'))
    slides = validate_deck(workbook, spec)
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ReportInputError('python-pptx missing; installation prohibited by dispatch') from exc
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(13.333333), Inches(7.5)
    sources = []
    workbook_path = str(Path(workbook).resolve())
    for number, definition in enumerate(slides,1):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.4), Inches(.15), Inches(12.5), Inches(.95))
        box.text_frame.text = definition['title']
        box.text_frame.word_wrap = True
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.name, paragraph.font.size = 'Arial', Pt(23)
        notes = speaker_notes(workbook, definition)
        for item in definition['values']:
            sources.append(dict(slide=number, workbook=workbook_path,
                                sheet=item['sheet'], cell=item['cell'], value=item['value']))
        assets = definition.get('figures', [])
        if definition.get('layout')=='micrographs':
            box.top,box.height=Inches(.03),Inches(.45)
            for p in box.text_frame.paragraphs:p.font.size=Pt(18)
            for i,label in enumerate(definition.get('row_labels',[])):
                rowbox=slide.shapes.add_textbox(Inches(.25),Inches(.48+3.45*i),Inches(12.8),Inches(.2))
                rowbox.text_frame.margin_top=rowbox.text_frame.margin_bottom=0
                rowbox.text_frame.text=label
                for p in rowbox.text_frame.paragraphs:p.font.size=Pt(10)
        for asset_no, asset in enumerate(assets):
            from PIL import Image
            with Image.open(asset['path']) as img:
                ratio = img.width / img.height
            micrograph = definition.get('layout') == 'micrographs'
            columns = 4 if micrograph else 1 if len(assets)==1 else 2 if len(assets) in (2,4) else 3
            rows = (len(assets)+columns-1)//columns
            cell_width = 12.73/columns
            cell_height = 3.45 if micrograph else 4.95/rows
            max_width, max_height = cell_width-.04, (3.12 if micrograph else cell_height-.27)
            width = min(max_width, max_height*ratio)
            height = width/ratio
            left = .30+(asset_no%columns)*cell_width+(cell_width-width)/2
            top = (.83 if micrograph else 1.32)+(asset_no//columns)*cell_height+(max_height-height)/2
            slide.shapes.add_picture(asset['path'], Inches(left), Inches(top), Inches(width), Inches(height))
            caption_text = asset.get('caption_value') or asset.get('caption','')
            if caption_text:
                caption = slide.shapes.add_textbox(Inches(.30+(asset_no%columns)*cell_width+.05), Inches(top-.24), Inches(cell_width-.10), Inches(.23))
                caption.text_frame.margin_top = caption.text_frame.margin_bottom = 0
                caption.text_frame.text = caption_text
                for paragraph in caption.text_frame.paragraphs:
                    paragraph.font.name, paragraph.font.size = 'Arial', Pt(9)
                    from pptx.enum.text import PP_ALIGN
                    paragraph.alignment=PP_ALIGN.CENTER
            sources.append(dict(slide=number, figure=asset['path'], sha256=asset['sha256'],
                                full_figure=asset.get('full_path', ''), full_sha256=asset.get('full_sha256', '')))
        if not definition.get('figures'):
            content = slide.shapes.add_textbox(Inches(.6), Inches(1.2), Inches(12), Inches(5.5))
            content.text_frame.text = definition.get('body','') + '\n' + '\n'.join(
                f'{v.get("label", "value")}: {v["value"]}' for v in definition['values'] if v.get('display',True))
            content.text_frame.word_wrap = True
            for p in content.text_frame.paragraphs:
                p.font.name, p.font.size = 'Arial', Pt(18)
        readouts = [v['value'] for v in definition['values'] if v.get('label') == 'readout']
        if definition.get('readout_template'):
            readouts = [definition['readout_template'].format(**{v['label']:v['value'] for v in definition['values']})]
        if assets and definition.get('body') and not readouts:
            readouts=[definition['body']]
        if readouts and assets and definition.get('layout')!='micrographs':
            box = slide.shapes.add_textbox(Inches(.5), Inches(6.25), Inches(12.3), Inches(1.0))
            box.text_frame.word_wrap = True
            box.text_frame.text = str(readouts[0])
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.name, paragraph.font.size = 'Arial', Pt(18)
        slide.notes_slide.notes_text_frame.text = notes
    destination.parent.mkdir(parents=True, exist_ok=True)
    ppt.save(destination)
    pd.DataFrame(sources).to_csv(sources_path, index=False)
    return destination


def simple_coloc_slides(sheets, assets, rna, partner, cohort):
    """The same two workbook-traced identities for either channel pair."""
    refs=[]
    for sheet in ('Simple coloc metrics','Simple coloc representatives'):
        if sheet not in sheets: continue
        table=sheets[sheet]
        for i,row in table.iterrows():
            for j,column in enumerate(table.columns):
                refs.append(dict(sheet=sheet,cell=f'{get_column_letter(j+1)}{i+3}',
                                 label=column,allow_missing=True,display=False))
    sentence='Pearson and ICQ describe intensity agreement; Manders describes directional signal overlap at converged Costes thresholds.'
    sheets['Simple coloc readout']=pd.DataFrame({'sentence':[sentence,'Each panel shows the nucleus closest to its arm’s median Pearson correlation.']})
    result=[]
    for i,(identity,title,path) in enumerate(zip(
            ['simple_coloc','cytofluorogram'],
            [f'Pixel colocalization, {rna} × {partner}',f'Cytofluorogram, {rna} × {partner}'],assets)):
        path=Path(path)
        candidates=sorted(path.parent.glob(('FIG_SIMPLE_COLOC_' if identity=='simple_coloc' else 'FIG_CYTOFLUOROGRAM_')+'*_focus.png'))
        candidates=[p for p in candidates if p.name!='FIG_SIMPLE_COLOC_focus.png']
        if candidates: path=candidates[0]
        asset=dict(path=str(path.resolve()),sha256=sha256(path),cohort=cohort)
        full=path.with_name(path.name.replace('_focus','_full'))
        if full!=path and full.is_file(): asset.update(full_path=str(full.resolve()),full_sha256=sha256(full))
        single_assets=[]
        for candidate in candidates:
            twin=candidate.with_name(candidate.name.replace('_focus','_full'))
            single_assets.append(dict(path=str(candidate.resolve()),sha256=sha256(candidate),cohort=cohort,
                full_path=str(twin.resolve()),full_sha256=sha256(twin),single_plot=True))
        result.append(dict(identity=identity,title=title,figures=single_assets or [asset],values=refs+[
            dict(sheet='Simple coloc readout',cell=f'A{i+3}',label='readout',display=True)]))
    return result


def append_per_well_micrographs(spec, sheets, micro):
    """Replace the overview micrograph with separately movable native panels."""
    sheets['Per-well micrograph selection'] = pd.DataFrame(micro['selection'])
    panel_rows = []
    replacements = []
    for index, item in enumerate(micro['slides']):
        assets = []
        for row in item['rows']:
            for panel in row['panels']:
                path = Path(panel['path'])
                panel_rows.append(dict(group=row['group'], well_id=row['well_id'],
                                       image=row['image'], **panel))
                assets.append(dict(path=str(path.resolve()), sha256=sha256(path),
                    cohort=spec['cohort'], caption=panel['panel']))
        replacements.append(dict(identity='micrographs_well_'+str(index+1),
            title='Per-well micrographs: '+item['title'], values=[], figures=assets, layout='micrographs',
            row_labels=[r['group']+' · '+r['well_id']+' · FOV '+Path(r['image']).stem.rsplit('_',1)[-1] for r in item['rows']]))
    sheets['Per-well micrograph panels'] = pd.DataFrame(panel_rows)
    if replacements:
        positions = [i for i,s in enumerate(spec['slides']) if s.get('identity')=='micrographs']
        position = positions[0] if positions else len(spec['slides'])
        spec['slides'] = [s for s in spec['slides'] if s.get('identity')!='micrographs']
        spec['slides'][position:position] = replacements
    return spec


def draw_micrograph_panel(ax, row, color="black"):
    """Display crop with a padded annotation band; source raster is unchanged.

    Omit the bottom six percent containing the verified native annotation, then
    redraw the calibrated bar in padding. Coordinates retain source-pixel scale.
    At the full-width deck placement, all annotation text is at least 18 pt.
    """
    pixels = figures.plt.imread(row['path'])
    height, width = pixels.shape[:2]
    stop = int(height * .94)
    ax.imshow(pixels[:stop], extent=(-.5, width-.5, stop-.5, -.5))
    ax.set_facecolor('black')
    ax.set_xlim(-.5, width-.5)
    ax.set_ylim(height * 1.18, -.5)
    from matplotlib.patches import Rectangle
    ax.add_patch(Rectangle((-.5, stop-.5), width, height * .25,
                           color='black', zorder=2))
    bar_px = float(row['bar_um']) / float(row['voxel_xy_um'])
    right = width * .94
    if bar_px > width * .88:
        raise ReportInputError('calibrated scale bar does not fit micrograph panel')
    ax.plot([right-bar_px, right], [height*1.10]*2, color='white', lw=3,
            solid_capstyle='butt', zorder=3)
    ax.text(right-bar_px/2, height*.99, f"{row['bar_um']:g} µm",
            ha='center', va='center', fontsize=18, color='white', zorder=3)
    ax.set_title(row['group'], fontsize=20, color=color)
    ax.axis('off')


def prepare_deck(template: dict, sheets: dict, out_dir: Path, data: dict,
                 panel, ctx, well: pd.DataFrame, field: pd.DataFrame,
                 contrasts: pd.DataFrame, endpoints: list) -> dict:
    """Resolve semantic endpoint requests only after report rows are available.

    Values copied to Slide values retain their exact origin cell. Speaker notes
    summarize contrasts; slide_sources.csv retains all cell-level provenance.
    """
    from . import figures as fig
    out_dir = guard_output(out_dir)
    for item in template['slides']:
        for key in ('title', 'body', 'qualification'):
            if _literal_number(item.get(key, '')):
                raise ReportInputError('literal numeric claim must use workbook cells: ' + item[key])
    selection = Path(template['nucleus_selection'])
    if not selection.is_file() or sha256(selection) != template['nucleus_selection_sha256']:
        raise ReportInputError('missing or stale nucleus selection source')
    micro_root = Path(template['micrographs'])
    if not micro_root.is_dir():
        raise ReportInputError('missing calibrated micrograph directory')
    selection_text = selection.read_text(encoding='utf-8')
    sheets['Deck specification']=pd.DataFrame([
        dict(title=item['title'],body=item.get('body',''),qualification=item.get('qualification',''))
        for item in template['slides']])
    counts = data['nuclei'].groupby('group',dropna=False).size()
    sel = [dict(stage='retained labels', value=len(data['nuclei']), source='Per nucleus; A0 retained roster'),
           dict(stage='pre-area candidates', value='missing', source=str(selection)),
           dict(stage='pre-border exclusions', value='missing', source=str(selection))]
    sel += [dict(stage=str(group), value=int(count), source='Per nucleus / group') for group,count in counts.items()]
    sheets['Nucleus selection'] = pd.DataFrame(sel)
    sheets['Selection evidence'] = pd.DataFrame([dict(file=str(selection),sha256=sha256(selection),text=selection_text)])
    micro = []
    overlays = panel.tables['overlays_index']
    for group in ctx.group_order:
        choices = overlays.loc[overlays.group==group].sort_values('well_id')
        if choices.empty:
            raise ReportInputError(f'missing micrograph for {group}')
        row = choices.iloc[0]
        path = micro_root / (row['stem']+'__merge_all.png')
        if not path.is_file():
            raise ReportInputError(f'missing micrograph: {path}')
        if sha256(path) != template.get('micrograph_hashes',{}).get(path.name):
            raise ReportInputError(f'missing or stale pinned micrograph: {path}')
        nuc = data['nuclei'].loc[data['nuclei'].image==row['image']]
        vox = pd.to_numeric(nuc.voxel_xy_um).dropna().unique()
        if len(vox)!=1 or not np.isfinite(vox[0]) or vox[0]<=0:
            raise ReportInputError('missing or ambiguous micrograph calibration')
        planes=panel.tables['per_nucleus'].loc[lambda x:x.image==row['image'],'z_plane'].dropna().unique()
        if len(planes)!=1 or planes[0]!=row['z_plane']:
            raise ReportInputError('micrograph plane differs from persisted nuclear plane')
        from PIL import Image
        with Image.open(path) as image:
            width,height = image.size
            pixels=np.asarray(image.convert('RGB'))
        bar_um=template.get('micrograph_bar_um',{}).get(path.name)
        if bar_um is None or not np.isfinite(bar_um) or bar_um<=0:
            raise ReportInputError('missing native micrograph scale-bar annotation')
        # The named native PNGs already carry a lower-right white scale bar.
        # Verify its raster length against the source voxel calibration, and
        # retain it instead of drawing a second, differently sized bar.
        bright=(pixels[int(height*.9):,int(width*.7):,:]>240).all(axis=2)
        runs=[]
        for scan in bright:
            transitions=np.diff(np.r_[False,scan,False].astype(int))
            runs.extend(np.flatnonzero(transitions==-1)-np.flatnonzero(transitions==1))
        native_bar_px=int(max(runs,default=0))
        if abs(native_bar_px-bar_um/float(vox[0]))>1:
            raise ReportInputError('native micrograph scale-bar length disagrees with source calibration')
        # Native image dimensions must equal the persisted mask geometry; a scaled
        # preview cannot inherit source micrometres-per-pixel without a transform.
        masks = [m for m in panel.manifest.get('persisted_masks',[]) if m['image']==row['image']]
        if len(masks)!=1 or masks[0]['shape'] != [height,width]:
            raise ReportInputError('micrograph calibration mismatch: image dimensions differ from retained mask')
        if sha256(Path(masks[0]['path'])) != masks[0]['sha256']:
            raise ReportInputError('stale micrograph calibration mask')
        micro.append(dict(group=group,image=row['image'],path=str(path.resolve()),sha256=sha256(path),
                          voxel_xy_um=float(vox[0]),bar_um=float(bar_um),native_bar_px=native_bar_px,z_plane=int(row['z_plane']),
                          width_px=width,height_px=height,calibration_source='nuclei_metrics.csv:voxel_xy_um',
                          plane_source='coloc_standard_panel.xlsx:overlays_index/z_plane',
                          selection='first persisted representative well in group; presentation only'))
    # Exact native publication PNGs use the run's recorded manual windows/LUTs.
    panel_rows = []
    for row in micro:
        native = fig.publication_panel_paths(ctx.run_dir/'publication_images',
                    Path(row['path']).name.removesuffix('__merge_all.png'),
                    data['cfg']['config_resolved'])
        row['panels'] = [dict(row, **entry) for entry in native]
        panel_rows.extend(dict(group=row['group'],image=row['image'],**entry) for entry in native)
    sheets['Micrograph panels'] = pd.DataFrame(panel_rows)
    sheets['Micrographs'] = pd.DataFrame([{k:v for k,v in row.items() if k != 'panels'} for row in micro])
    from .aggregate import reconcile_localization
    spots = pd.read_csv(ctx.run_dir/'spot_metrics.csv')
    audit = reconcile_localization(data['nuclei'], spots)
    if audit['status'] != 'ok':
        raise ReportInputError('Localization reconciliation blocked: ' + '; '.join(audit['errors']))
    for key, name in [('per_nucleus','Localization counts'), ('unassigned','Localization unassigned'),
                      ('checks','Localization checks'), ('territory','Localization territory')]:
        sheets[name] = audit[key]
    if 'dapi' in data:
        from .dapi_mask import render_dapi
        localization = render_dapi(ctx, data['dapi'], data['nuclei'], well, field, contrasts, out_dir/'localization')
    else:
        localization = fig.render_localization(ctx, well, field, data['nuclei'], spots, contrasts,
                                                out_dir/'localization', pub_dir=micro_root)
    sheets['Localization crops'] = pd.DataFrame(localization['crops'])
    sheets['Endpoint coverage'] = (field.groupby(['endpoint','group'],as_index=False)
        .agg(defined_nuclei=('n_nuclei_nonmissing','sum'),
             eligible_nuclei=('n_nuclei_total','sum'),fovs=('image','nunique')))
    by_endpoint = {e.name:e for e in endpoints}
    values, figure_rows, resolved = [], [], []
    figure_dir = guard_output(out_dir/'figures')
    figure_dir.mkdir(parents=True,exist_ok=True)
    fig.set_style()
    for slide_no, item in enumerate(template['slides'],1):
        refs = []
        values.append(dict(slide=slide_no,endpoint='',label='title',value=item['title'],
                           source_sheet='Deck specification',source_cell=f'A{slide_no+2}'))
        title_ref=dict(sheet='Slide values',cell=f'D{len(values)+2}')
        refs.append(dict(**title_ref,label='title',display=False))
        refs.append(dict(sheet='Deck specification',cell=f'A{slide_no+2}',label='title source',display=False))
        names = item.get('endpoints',[])
        for name in names:
            match = contrasts.index[contrasts.endpoint==name].tolist()
            if len(match)!=1:
                raise ReportInputError(f'missing or ambiguous slide endpoint: {name}')
            i = match[0]
            for col_index, col in enumerate(contrasts.columns,1):
                value = contrasts.loc[i,col]
                values.append(dict(slide=slide_no,endpoint=name,label=col,value=value,
                                   source_sheet='Contrasts',source_cell=f'{get_column_letter(col_index)}{i+3}'))
                refs.append(dict(sheet='Slide values',cell=f'D{len(values)+2}',label=f'{name}: {col}',allow_missing=True,display=False))
                refs.append(dict(sheet='Contrasts',cell=f'{get_column_letter(col_index)}{i+3}',label=f'{name}: {col}',allow_missing=True,display=False))
            direction = 'higher' if contrasts.loc[i,'diff']>0 else 'lower' if contrasts.loc[i,'diff']<0 else 'unchanged' if contrasts.loc[i,'diff']==0 else 'missing'
            values.append(dict(slide=slide_no,endpoint=name,label='MEASURED direction',value=direction,
                               source_sheet='Contrasts',source_cell=f'{get_column_letter(contrasts.columns.get_loc("diff")+1)}{i+3}'))
            refs.append(dict(sheet='Slide values',cell=f'D{len(values)+2}',label=f'{name}: MEASURED direction',display=True))
            for wi,row in well.loc[well.endpoint==name].iterrows():
                # iterrows retains the aggregation index, while the workbook uses
                # positional row order. Resolve position explicitly.
                position=well.index.get_loc(wi)
                for column in ('group','well_id','well_mean_of_field_values'):
                    cell=f'{get_column_letter(well.columns.get_loc(column)+1)}{position+3}'
                    refs.append(dict(sheet='Per well',cell=cell,label=f'{name}: {column}',allow_missing=True,display=False))
            values.append(dict(slide=slide_no,endpoint=name,label='display multiplier',
                               value=fig.fraction_scale(name),
                               source_sheet='Endpoint definitions',source_cell=f'A{list(by_endpoint).index(name)+3}'))
            refs.append(dict(sheet='Slide values',cell=f'D{len(values)+2}',label=f'{name}: display multiplier',display=False))
        extra_sheets = list(item.get('sheets',[]))
        if item.get('micrographs') or item.get('kind') == 'micrographs':
            extra_sheets.append('Micrograph panels')
        if names:
            coverage = sheets['Endpoint coverage']
            for ci,row in coverage.loc[coverage.endpoint.isin(names)].iterrows():
                for column_index,column in enumerate(coverage.columns,1):
                    refs.append(dict(sheet='Endpoint coverage',cell=f'{get_column_letter(column_index)}{ci+3}',
                                     label=f'{row.endpoint}: {row.group}: {column}',display=False))
        for sheet_name in extra_sheets:
            if sheet_name not in sheets:
                raise ReportInputError(f'missing requested slide sheet: {sheet_name}')
            for row_i, row in enumerate(sheets[sheet_name].itertuples(index=False,name=None),3):
                for col_i,value in enumerate(row,1):
                    refs.append(dict(sheet=sheet_name,cell=f'{get_column_letter(col_i)}{row_i}',
                                     label=f'{sheet_name}: {sheets[sheet_name].columns[col_i-1]}',allow_missing=True,display=False))
        assets = []
        stem = item['asset']
        if item.get('kind') == 'localization':
            for record in localization['figures']:
                if record.get('is_composite'): continue
                path = out_dir/'localization'/record['png']
                full = out_dir/'localization'/record.get('full_png', record['png'])
                assets.append(dict(path=str(path.resolve()), sha256=sha256(path),
                    full_path=str(full.resolve()), full_sha256=sha256(full),
                    cohort=template['cohort'], caption=record.get('description',''), single_plot=True))
        elif names:
            for name in names:
                e = by_endpoint[name]
                scale = 100. if name == 'nuclear_spot_fraction_dapi' else fig.fraction_scale(name)
                records = []
                record = fig.superplot_standalone(ctx, name, e.pretty(ctx.channel_labels),
                    'percent (%)' if scale == 100 else e.unit, well, field, data['nuclei'],
                    contrasts, e.column, figure_dir, stem+'__'+name, records, scale=scale)
                path = figure_dir/record['png']
                full = figure_dir/record.get('full_png', record['png'])
                assets.append(dict(path=str(path.resolve()), sha256=sha256(path),
                    full_path=str(full.resolve()), full_sha256=sha256(full),
                    cohort=template['cohort'], caption=e.pretty(ctx.channel_labels), single_plot=True))
        elif item.get('kind')=='micrographs':
            assets = [dict(path=entry['path'], sha256=sha256(Path(entry['path'])),
                cohort=template['cohort'], caption=row['group']+' · '+entry['panel'])
                for row in micro for entry in row['panels']]
        else:
            import textwrap
            f = fig.plt.figure(figsize=(12.4,5.6))
            f.text(.06,.88,item['title'],fontsize=18)
            f.text(.06,.72,'\n'.join(textwrap.wrap(item.get('body',''),95)),fontsize=13,va='top',linespacing=1.6)
            if item.get('kind') == 'selection':
                if 'dapi' in data:
                    rows = data['dapi']['arms']
                    lines = [f'{r.group}: {r.retained_nuclei} retained nuclei; {r.unretained_dapi_objects} unretained DAPI objects; '
                             f'{r.below_min_area_px} below {r.nucleus_min_area_px:g} px' for r in rows.itertuples()]
                    lines += ['DAPI area median [IQR], px: ' + '; '.join(
                        f'{r.group} {r.unretained_area_median_px:g} [{r.unretained_area_q25_px:g}, {r.unretained_area_q75_px:g}]'
                        for r in rows.itertuples()),
                        'Connected components may merge touching nuclei; no claim of recovered segmentation exclusions.']
                    f.text(.06,.42,'\n\n'.join(lines),fontsize=11,va='top')
                else:
                    rows = sheets['Nucleus selection']
                    f.text(.06,.32,'\n'.join(f'{r.stage}: {r.value}' for r in rows.itertuples()),fontsize=12,va='top')
            fig.save(f,figure_dir,stem,[],item['title'],'; '.join(extra_sheets))
            path=figure_dir/f'{stem}.png'
            assets=[dict(path=str(path.resolve()),sha256=sha256(path),cohort=template['cohort'])]
        for asset in assets:
            figure_rows.append(dict(slide=slide_no,**asset,source_sheet='Contrasts' if names else 'Micrographs',
                                    source_cells=';'.join(f'{r["sheet"]}!{r["cell"]}' for r in refs),
                                    endpoints=';'.join(names),filter='all retained; endpoint finite/usability masks',
                                    test='two-sided Welch on well means'))
        # Non-figure text is workbook material too: no untraced scientific literals.
        if item.get('body'):
            values.append(dict(slide=slide_no,endpoint='',label='readout',value=item['body'],
                               source_sheet='Deck specification',source_cell=f'B{slide_no+2}'))
            refs.append(dict(sheet='Slide values',cell=f'D{len(values)+2}',label='readout',display=True))
        resolved.append(dict(identity=item['identity'],title=item['title'],question=item.get('question', ''),title_cell=title_ref,values=refs,figures=assets, layout='micrographs' if item.get('kind')=='micrographs' else 'grid'))
    sheets['Slide values']=pd.DataFrame(values)
    sheets['Figure sources']=pd.DataFrame(figure_rows)
    if template.get('simple_coloc_csv'):
        from .coloc_existing import (read_simple_metrics, simple_statistics,
            representative_nuclei, render_simple_coloc, load_representative_pixels,
            render_cytofluorogram)
        nuclei=read_simple_metrics(Path(template['simple_coloc_csv']))
        fields,wells,tests=simple_statistics(nuclei,ctx.reference)
        selected=representative_nuclei(nuclei)
        rna=template.get('simple_coloc_rna_label',ctx.channel_labels['rna1'])
        partner=template.get('simple_coloc_partner_label',ctx.channel_labels['protein'])
        pixels,pixel_sources=load_representative_pixels(selected,ctx.run_dir,
                                                       template['simple_coloc_image_paths'])
        render_simple_coloc(nuclei,fields,wells,tests,figure_dir,rna,partner,ctx.reference)
        render_cytofluorogram(selected,pixels,figure_dir,rna,partner)
        sheets.update({'Simple coloc metrics':tests,'Simple coloc nuclei':nuclei,
                       'Simple coloc FOVs':fields,'Simple coloc wells':wells,
                       'Simple coloc representatives':selected,'Simple coloc pixel sources':pixel_sources})
        resolved.extend(simple_coloc_slides(sheets,[figure_dir/'FIG_SIMPLE_COLOC_focus.png',
            figure_dir/'FIG_CYTOFLUOROGRAM.png'],rna,partner,template['cohort']))
    return dict(schema='resolved-deck-1',cohort=template['cohort'],slides=resolved)
