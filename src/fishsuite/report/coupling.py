"""Descriptive, well-replicated single-plane MIAT/QKI coupling report."""
from __future__ import annotations

import colorsys
import hashlib
import shlex
import textwrap
from pathlib import Path

import numpy as np
import pandas as pd

from . import figures as figlib
from . import workbook
from .coupling_stats import METRICS, prepare_data, summarize

ABUNDANCE = ('n_miat_spots', 'integrated_nuclear_miat')
ASSOCIATION = ('obs_minus_null_frac_miat_spots_qki_pos', 'qki_at_spots_minus_nuclear')
SHEET_ORDER = ['README', 'per_well', 'per_arm', 'descriptive_qc', 'contrast', 'ratio_of_ratios',
               'within_well_correlation', 'sensitivity', 'all_arms']
PERMUTATION_NOTE = '2 v 2 wells: exact permutation has 6 allocations; smallest two-sided p = 0.333'
MARKERS = ('o', '^', 's', 'D')
PLOT_METRICS = (*METRICS, 'mean_nuclear_miat', 'n_miat_spots_per_100um2')
METRIC_LABELS = {
    'n_miat_spots': 'MIAT puncta per nucleus',
    'mean_nuclear_miat': 'Mean nuclear MIAT intensity (a.u.)',
    'qki_pos_area_frac': 'Fraction of nuclear area QKI-positive',
    'miat_pos_area_frac': 'Fraction of nuclear area MIAT-positive',
    'frac_qki_area_on_miat_area': 'Fraction of QKI-positive area on MIAT-positive pixels (raw)',
    'frac_qki_area_on_miat_footprints': 'Fraction of QKI-positive area on MIAT puncta (raw, uncorrected)',
    'n_miat_spots_per_100um2': 'MIAT puncta per 100 µm² nuclear area',
    'integrated_nuclear_miat': 'Total nuclear MIAT intensity (a.u.)',
    'frac_miat_spots_qki_pos': 'Fraction of MIAT puncta with QKI ≥ threshold (single plane)',
    'obs_minus_null_frac_miat_spots_qki_pos': 'QKI-positive MIAT puncta: observed − chance (single plane)',
    'frac_miat_footprint_area_qki_pos': 'Fraction of MIAT punctum area that is QKI-positive (single plane)',
    'obs_minus_null_frac_miat_footprint_area_qki_pos': 'QKI-positive MIAT area: observed − chance (single plane)',
    'mean_qki_at_miat_spots': 'Mean QKI intensity at MIAT puncta (a.u.)',
    'qki_at_spots_minus_nuclear': 'QKI at MIAT puncta − nuclear mean QKI (a.u.)',
    'mean_nuclear_qki': 'Mean nuclear QKI intensity (a.u.)',
    'nuclear_area_um2': 'Nuclear area (µm²)',
    'n_miat_spots_qki_pos': 'QKI-positive MIAT puncta per nucleus (single plane)',
    'sat_frac_miat': 'Fraction of saturated MIAT pixels',
    'sat_frac_qki': 'Fraction of saturated QKI pixels',
}


def _plot_metrics(available):
    available = set(available)
    return [metric for metric in PLOT_METRICS if metric in available]


def _well_markers(df, order):
    return {(arm, well): MARKERS[i % len(MARKERS)] for arm in order
            for i, well in enumerate(sorted(df.loc[df.condition.eq(arm), 'well'].unique()))}


def _well_key(df, order):
    rows = []
    for arm in order:
        for i, well in enumerate(sorted(df.loc[df.condition.eq(arm), 'well'].unique()), 1):
            images = sorted(df.loc[df.condition.eq(arm) & df.well.eq(well), 'image'].dropna().unique())
            for image in images:
                rows.append(dict(condition=arm, well_number=i, well=well, image=image,
                                 marker=MARKERS[(i - 1) % len(MARKERS)]))
    return pd.DataFrame(rows)


def _label(metric):
    return METRIC_LABELS.get(metric, metric)


def _arm_label(arm):
    return '\n'.join(textwrap.fill(part, 12) for part in str(arm).replace(' (', '\n(').splitlines())


def _legend(ax, df, order, means=True):
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    count = max((df.loc[df.condition.eq(arm), 'well'].nunique() for arm in order), default=0)
    handles = [Line2D([], [], linestyle='none', marker=MARKERS[i % len(MARKERS)],
                      color='#777777', markersize=4, label=f'well {i + 1}') for i in range(count)]
    if means:
        handles += [Line2D([], [], linestyle='none', marker='o', markerfacecolor='white',
                           markeredgecolor='#555555', markersize=7, label='well mean'),
                    Patch(facecolor='#d0d0d0', edgecolor='#555555', label='arm mean of well means')]
    ax.figure.legend(handles=handles, loc='lower center', bbox_to_anchor=(.5, .135),
                     ncol=min(6, len(handles)), frameon=False, fontsize=7)


def _chance(ax, metric):
    if metric.startswith('obs_minus_null_'):
        ax.axhline(0, color='#777777', linewidth=.7, linestyle='--', label='chance', zorder=1)
        ax.annotate('chance', xy=(1, 0), xycoords=('axes fraction', 'data'),
                    xytext=(-3, 3), textcoords='offset points', ha='right', fontsize=7)


def _configuration(config, arms, treated):
    import yaml
    from matplotlib.colors import to_rgb

    cfg = yaml.safe_load(Path(config).read_text(encoding='utf-8')) if config else {}
    if cfg is None:
        cfg = {}
    if not isinstance(cfg, dict):
        raise ValueError('config must be a YAML mapping')
    conditions = cfg.get('conditions', {})
    if not isinstance(conditions, dict):
        raise ValueError('conditions must be a YAML mapping')
    requested = cfg.get('condition_order', conditions.get('condition_order', []))
    overrides = cfg.get('group_colors', conditions.get('group_colors', {}))
    if not isinstance(requested, list) or not isinstance(overrides, dict):
        raise ValueError('condition_order must be a list and group_colors a mapping')
    order = list(dict.fromkeys([str(g) for g in requested if str(g) in arms] + list(arms)))
    defaults = {g: figlib.OKABE_ITO['sky'] if g == treated else '#7F7F7F' for g in order}
    defaults.update({str(k): str(v) for k, v in overrides.items()})
    colors = figlib.group_colors(order, defaults)
    hues = [colorsys.rgb_to_hsv(*to_rgb(c)) for c in colors.values()]
    red = any((h < .08 or h > .94) and s > .3 for h, s, v in hues)
    green = any(.20 < h < .48 and s > .3 for h, s, v in hues)
    if red and green:
        raise ValueError('group_colors must not pair red and green')
    return order, colors


def _axis_variants(ax, values):
    """Register display-only bounds with the shared full/focus exporter."""
    vals = np.asarray(values, dtype=float)
    vals = vals[np.isfinite(vals)]
    lo, hi = (float(vals.min()), float(vals.max())) if len(vals) else (0., 1.)
    pad = max(hi - lo, abs(hi) * .1, .01) * .12

    def setter(variant):
        ax.set_ylim(min(0., lo - pad) if variant == 'full' else lo - pad, hi + pad)
    ax._replicate_simple_axis = setter


def _finish(fig, ax, destination, stem, title, filters, footer, manifest):
    figlib.no_box(fig, ax)
    fig.suptitle(title, y=.97, fontsize=10, fontweight='bold')
    ax.set_ylabel('\n'.join(textwrap.fill(line, 42) for line in ax.get_ylabel().splitlines()))
    ax.set_xlabel(textwrap.fill(ax.get_xlabel(), 65))
    foot = f'Descriptive; wells are the replicates. {filters}; {footer}; axis: focus window'
    # Measure at the final font size: keep provenance in at most two lines.
    probe = fig.text(.025, .025, foot, fontsize=7)
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    words, lines, line = foot.split(), [], ''
    for word in words:
        candidate = (line + ' ' + word).strip()
        probe.set_text(candidate)
        if probe.get_window_extent(renderer).width > fig.bbox.width * .95 and line:
            lines.append(line)
            line = word
        else:
            line = candidate
    lines.append(line)
    if len(lines) > 2:
        fig.set_figwidth(fig.get_figwidth() * len(lines) / 2)
        probe.remove()
        return _finish(fig, ax, destination, stem, title, filters, footer, manifest)
    probe.set_text('\n'.join(lines))
    probe.set_va('bottom')
    fig.subplots_adjust(left=1.1 / fig.get_figwidth(), right=.97, top=.88, bottom=.32)
    figlib.save(fig, destination, stem, manifest, title + '; ' + filters, footer)


def _thresholds(df):
    parts = []
    for key in ('threshold_multiplier', 'miat_min_used', 'qki_min_used'):
        values = ','.join(f'{v:g}' for v in sorted(pd.to_numeric(df[key], errors='coerce').dropna().unique())) if key in df else 'missing'
        parts.append(f'{key}={values or "missing"}')
    return '; '.join(parts)


def _distribution(ax, df, base, metric, order, colors, rng, nuclei=True):
    from matplotlib.colors import to_rgba
    shapes = _well_markers(df, order)
    for x, arm in enumerate(order):
        wells = sorted(df.loc[df.condition.eq(arm), 'well'].unique())
        slot = .72 / max(1, len(wells))
        means = []
        for i, well in enumerate(wells):
            offset = -.36 + slot * (i + .5)
            vals = base.loc[base.condition.eq(arm) & base.well.eq(well), metric].dropna().to_numpy()
            if nuclei:
                ax.scatter(x + offset + rng.uniform(-slot * .35, slot * .35, len(vals)), vals,
                           s=10, marker=shapes[arm, well], color=colors[arm], alpha=.45, zorder=3)
            if len(vals):
                means.append(vals.mean())
                ax.scatter(x + offset, vals.mean(), s=65, marker=shapes[arm, well],
                           facecolors='white', edgecolors=colors[arm], linewidths=1.3, zorder=5)
            else:
                means.append(np.nan)
                ax.annotate(f'well {i + 1}: NA', (x + offset, .02),
                            xycoords=('data', 'axes fraction'), ha='center', fontsize=7)
        if means and np.isfinite(means).all():
            ax.bar(x, np.mean(means), width=.8, facecolor=to_rgba(colors[arm], .35),
                   edgecolor=colors[arm], linewidth=1, zorder=0)
    ax.set_xticks(range(len(order)), [_arm_label(arm) for arm in order])
    ax.set_xlim(-.6, len(order) - .4)
    ax.set_ylabel(_label(metric))
    _legend(ax, df, order)
    _chance(ax, metric)
    # All threshold levels contribute bounds, so a metric retains its scale.
    bounds = df.loc[df.condition.isin(order), metric].tolist()
    _axis_variants(ax, [0, *bounds] if metric.startswith('obs_minus_null_') else bounds)


def _figures(df, sheets, treated, control, order, colors, source, out, seed):
    import matplotlib.pyplot as plt

    figlib.set_style()
    rng = np.random.default_rng(seed)
    base = df.loc[df.threshold_multiplier.eq(1.0)]
    footer = f'Input: {"/".join(Path(source).parts[-2:])}; single-plane'
    manifest = []
    shapes = _well_markers(df, order)
    pairs = list(base[['condition', 'well']].drop_duplicates().itertuples(index=False, name=None))
    metrics = _plot_metrics(sheets['sensitivity'].metric)
    size = (max(6.5, 1.1 * len(order) + 1.5), 4.8)
    for metric in metrics:
        for kind, arms, nuclei in (('distribution', order, True),
                                   ('well_means', [a for a in order if a in (treated, control)], False)):
            fig, ax = plt.subplots(figsize=size)
            _distribution(ax, df, base, metric, arms, colors, rng, nuclei=nuclei)
            _finish(fig, ax, out, f'{kind}_{metric}', _label(metric).removesuffix(' (single plane)'),
                    'Filter: threshold_multiplier=1; finite metric; N=0 retained when defined', footer, manifest)

        fig, ax = plt.subplots(figsize=size)
        row = sheets['sensitivity'].set_index('metric').loc[metric]
        levels = sorted(df.threshold_multiplier.unique())
        positions = np.arange(len(levels))
        differences = [row[f'difference_multiplier_{level:g}'] for level in levels]
        ax.plot(positions, differences, marker='o', color=colors[treated])
        ax.set_xticks(positions, [f'{level:g}' for level in levels])
        ax.set_xlim(-.5, len(levels) - .5)
        for position, level, difference in zip(positions, levels, differences):
            if not np.isfinite(difference):
                reason = str(row[f'na_reason_multiplier_{level:g}'])
                ax.scatter([position], [0], transform=ax.get_xaxis_transform(),
                           marker='o', s=45, facecolors='none', edgecolors=colors[treated],
                           linewidths=1.2, clip_on=False, zorder=6)
                ax.annotate('NA\n' + textwrap.fill(reason, 25), xy=(position, 0),
                            xycoords=ax.get_xaxis_transform(), xytext=(0, 8),
                            textcoords='offset points', ha='center', va='bottom', fontsize=7,
                            bbox={'facecolor': 'white', 'edgecolor': 'none', 'pad': 1})
        if metric.startswith('obs_minus_null_'):
            _chance(ax, metric)
        else:
            ax.axhline(0, color='#7F7F7F', linewidth=.7)
        ax.set_xlabel('Threshold multiplier')
        ax.set_ylabel('Treated − control\n' + _label(metric))
        _axis_variants(ax, [0, *differences])
        _finish(fig, ax, out, f'sensitivity_{metric}', _label(metric).removesuffix(' (single plane)'),
                'Filter: treated/control; threshold_multiplier=all; finite metric; equal well weights', footer, manifest)

    for abundance in ABUNDANCE:
        for association in ASSOCIATION:
            fig, ax = plt.subplots(figsize=size)
            for arm, well in pairs:
                subset = base.loc[base.condition.eq(arm) & base.well.eq(well) & base.n_miat_spots.gt(0)]
                subset = subset.dropna(subset=[abundance, association])
                ax.scatter(subset[abundance], subset[association], marker=shapes[arm, well],
                           s=12, alpha=.45, color=colors[arm])
            ax.set_xlabel(_label(abundance))
            ax.set_ylabel(_label(association))
            _legend(ax, df, order, means=False)
            _chance(ax, association)
            bounds = df.loc[df.n_miat_spots.gt(0), association].tolist()
            _axis_variants(ax, [0, *bounds] if association.startswith('obs_minus_null_') else bounds)
            _finish(fig, ax, out, f'scatter_{abundance}_vs_{association}', 'Within-well abundance versus association',
                    'Filter: threshold_multiplier=1; N>0; finite pairs', footer, manifest)

    fig, ax = plt.subplots(figsize=size)
    for arm, well in pairs:
        subset = base.loc[base.condition.eq(arm) & base.well.eq(well) & base.n_miat_spots.gt(0)].dropna(subset=['frac_miat_spots_qki_pos'])
        ax.scatter(subset.n_miat_spots, subset.frac_miat_spots_qki_pos, marker=shapes[arm, well],
                   color=colors[arm], alpha=.45, s=12)
    ax.set_xlabel(_label('n_miat_spots'))
    ax.set_ylabel(_label('frac_miat_spots_qki_pos'))
    _legend(ax, df, order, means=False)
    _axis_variants(ax, [0, 1, *df.frac_miat_spots_qki_pos.dropna()])
    _finish(fig, ax, out, 'precision_funnel', 'Precision funnel',
            'Filter: threshold_multiplier=1; N>0; finite fraction', footer, manifest)
    pd.DataFrame(manifest).to_csv(out / 'figure_manifest.csv', index=False)


def build_coupling(assoc_csv, treated, control, out, config=None, seed=0, n_boot=2000):
    """Build a new report from saved per-nucleus association measurements."""
    from ..core import repro
    from ..core.qki_association import COLUMN_DEFINITIONS

    source = Path(assoc_csv).resolve()
    destination = Path(out).resolve()
    if destination.exists() and any(destination.iterdir()):
        raise ValueError('out must be new or empty; existing report files are not overwritten')
    raw = pd.read_csv(source, dtype={'condition': str, 'well': str, 'image': str})
    sheets = summarize(raw, treated, control, seed=seed, n_boot=n_boot)
    df = prepare_data(raw)
    order, colors = _configuration(config, list(dict.fromkeys(df.condition)), treated)
    notes = [dict(topic='input', detail=str(source)),
             dict(topic='input_sha256', detail=hashlib.sha256(source.read_bytes()).hexdigest()),
             dict(topic='thresholds used', detail=_thresholds(df)),
             dict(topic='display', detail=f'condition_order={order}; group_colors={colors}'),
             dict(topic='bootstrap settings', detail=f'seed={seed}; n_boot={n_boot}; with fewer than 3 wells per arm no interval is reported; read the individual well means. For >=3 wells per arm, exploratory 2.5/97.5 percentiles of treated-control difference; not confirmatory')]
    well_key = _well_key(df, order)
    for row in well_key.itertuples(index=False):
        notes.append(dict(topic=f'figure well key: {row.condition}; well {row.well_number}',
                          detail=f'well={row.well}; image={row.image}; marker={row.marker}'))
    for metric in METRICS:
        notes.append(dict(topic=metric, detail=COLUMN_DEFINITIONS.get(metric, 'Definition missing')))
    for column in df:
        if column.startswith('na_reason_'):
            notes.append(dict(topic=column, detail=str(df[column].fillna('').value_counts().to_dict())))
    # Normalize the README's two descriptive columns without discarding any rows.
    readme = sheets['README'].copy()
    if len(readme.columns) == 2:
        readme.columns = ['topic', 'detail']
    sheets['README'] = pd.concat([readme, pd.DataFrame(notes)], ignore_index=True)
    descriptions = {name: f'{name}: single-plane; biological wells are replicates; descriptive, no nucleus-level tests.' for name in SHEET_ORDER}
    descriptions['contrast'] += ' ' + PERMUTATION_NOTE
    destination.mkdir(parents=True, exist_ok=True)
    well_key.to_csv(destination / 'figure_well_key.csv', index=False)
    workbook.write(destination / 'coupling_summary.xlsx', sheets, order=SHEET_ORDER, descriptions=descriptions)
    # The common writer reserves README wrapping for its legacy 'Read me' name.
    # Keep its table/header formatting and make this command's README readable too.
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    book = load_workbook(destination / 'coupling_summary.xlsx')
    readme_sheet = book['README']
    readme_sheet.column_dimensions['A'].width = 48
    readme_sheet.column_dimensions['B'].width = 110
    for row in readme_sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        lines = max(len(textwrap.wrap(str(cell.value or ''), 46 if cell.column == 1 else 105)) for cell in row)
        readme_sheet.row_dimensions[row[0].row].height = max(30, 15 * lines)
    book.save(destination / 'coupling_summary.xlsx')
    book.close()
    _figures(df, sheets, treated, control, order, colors, source, destination / 'figures', seed)
    command = ['fishsuite', 'coupling', '--assoc-csv', str(source), '--treated', treated,
               '--control', control, '--seed', str(seed), '--n-boot', str(n_boot), '--out', str(destination)]
    if config:
        command += ['--config', str(Path(config).resolve())]
    if not repro.write_command_log(destination, str(Path(config).resolve()) if config else 'none', destination, seed,
                                  extra={'coupling_command': shlex.join(command), 'quantitation': 'single-plane'}):
        raise OSError('failed to write command.log')
    if not repro.write_versions_txt(destination, seed):
        raise OSError('failed to write versions.txt')
    return destination
