"""Persisted fixed-10 MIAT/QKI estimands; no detection or null generation.

R is the ratio of the associated and total KD/NT ratios, not a loss ratio.
The eight q95 tests are exploratory and distinct from the historical gate-nine.
"""
from __future__ import annotations

from itertools import combinations, product
from pathlib import Path
import hashlib
import re
from types import SimpleNamespace

import numpy as np
import pandas as pd
from scipy import stats

POLICIES = (
    'historical_nucleoplasm_centroid_rotation',
    'whole_nucleus_centroid_rotation',
    'compartment_matched_centroid_rotation',
    'whole_nucleus_uniform_center_position',
)
LABELS = ('Historical nucleoplasm', 'Whole-nucleus rotation',
          'Compartment-matched', 'Uniform center')
MEASUREMENTS = ('spot_count', 'miat_intensity')
KEYS = ['biological_set', 'slide', 'arm']
METHOD = 'multivariate_delta_within_arm_covariance'
COHORT = 'biological_fixed10'
COLORS = {'NT': '#595959', 'KD': '#CC79A7'}


def holm(values):
    """Fixed membership, including unavailable tests in the family size."""
    p = np.asarray(values, float)
    out = np.full(len(p), np.nan)
    order = np.argsort(np.where(np.isfinite(p), p, np.inf))
    running = 0.
    for rank, i in enumerate(order):
        if np.isfinite(p[i]):
            running = max(running, min(1., (len(p)-rank)*p[i]))
            out[i] = running
    return out


def ratio_interval(pairs, denominator_source=COHORT, *, expected_wells=6):
    """Six matched (A,T) wells per arm; pairing is within well, never NT/KD."""
    if 'count' in str(denominator_source).lower():
        raise ValueError('COUNT localization cannot supply a ratio denominator')
    if denominator_source not in (COHORT, 'biological_fixed10_usable', 'run_retained'):
        raise ValueError('denominator must be the fixed-10 coloc cohort')
    out = dict.fromkeys(('R', 'rT', 'rA', 'ci_low', 'ci_high', 'se_log_R',
                         'variance_log_R', 'R_MDE_05', 'R_MDE_05_8'), np.nan)
    out.update(reason='', ci_method=METHOD, denominator_source=denominator_source)
    if not set(KEYS + ['A', 'T']).issubset(pairs):
        return dict(out, reason='missing matched well columns')
    if pairs.biological_set.duplicated().any() or pairs[KEYS].isna().any().any():
        return dict(out, reason='duplicate well or missing well metadata')
    if len(pairs) != 2*expected_wells or set(pairs.arm) != {'NT', 'KD'}:
        return dict(out, reason=f'incomplete matched cohort: require {expected_wells} wells per arm')
    variance = 0.
    for arm in ('NT', 'KD'):
        x = pairs.loc[pairs.arm == arm, ['A', 'T']].to_numpy(float)
        out['n_' + arm] = len(x)
        if x.shape != (expected_wells, 2) or not np.isfinite(x).all() or (x < 0).any():
            return dict(out, reason='incomplete/nonfinite/negative matched well values')
        a, t = x.mean(axis=0)
        if not np.isfinite([a, t]).all() or a <= 0 or t <= 0:
            return dict(out, reason='nonpositive/nonfinite arm mean A or denominator T')
        cov = np.cov(x, rowvar=False, ddof=1)
        terms = [cov[0, 0]/(expected_wells*a*a), cov[1, 1]/(expected_wells*t*t),
                 -2*cov[0, 1]/(expected_wells*a*t)]
        v = sum(terms)
        # Only floating-point cancellation may be clamped; not material failure.
        if not np.isfinite(v) or v < -1e-12 * max(1., sum(abs(z) for z in terms)):
            return dict(out, reason='invalid negative/nonfinite log variance')
        variance += max(0., v)
        out.update({f'A_{arm}': a, f'T_{arm}': t, f'var_A_{arm}': cov[0, 0],
                    f'var_T_{arm}': cov[1, 1], f'cov_AT_{arm}': cov[0, 1],
                    f'variance_log_{arm}': max(0., v)})
    rt, ra = out['T_KD']/out['T_NT'], out['A_KD']/out['A_NT']
    r, se = ra/rt, np.sqrt(variance)
    if not np.isfinite([rt, ra, r, se]).all() or min(rt, ra, r) <= 0:
        return dict(out, reason='invalid nonpositive/nonfinite ratio or SE')
    lo, hi = np.exp(np.log(r) + np.array([-1., 1.])*stats.norm.ppf(.975)*se)
    out.update(R=r, rT=rt, rA=ra, ci_low=lo, ci_high=hi,
               variance_log_R=variance, se_log_R=se,
               R_MDE_05=np.exp((stats.norm.ppf(.975)+stats.norm.ppf(.8))*se),
               R_MDE_05_8=np.exp((stats.norm.ppf(1-.05/16)+stats.norm.ppf(.8))*se))
    return out


def unblocked_exact_ratio(pairs, *, expected_wells=3):
    """Complete balanced well-label enumeration for a single unblocked design.

    Uses the same ratio-of-arm-means estimand and |log R| two-sided statistic
    as ``exact_ratio``. The one-sided R>1 tail is separately reported.
    """
    observed = ratio_interval(pairs, 'run_retained', expected_wells=expected_wells)
    if observed['reason']:
        return dict(p_exact=np.nan, n_assignments=0, exact_reason=observed['reason']), pd.DataFrame()
    x = pairs.sort_values('biological_set').reset_index(drop=True)
    rows = []
    for number, selected in enumerate(combinations(range(len(x)), expected_wells)):
        assigned = x.index.isin(selected)
        kd = x.loc[assigned, ['A', 'T']].mean()
        nt = x.loc[~assigned, ['A', 'T']].mean()
        r = (kd['A']/nt['A'])/(kd['T']/nt['T'])
        rows.append(dict(assignment=number, assigned_KD=';'.join(x.loc[assigned, 'biological_set']), R=r))
    registry = pd.DataFrame(rows)
    if not np.isfinite(registry.R).all() or (registry.R <= 0).any():
        return dict(p_exact=np.nan, n_assignments=len(rows), exact_reason='invalid permuted ratio; no assignments dropped'), registry
    extreme = (np.abs(np.log(registry.R)) >= abs(np.log(observed['R']))-1e-12)
    greater = registry.R >= observed['R']-1e-12
    return dict(p_exact=float(extreme.mean()), p_exact_greater=float(greater.mean()),
                n_assignments=len(rows), n_extreme=int(extreme.sum()), exact_reason='',
                exact_statistic='abs(log(R)); complete balanced well-label enumeration; no plus-one'), registry


def assignments(pairs):
    """Enumerate the complete recorded design, with identifiers on every assignment."""
    groups = [g.sort_values('biological_set') for _, g in pairs.groupby('slide')]
    if len(groups) != 2 or any(len(g) != 6 or g.arm.value_counts().to_dict() != {'KD': 3, 'NT': 3} for g in groups):
        raise ValueError('exact test requires two slides, each with three NT and three KD wells')
    rows = []
    for number, choices in enumerate(product(combinations(range(6), 3), repeat=2)):
        for g, selected in zip(groups, choices):
            for i, row in enumerate(g.itertuples(index=False)):
                rows.append(dict(assignment=number, biological_set=row.biological_set,
                                 slide=row.slide, original_arm=row.arm,
                                 assigned_arm='KD' if i in selected else 'NT'))
    return pd.DataFrame(rows)


def exact_ratio(pairs, registry):
    """Two-sided |log R| tail, >= observed - 1e-12; no plus-one correction.

    Recovered from B0 audit_checks.py checks(), which reproduces EFFECTS.
    The explicit original assignment file was not delivered; all assignments
    of the attested design are enumerated and exported here.
    """
    observed = ratio_interval(pairs)
    if observed['reason']:
        return dict(p_exact=np.nan, n_assignments=0, exact_reason=observed['reason'])
    joined = registry.merge(pairs[KEYS+['A', 'T']], on=['biological_set', 'slide'],
                            validate='many_to_one')
    means = joined.groupby(['assignment', 'assigned_arm'])[['A', 'T']].mean()
    if not np.isfinite(means.to_numpy()).all() or (means <= 0).any().any():
        return dict(p_exact=np.nan, n_assignments=400, exact_reason='invalid permuted arm mean; assignments not dropped')
    kd, nt = means.xs('KD', level=1), means.xs('NT', level=1)
    r = (kd['A']/nt['A'])/(kd['T']/nt['T'])
    extreme = int((np.abs(np.log(r)) >= abs(np.log(observed['R']))-1e-12).sum())
    return dict(p_exact=extreme/len(r), n_assignments=len(r), n_extreme=extreme,
                exact_reason='', exact_statistic='abs(log(R)); >= observed - 1e-12; extreme/400')


def scalar_contrast(pairs, column):
    a, b = [pairs.loc[pairs.arm == arm, column].to_numpy(float) for arm in ('KD', 'NT')]
    if len(a) != 6 or len(b) != 6 or not np.isfinite(np.r_[a, b]).all():
        return dict(p_welch=np.nan, hedges_g=np.nan, reason='incomplete well vectors')
    pooled = np.sqrt((a.var(ddof=1)+b.var(ddof=1))/2)
    return dict(mean_KD=a.mean(), mean_NT=b.mean(), n_KD=6, n_NT=6,
                p_welch=stats.ttest_ind(a, b, equal_var=False).pvalue,
                hedges_g=(1-3/39)*(a.mean()-b.mean())/pooled if pooled > 0 else np.nan,
                reason='' if pooled > 0 else 'zero pooled SD: standardized effect unavailable')


def _read(path):
    df = pd.read_csv(path)
    df['source_file'] = str(Path(path).resolve())
    df['source_row'] = np.arange(len(df))+2
    return df


def _one(df, label):
    if len(df) != 1:
        raise ValueError(f'{label}: expected exactly one source row, got {len(df)}')
    return df.iloc[0]


def _hierarchy(nuclei, columns):
    # No silent NA-dropping: missing measurements invalidate the complete well.
    mean = lambda x: x.mean(skipna=False)
    fov = nuclei.groupby(KEYS+['image_key'], dropna=False)[columns].agg(mean)
    return fov.groupby(KEYS)[columns].agg(mean).reset_index()


def analyze(data_dir, count_dir=None):
    """Validate persisted sources, returning complete workbook-ready tables."""
    data_dir = Path(data_dir).resolve()
    if 'MIAT_COUNT' in str(data_dir).upper():
        raise ValueError('COUNT data cannot supply coloc ratio denominators')
    values = _read(data_dir/'biological_endpoint_values.csv')
    effects = _read(data_dir/'biological_effect_summary.csv')
    # Cell citations are checked against the delivered workbook, not inferred
    # merely from the CSV order. The workbook is never saved or altered.
    import openpyxl
    book = openpyxl.load_workbook(data_dir.parent/'MIAT_QKI_RESULTS.xlsx', read_only=True, data_only=True)
    try:
        for name, csv in (('Bio_Endpoints', values), ('Bio_Effects', effects)):
            rows = list(book[name].values)
            delivered = pd.DataFrame(rows[1:], columns=rows[0])
            if len(delivered) != len(csv):
                raise ValueError(f'{name}: workbook/CSV row count mismatch')
            for column in ('tier', 'null_policy', 'measurement'):
                if not delivered[column].equals(csv[column]):
                    raise ValueError(f'{name}: workbook/CSV selector mismatch: {column}')
            for column in csv.select_dtypes(include='number'):
                if column != 'source_row':
                    np.testing.assert_allclose(pd.to_numeric(delivered[column]), csv[column], rtol=0, atol=1e-9, equal_nan=True)
    finally:
        book.close()
    hist = _read(data_dir/'pooled_Fig3_source_data.csv')
    nuc = _read(data_dir/'spatial_null_nuclei.csv.gz')
    nuc = nuc[nuc.cohort == COHORT].copy()
    if set(nuc.null_policy) != set(POLICIES):
        raise ValueError('fixed-10 policy roster differs from the frozen four policies')
    if nuc.duplicated(['null_policy', 'nucleus_uid']).any():
        raise ValueError('duplicate policy nucleus')
    for policy, n in nuc.groupby('null_policy'):
        if len(n) != 370 or not n.groupby('image_key').size().eq(10).all():
            raise ValueError('fixed-10 requires 370 nuclei and ten per FOV')
        if n.arm.value_counts().to_dict() != {'KD': 190, 'NT': 180}:
            raise ValueError('fixed-10 arm roster differs')
        if (n.observed_spot_count == 0).sum() != 4:
            raise ValueError('four restored zero-spot nuclei required')
        if not n.loc[n.observed_spot_count == 0, 'roster_zero_spot_filled'].astype(str).str.lower().eq('true').all():
            raise ValueError('restored zero-spot status missing')
        np.testing.assert_allclose(n.observed_spot_count, n.n_spots_all, rtol=0, atol=1e-9)
        c = n[['observed_spot_count', 'candidate_spot_count', 'usable_spot_count', 'q95_positive_spot_count']].to_numpy(float)
        if not np.isfinite(c).all() or (c < 0).any() or (np.diff(c, axis=1) > 0).any():
            raise ValueError('invalid coverage partition: require observed >= candidate >= usable >= positive >= 0')
    roster = nuc[nuc.null_policy == POLICIES[0]].sort_values('nucleus_uid')
    historical_roster = hist[(hist.panel == 'A-total') & (hist.tier == 'nucleus') & (hist.endpoint == 'n_spots_floor')].copy()
    historical_roster['nucleus_uid'] = historical_roster.image_key+':nucleus:'+historical_roster.nucleus_id.astype(int).astype(str)
    joined_roster = roster.merge(historical_roster[['nucleus_uid', 'value']], on='nucleus_uid', how='outer', validate='one_to_one')
    if len(joined_roster) != 370 or joined_roster.value.isna().any():
        raise ValueError('historical fixed-10 nucleus identity reconciliation failed')
    np.testing.assert_allclose(joined_roster.observed_spot_count, joined_roster.value, rtol=0, atol=1e-9)
    for policy in POLICIES[1:]:
        other = nuc[nuc.null_policy == policy].sort_values('nucleus_uid')
        pd.testing.assert_frame_equal(roster[['nucleus_uid', 'image_key']+KEYS].reset_index(drop=True),
                                      other[['nucleus_uid', 'image_key']+KEYS].reset_index(drop=True))
    well = values[values.tier == 'biological_set_mean']
    ratios, wells, scalars, sensitivities = [], [], [], []
    registry = None
    # Historical count and intensity gates are evaluated before any new interval.
    for policy in POLICIES:
        for measurement in MEASUREMENTS:
            total = well[(well.measurement == measurement) & (well.scope == 'all_detected_miat') &
                         (well.null_policy == 'not_applicable_global')]
            associated = well[(well.measurement == measurement) & (well.scope == 'qki_enriched_exact_miat_footprints') &
                              (well.null_policy == policy)]
            p = total[KEYS+['value', 'source_row']].merge(associated[KEYS+['value', 'source_row']],
                    on=KEYS, how='outer', suffixes=('_T', '_A'), validate='one_to_one').rename(columns={'value_T':'T', 'value_A':'A'})
            p = p.sort_values(['slide', 'biological_set']).reset_index(drop=True)
            result = ratio_interval(p)
            if result['reason']:
                raise ValueError('delivered well reconstruction failed: '+result['reason'])
            if registry is None:
                registry = assignments(p)
            result.update(null_policy=policy, measurement=measurement)
            delivered = _one(effects[(effects.tier == 'biological_set_mean') &
                                    (effects.null_policy == policy) & (effects.measurement == measurement)], 'EFFECTS')
            for key, source in [('R','ratio_of_ratios'), ('rT','global_kd_over_nt'), ('rA','associated_kd_over_nt'),
                                ('T_NT','global_mean_nt'), ('T_KD','global_mean_kd'),
                                ('A_NT','associated_mean_nt'), ('A_KD','associated_mean_kd')]:
                np.testing.assert_allclose(result[key], delivered[source], rtol=0, atol=1e-9)
            if policy == POLICIES[0]:
                endpoint = ('threshold_positive_spots_per_nucleus_q95' if measurement == 'spot_count'
                            else 'miat_footprint_mass_q95_positive_spot_summed')
                h = _one(hist[(hist.panel == 'C') & (hist.cohort == 'sampled_primary') &
                              (hist.thr == 'q95') & (hist.numerator_endpoint == endpoint)], 'historical CI')
                for key, source in [('R','ratio_of_ratios'), ('ci_low','ratio_of_ratios_ci95_low'), ('ci_high','ratio_of_ratios_ci95_high')]:
                    np.testing.assert_allclose(result[key], h[source], rtol=0, atol=1e-9)
                result['historical_ci_source'] = f"{h.source_file}:row {h.source_row}"
                result['historical_delta_p'] = h.p_two_sided
            result.update(exact_ratio(p, registry))
            result['p_exact_delivered'] = delivered.ratio_of_ratios_exact_permutation_p_two_sided
            np.testing.assert_allclose(result['p_exact'], result['p_exact_delivered'], rtol=0, atol=1e-9)
            if delivered.ratio_of_ratios_n_permutations != 400 or not delivered.ratio_of_ratios_exact_permutation_supported:
                raise ValueError('delivered exact-test design unsupported')
            result['source_effect_cells'] = f"{data_dir.parent/'MIAT_QKI_RESULTS.xlsx'}:Bio_Effects!T{delivered.source_row},Y{delivered.source_row}"
            result['source_effect_csv'] = f"{delivered.source_file}:row {delivered.source_row}"
            # Policy-specific columns ONLY; legacy aliases repeat historical data.
            cols = ['observed_spot_count', 'q95_positive_spot_count', 'q95_associated_miat_intensity']
            recon = _hierarchy(nuc[nuc.null_policy == policy], cols)
            check = p.merge(recon, on=KEYS, validate='one_to_one')
            acol = 'q95_positive_spot_count' if measurement == 'spot_count' else 'q95_associated_miat_intensity'
            np.testing.assert_allclose(check.A, check[acol], rtol=0, atol=1e-9)
            if measurement == 'spot_count':
                np.testing.assert_allclose(check['T'], check.observed_spot_count, rtol=0, atol=1e-9)
            p['null_policy'], p['measurement'] = policy, measurement
            p['source_file'] = str(data_dir/'biological_endpoint_values.csv')
            wells.append(p)
            ratios.append(result)
            for col in ('A', 'T') if policy == POLICIES[0] else ('A',):
                endpoint = f'{col}:{measurement}'+(f':{policy}' if col == 'A' else '')
                scalars.append(dict(endpoint=endpoint, measurement=measurement, pool=col,
                                    null_policy=policy if col == 'A' else 'not_applicable_global',
                                    family='exploratory_scalar_10', **scalar_contrast(p, col),
                                    p_exact_historical=delivered.associated_exact_permutation_p_two_sided if col == 'A' else delivered.global_exact_permutation_p_two_sided))
    ratios = pd.DataFrame(ratios)
    ratios['p_exact_holm8'] = holm(ratios.p_exact)
    ratios['interval_scope'] = 'marginal 95%, not simultaneous; normal log-delta, not exact-test inversion'
    for r in ratios.to_dict('records'):
        for alpha, mde in ((.05, r['R_MDE_05']), (.05/8, r['R_MDE_05_8'])):
            sensitivities.append(dict(null_policy=r['null_policy'], measurement=r['measurement'],
                alpha=alpha, power=.8, se_log_R=r['se_log_R'], R_MDE=mde,
                percent_relative_retention=100*(mde-1),
                formula='exp((norm.ppf(1-alpha/2)+norm.ppf(0.8))*SE(log R))',
                method='plug-in normal detection threshold around R=1, not exact-permutation power',
                **{k:v for k,v in r.items() if k.startswith(('cov_', 'var_'))}))
    scalar = pd.DataFrame(scalars)
    scalar['p_welch_holm10'] = holm(scalar.p_welch)
    coverage, usable = coverage_tables(nuc)
    multiplicity = historical_registry(data_dir)
    for r in ratios.itertuples():
        multiplicity.append(dict(family='exploratory_R_8', endpoint=f'{r.measurement}:{r.null_policy}',
                                 family_size=8, raw_p=r.p_exact, adjusted_p=r.p_exact_holm8,
                                 test='two-sided exact blocked |log R|', disposition='exploratory; original gate failed'))
    for r in scalar.itertuples():
        multiplicity.append(dict(family=r.family, endpoint=r.endpoint, family_size=10,
                                 raw_p=r.p_welch, adjusted_p=r.p_welch_holm10, test='raw Welch, six vs six wells',
                                 disposition='new exploratory scalar family; T deduplicated by measurement'))
    union = hist[(hist.panel == 'C') & (hist.cohort == 'sampled_primary') & (hist.thr == 'q95') &
                 (hist.numerator_endpoint == 'miat_footprint_mass_q95_positive_union_deduplicated')].copy()
    return {'Ratio intervals': ratios, 'Retention sensitivity': pd.DataFrame(sensitivities),
            'Coverage': coverage, 'Usable pool ratios': usable, 'Ratio well values': pd.concat(wells, ignore_index=True),
            'Nucleus roster': nuc, 'Assignments': registry, 'Scalar contrasts': scalar,
            'Multiplicity plan': pd.DataFrame(multiplicity), 'COUNT localization': count_localization(count_dir),
            'Historical union sensitivity': union}


def coverage_tables(nuc):
    rows, matched = [], []
    for policy in POLICIES:
        n = nuc[nuc.null_policy == policy].copy()
        cols = ['observed_spot_count', 'candidate_spot_count', 'usable_spot_count', 'q95_positive_spot_count']
        w = _hierarchy(n, cols)
        w['noncandidate'] = w.observed_spot_count-w.candidate_spot_count
        w['candidate_unusable'] = w.candidate_spot_count-w.usable_spot_count
        w['usable_negative'] = w.usable_spot_count-w.q95_positive_spot_count
        w['usable_positive'] = w.q95_positive_spot_count
        for endpoint, num, den in [('candidate/observed', cols[1], cols[0]),
                                    ('usable/candidate', cols[2], cols[1]), ('usable/observed', cols[2], cols[0])]:
            w['value'] = w[num]/w[den].where(w[den] > 0)
            for r in w.to_dict('records'):
                rows.append(dict(r, null_policy=policy, endpoint=endpoint, tier='well',
                                 reason='' if np.isfinite(r['value']) else 'nonpositive coverage denominator'))
            for arm, g in w.groupby('arm'):
                x = g.value[np.isfinite(g.value)]
                se = x.std(ddof=1)/np.sqrt(len(x)) if len(x) > 1 else np.nan
                rows.append(dict(null_policy=policy, endpoint=endpoint, tier='arm', arm=arm,
                                 value=x.mean(), sd=x.std(ddof=1), n_finite=len(x), n_wells=len(g),
                                 ci_low=x.mean()-stats.t.ppf(.975, len(x)-1)*se,
                                 ci_high=x.mean()+stats.t.ppf(.975, len(x)-1)*se,
                                 reason='t interval on well coverage ratios; not clipped to [0,1]'))
        pair = w[KEYS+['q95_positive_spot_count', 'usable_spot_count']].rename(
            columns={'q95_positive_spot_count':'A', 'usable_spot_count':'T'})
        matched.append(dict(null_policy=policy, measurement='spot_count', **ratio_interval(pair, 'biological_fixed10_usable')))
        matched.append(dict(null_policy=policy, measurement='miat_intensity', R=np.nan,
                            reason='missing summed MIAT mass on the same usable footprint support; not supplied by NUC'))
    return pd.DataFrame(rows), pd.DataFrame(matched)


def historical_registry(data_dir):
    """B0-resolved named sources only; missing registry evidence stays explicit."""
    records = []
    nine_path = data_dir.parent/'figures_v2_2026-09-03/derived_nine_test_holm_family.csv'
    gate_path = data_dir.parent.parent/'DELIVERY_MIAT_QKI_CANONICAL_2026-08-31/data/compact/base_coloc_inference_primary_sampled_primary.csv'
    if nine_path.is_file():
        nine = _read(nine_path)
        if len(nine) != 9 or nine.endpoint.nunique() != 9:
            raise ValueError('historical registry must contain nine distinct endpoints')
        np.testing.assert_allclose(holm(nine.raw_p), nine.holm_p_m9, rtol=0, atol=1e-9)
        for r in nine.itertuples():
            records.append(dict(family='frozen_historical_9', endpoint=r.endpoint, family_size=9,
                raw_p=r.raw_p, adjusted_p=r.holm_p_m9, test='historical exact scalar test',
                disposition='downstream exploratory: original gate failed; no temporal prespecification claim',
                source_file=r.source_file, source_row=r.source_row))
    else:
        records.append(dict(family='frozen_historical_9', family_size=9, disposition='missing original nine-test registry', source_file=str(nine_path)))
    if gate_path.is_file():
        gate = _one(_read(gate_path), 'original primary gate')
        records.append(dict(family='original_primary_gate', endpoint=gate.endpoint, family_size=1,
            raw_p=gate.permutation_p_exact_two_sided, adjusted_p=gate.permutation_p_holm,
            p_welch=gate.welch_p_two_sided, observed_difference=gate.permutation_observed_difference_kd_minus_nt,
            n_extreme=gate.permutation_extreme_permutations_two_sided, n_assignments=gate.permutation_n_permutations,
            disposition='failed original gate; no temporal prespecification claim',
            source_file=gate.source_file, source_row=gate.source_row))
    else:
        records.append(dict(family='original_primary_gate', disposition='missing original gate source', source_file=str(gate_path)))
    return records


def count_localization(count_dir):
    if count_dir is None:
        return pd.DataFrame([dict(reason='COUNT report directory not supplied; separate readout only', denominator_allowed=False)])
    path = Path(count_dir)/'per_well.csv'
    if not path.is_file():
        return pd.DataFrame([dict(reason=f'missing {path}', denominator_allowed=False)])
    d = _read(path)
    d = d[d.endpoint == 'rna1_nuclear_spot_fraction'].copy()
    d['denominator_allowed'] = False
    d['cohort'] = 'COUNT all-eligible, separate detection/run; never a coloc denominator'
    return d


def policy_caption(r, label):
    prefix = (f"Detected nuclear MIAT retained rT={r['rT']:.4f} of NT; the {label} q95-positive pool retained "
              f"rA={r['rA']:.4f}. The tested protection prediction is R > 1. ")
    detected = r['R'] > 1 and r['p_exact_holm8'] < .05
    conclusion = ('Preferential spatial retention was detected' if detected else
                  'Preferential spatial retention was not detected')
    text = (f"{conclusion} under {label}; R={r['R']:.4f}, marginal 95% log-delta CI "
            f"[{r['ci_low']:.4f}, {r['ci_high']:.4f}], exact within-slide p={r['p_exact']:.4g}, "
            f"exploratory Holm p={r['p_exact_holm8']:.4g}. ")
    if r['ci_high'] > 1:
        text += f"Retention above {100*(r['ci_high']-1):.1f}% is outside this interval; "
        text += 'smaller protection remains compatible. ' if r['ci_low'] <= 1 else 'the interval lies above R=1. '
    else:
        text += 'This interval does not include positive relative retention. '
    if r['ci_low'] > 1 and not detected:
        text += 'The log-delta interval and exact/Holm decision differ. '
    text += (f"Minimum detectable retention at 80% power = {100*(r['R_MDE_05']-1):.1f}% "
             f"(R={r['R_MDE_05']:.4f}, alpha=.05) and {100*(r['R_MDE_05_8']-1):.1f}% "
             f"(R={r['R_MDE_05_8']:.4f}, alpha=.05/8). ")
    return prefix+text+'This is not evidence of equivalence, absence of binding, or molecular protection.'


def _context(data_dir, colors):
    return SimpleNamespace(group_order=['NT', 'KD'], reference='NT', colors=colors,
                           technical_layer='none', run_path=str(data_dir), footer=lambda x:x)


def make_figure(sheets, data_dir, measurement='spot_count', variant='full', colors=None):
    """Reuse the locked mean/SD/well renderer; one R per policy, no pseudo-dots."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FixedLocator, FuncFormatter, NullFormatter
    from . import figures
    figures.set_style()
    colors = colors or COLORS
    fig = plt.figure(figsize=(15, 10))
    grid = fig.add_gridspec(2, 5, top=.90, bottom=.27, left=.065, right=.97,
                           height_ratios=[1, 1.12], hspace=.65, wspace=.85)
    ctx = _context(data_dir, colors)
    w = sheets['Ratio well values']
    w = w[w.measurement == measurement]
    scalar = sheets['Scalar contrasts']
    scalar = scalar[scalar.measurement == measurement]
    unit = 'puncta per nucleus' if measurement == 'spot_count' else 'MIAT footprint intensity\nper nucleus (a.u.)'
    axes = []
    for i, policy in enumerate((None,)+POLICIES):
        col, p = ('T', POLICIES[0]) if policy is None else ('A', policy)
        pw = w[w.null_policy == p].rename(columns={'arm':'group', 'biological_set':'well_id', col:'well_mean_of_field_values'}).copy()
        pw['endpoint'] = 'miat_scalar'
        c = scalar[(scalar.pool == col) & ((scalar.null_policy == p) if col == 'A' else True)].copy()
        c['endpoint'], c['test_group'], c['reference_group'] = 'miat_scalar', 'KD', 'NT'
        c['p_welch_holm_within_family'] = c.p_welch_holm10
        ax = fig.add_subplot(grid[0, i]); axes.append(ax)
        figures.draw_replicate_simple(ax, ctx, 'miat_scalar', pw, pd.DataFrame(), pd.DataFrame(), c,
                                      (unit if i == 0 else 'q95-positive '+unit if i == 1 else ''), None, compact=False)
        ax._replicate_simple_axis(variant)
        if variant == 'focus':
            lows = []
            for _, group in pw.groupby('group'):
                v = group.well_mean_of_field_values.to_numpy(float)
                lows.extend([v.min(), v.mean()-v.std(ddof=1)])
            floor = max(0., .8*min(lows))
            ax.set_ylim(floor, ax.get_ylim()[1])
            for bar in ax.patches:
                if (bar.get_gid() or '').startswith('group-mean:'):
                    mean = bar.get_y()+bar.get_height()
                    bar.set_y(floor); bar.set_height(mean-floor)
        ax.set_title('A  Total nuclear MIAT' if policy is None else f'B{i}  {LABELS[i-1]}', fontsize=9, pad=15)
        r = sheets['Ratio intervals']
        r = r[(r.measurement == measurement) & (r.null_policy == p)].iloc[0]
        ax.text(.5, -.24, f"{'rT' if col == 'T' else 'rA'} = {r['rT' if col == 'T' else 'rA']:.3f}; 6 wells/arm\n"
                f"Holm10 p={c.iloc[0].p_welch_holm10:.3g}; g={c.iloc[0].hedges_g:.2f}",
                transform=ax.transAxes, ha='center', fontsize=7.5, linespacing=1.6)
    # Identical associated-pool axis scale across all four facets, including SD/brackets.
    lower = min(ax.get_ylim()[0] for ax in axes[1:]); upper = max(ax.get_ylim()[1] for ax in axes[1:])
    for ax in axes[1:]:
        ax.set_ylim(lower, upper)
        for bar in ax.patches:
            if (bar.get_gid() or '').startswith('group-mean:'):
                mean = bar.get_y()+bar.get_height()
                bar.set_y(lower)
                bar.set_height(mean-lower)
    forest = fig.add_subplot(grid[1, :3])
    forest.set_position([.17, .30, .40, .26])
    selected = sheets['Ratio intervals'].set_index(['measurement', 'null_policy'])
    all_limits = [1.]
    for i, policy in enumerate(POLICIES):
        r = selected.loc[(measurement, policy)]
        forest.errorbar(r.R, 3-i, xerr=[[r.R-r.ci_low], [r.ci_high-r.R]],
                        fmt='o', color=colors['KD'], capsize=4, markersize=6)
        forest.plot(r.R_MDE_05, 3-i+.12, marker='|', color='#0072B2', markersize=16,
                    linestyle='none', label='80% power MDE, alpha=.05' if i == 0 else None)
        forest.plot([r.R_MDE_05_8]*2, [3-i-.2, 3-i+.2], '--', color='#0072B2',
                    linewidth=1.2, label='80% power MDE, alpha=.05/8' if i == 0 else None)
        forest.text(r.R, 3-i-.24, f"exact p={r.p_exact:.3g}; Holm8={r.p_exact_holm8:.3g}",
                    ha='center', va='top', fontsize=7)
        all_limits.extend([r.ci_low, r.ci_high, r.R_MDE_05, r.R_MDE_05_8])
    forest.axvline(1, color='#595959', linestyle=':', linewidth=1)
    forest.set_xscale('log'); forest.set_ylim(-.5, 3.6)
    forest.set_xlim(min(all_limits)*.9, max(all_limits)*1.14)
    ticks = [.4, .5, .6, .8, 1., 1.2, 1.5, 2., 2.5, 3., 4.]
    forest.xaxis.set_major_locator(FixedLocator([x for x in ticks if min(all_limits)*.9 < x < max(all_limits)*1.14]))
    forest.xaxis.set_major_formatter(FuncFormatter(lambda value, pos: f'{value:g}'))
    forest.xaxis.set_minor_formatter(NullFormatter())
    forest.set_yticks([3, 2, 1, 0], LABELS, fontsize=8)
    forest.set_xlabel('Relative retention R = rA / rT (log axis)', fontsize=10)
    forest.set_title('C  Policy-specific relative retention and marginal 95% CI', fontsize=10, pad=12)
    forest.legend(loc='upper center', bbox_to_anchor=(.5, -.18), ncol=2, fontsize=7, frameon=False)
    note = fig.add_subplot(grid[1, 3:]); note.axis('off')
    lines = ['Observed q95-positive pool among all detected MIAT', 'Coverage (mean of six well ratios): usable / observed']
    cov = sheets['Coverage']
    for policy, label in zip(POLICIES, LABELS):
        c = cov[(cov.null_policy == policy) & (cov.tier == 'arm') & (cov.endpoint == 'usable/observed')].set_index('arm')
        lines += [f"{label}: NT {100*c.loc['NT','value']:.1f}%, KD {100*c.loc['KD','value']:.1f}%"]
    lines += ['', 'A is an operational spatial category, not bound molecules.',
              'Null-unusable calls are unknown, never biological negatives.',
              'Matched usable-pool sensitivity is reported separately.',
              'Raw Welch brackets; scalar Holm family = 10.',
              'R family = 8 exploratory exact tests (count + intensity).',
              'Original primary gate failed; historical Holm-nine is separate.']
    note.text(0, 1, '\n'.join(lines), va='top', fontsize=8, linespacing=1.7)
    fig.suptitle('MIAT level to QKI spatial association' + (' | intensity sensitivity' if measurement != 'spot_count' else ''), fontsize=15, y=.975)
    fig.text(.5, .935, ('Full view: zero-baseline scalar columns; all wells, SD, CIs and MDEs' if variant == 'full' else
                      'Focus view: expanded scalar axes; identical wells and inference; no points, CIs or MDEs clipped'), ha='center', fontsize=9)
    # Compact policy wording on the figure; complete wording is in READOUT.md.
    y = .205
    for policy, label in zip(POLICIES, LABELS):
        r = selected.loc[(measurement, policy)]
        fig.text(.04, y, f"{label}: preferential spatial retention not detected; R={r.R:.3f}, 95% CI [{r.ci_low:.3f}, {r.ci_high:.3f}]. "
                 f"Retention above {100*(r.ci_high-1):.1f}% is outside this interval; smaller protection remains compatible.\n"
                 f"Minimum detectable retention at 80% power: {100*(r.R_MDE_05-1):.1f}% (alpha=.05), {100*(r.R_MDE_05_8-1):.1f}% (alpha=.05/8).", fontsize=7.4, linespacing=1.3)
        y -= .036
    fig.text(.04, .050, 'The tested protection prediction is R > 1. MDE markers are plug-in normal 80% power thresholds, not exclusion bounds or exact-test power.\n'
             'Fixed-10 coloc: 370 nuclei / 37 FOVs (NT 180, KD 190); four restored zero nuclei; six wells per arm; equal-weight nucleus → FOV → well means.\n'
             'Shared single-plane nuclear detection; columns ± SD, R ± marginal log-delta CI (not exact-test inversion). COUNT has a separate cohort and denominator.\n'
             'Independent plate provenance, orthogonal KD, laser constancy and specificity/registration/PSF validation remain unverified. Spatial association is not molecular protection.',
             fontsize=6.8, va='top', linespacing=1.3)
    return fig


def render(sheets, data_dir, out_dir, colors):
    import matplotlib.pyplot as plt
    from .provenance import guard_output
    out_dir = guard_output(Path(out_dir))
    out_dir.mkdir(parents=True, exist_ok=True)
    records = []
    for measurement in MEASUREMENTS:
        for variant in ('full', 'focus'):
            fig = make_figure(sheets, data_dir, measurement, variant, colors)
            stem = 'FIG_MIAT_LEVEL_TO_QKI_ASSOCIATION'+('_intensity' if measurement == 'miat_intensity' else '')+'_'+variant
            for ext in ('png', 'svg'):
                path = guard_output(out_dir/(stem+'.'+ext))
                fig.savefig(path, dpi=600, facecolor='white')
                records.append(dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                                    measurement=measurement, variant=variant, figure='ratio',
                                    source_sheet='Ratio intervals; Ratio well values; Coverage; Scalar contrasts',
                                    cohort=COHORT, NT=colors['NT'], KD=colors['KD']))
            plt.close(fig)
    # Coverage has its own well dots and finite-n uncertainty; no ratio pseudo-replicates.
    from . import figures
    for variant in ('full', 'focus'):
        fig, axes = plt.subplots(3, 4, figsize=(13, 10))
        cov = sheets['Coverage']
        for i, endpoint in enumerate(('candidate/observed', 'usable/candidate', 'usable/observed')):
            for j, policy in enumerate(POLICIES):
                d = cov[(cov.endpoint == endpoint) & (cov.null_policy == policy) & (cov.tier == 'well')].copy()
                d = d.rename(columns={'arm':'group', 'biological_set':'well_id', 'value':'well_mean_of_field_values'})
                d['endpoint'] = 'coverage'
                figures.draw_replicate_simple(axes[i, j], _context(data_dir, colors), 'coverage', d,
                    pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), endpoint if j == 0 else '', None)
                axes[i, j]._replicate_simple_axis(variant)
                bottom, top = axes[i, j].get_ylim()
                axes[i, j].set_ylim(bottom, top+.04*(top-bottom))
                axes[i, j].set_title(LABELS[j]+f' | finite n={d.groupby("group").well_mean_of_field_values.count().min()}/arm', fontsize=8)
        fig.suptitle('Coverage of spatial-null calls | fixed-10 | '+variant, fontsize=14)
        fig.text(.05, .02, 'Six well means per arm; columns ± SD. Arm t intervals and finite n in Coverage sheet. '
                 'Unknown calls are not negatives. Full and focus retain identical points.', fontsize=8)
        fig.subplots_adjust(top=.92, bottom=.09, hspace=.42, wspace=.4)
        for ext in ('png', 'svg'):
            path = guard_output(out_dir/f'FIG_MIAT_QKI_COVERAGE_{variant}.{ext}')
            fig.savefig(path, dpi=600, facecolor='white')
            records.append(dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                                measurement='coverage', variant=variant, figure='coverage', source_sheet='Coverage',
                                cohort=COHORT, NT=colors['NT'], KD=colors['KD']))
        plt.close(fig)
    return pd.DataFrame(records)


def build(data_dir, out_dir, count_dir=None, make_figures=True):
    from openpyxl.utils import get_column_letter
    from . import workbook
    from .provenance import guard_output
    out_dir = guard_output(Path(out_dir).resolve())
    out_dir.mkdir(parents=True, exist_ok=True)
    sheets = analyze(data_dir, count_dir)
    colors = COLORS.copy()
    index = Path(data_dir).parent/'coloc_standard_panel_2026-09-04_2300/FIGURE_INDEX.md'
    if index.is_file():
        match = re.search(r'KD[:]\s*`(#[0-9A-Fa-f]{6})`', index.read_text(encoding='utf-8'))
        if match:
            colors['KD'] = match.group(1)
    figures = render(sheets, data_dir, out_dir/'figures', colors) if make_figures else pd.DataFrame([dict(reason='figures disabled')])
    sheets['Figure sources'] = figures
    # Addresses are resolved against the exact final row/column order (+ description/header).
    cells = []
    for name in ('Ratio intervals', 'Retention sensitivity', 'Ratio well values', 'Coverage', 'COUNT localization', 'Scalar contrasts'):
        for row_index, (_, row) in enumerate(sheets[name].iterrows()):
            for col_index, (column, value) in enumerate(row.items(), 1):
                if isinstance(value, (int, float, np.number)) and np.isfinite(value):
                    cells.append(dict(sheet=name, cell=f'{get_column_letter(col_index)}{row_index+3}',
                        quantity=column, value=value, null_policy=row.get('null_policy', ''),
                        measurement=row.get('measurement', ''), biological_set=row.get('biological_set', ''),
                        arm=row.get('arm', ''), source=row.get('source_effect_csv', row.get('source_file', 'derived from matched Ratio well values'))))
    sheets['Slide values'] = pd.DataFrame(cells)
    sources = []
    for path in [Path(data_dir)/f for f in ('biological_endpoint_values.csv', 'biological_effect_summary.csv', 'pooled_Fig3_source_data.csv', 'spatial_null_nuclei.csv.gz')]+[index, Path(data_dir).parent/'MIAT_QKI_RESULTS.xlsx']:
        if path.is_file():
            sources.append(dict(path=str(path.resolve()), sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    for path in sheets['Multiplicity plan'].source_file.dropna().unique():
        p = Path(path)
        sources.append(dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    if count_dir:
        p = Path(count_dir)/'per_well.csv'
        if p.is_file():
            sources.append(dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    sheets['Source provenance'] = pd.DataFrame(sources)
    xlsx = guard_output(out_dir/'REPORT.xlsx')
    workbook.write(xlsx, sheets, order=list(sheets), descriptions={
        'Multiplicity plan': 'Separate frozen historical primary gate and Holm-nine scalar registry, exploratory Holm-eight R family, and new Holm-ten scalar family. The original gate failed; q95 remains exploratory. No prespecification inferred from timestamps.',
        'Figure sources': 'Rendered MIAT/QKI count, intensity and coverage figures with SHA256 hashes and resolved NT/KD palette. Values resolve through Slide values to Ratio intervals, Ratio well values, Scalar contrasts and Coverage cells.',
        'Slide values': 'Every finite numeric reporting value with its exact worksheet cell after the description and header rows. Ratio estimates, intervals, p-values and MDEs trace to paired fixed-10 wells; COUNT cells belong only to the separate localization readout.'})
    caption = ['# MIAT level to QKI spatial association', '',
        'Fixed-10 coloc: 370 nuclei, 37 FOVs, NT 180/KD 190, four restored zero-spot nuclei; six biological wells per arm. Nucleus means within FOV, equal-weight FOV means within well, equal-weight well means within arm. Shared single-plane nuclear detection. A is the observed q95-positive pool among all detected MIAT, an operational spatial category.', '',
        'Intervals use within-arm A,T covariance, assume independent arm well vectors and do not automatically model slide clustering; normal log-delta uncertainty is substantial with six wells/arm. The exact test moves intact A,T pairs within slides. The complete 400-assignment design was enumerated using the B0-recovered absolute log-R statistic and 1e-12 tie tolerance; all eight delivered p-values reconcile. The original assignment file was not delivered.', '',
        'The original primary gate failed. Frozen historical Holm-nine remains separate from exploratory Holm-eight on R; marginal CIs are not simultaneous. New scalar Welch/Hedges-g contrasts form a separate ten-test Holm family (two total endpoints, eight associated endpoints). Scalar historical exact p-values retain their labels.', '',
        'Full/focus figures use the same observations and inference; scalar focus axes expand without dropping points, all forest CI/MDE extents remain. SD describes the six well dots; CI describes R. COUNT is a separate all-eligible localization readout and never supplies a denominator. Coverage ratios are ratios of consistently aggregated well pool counts; matched usable-pool R uses usable T and positive A on that same support. Usable intensity denominator mass is missing.', '',
        'Independent well provenance, orthogonal KD, numeric laser constancy, specificity/registration/PSF validation remain unverified. Adaptive detection does not establish a lower bound on biological MIAT loss. Classical coefficients require individual interpretation; Pearson is invariant to positive scaling, while thresholded fractions can change.', '']
    for r in sheets['Ratio intervals'].to_dict('records'):
        caption += [r['measurement']+': '+policy_caption(r, LABELS[POLICIES.index(r['null_policy'])]), '']
    guard_output(out_dir/'READOUT.md').write_text('\n'.join(caption), encoding='utf-8')
    for name in ('Ratio intervals', 'Ratio well values', 'Retention sensitivity', 'Usable pool ratios', 'Scalar contrasts', 'Slide values', 'Source provenance'):
        sheets[name].to_csv(guard_output(out_dir/(name.lower().replace(' ', '_')+'.csv')), index=False)
    return dict(out_dir=out_dir, xlsx=xlsx, figures=figures.to_dict('records'), group_order=['NT', 'KD'],
                reference='NT', absent=[], sheets=sheets)
