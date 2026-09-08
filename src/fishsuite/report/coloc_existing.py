"""Read and render persisted standard-coloc results. No measurement/null engine imports."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .aggregate import ReportInputError
from . import endpoints as ep
from .provenance import sha256, guard_output


SIMPLE_METRICS = ('pearson_r_csp', 'manders_m1_costes_only',
                  'manders_m2_costes_only', 'li_icq_csp')


def read_simple_metrics(path: Path) -> pd.DataFrame:
    """Read named persisted nucleus rows, retaining source identity and missingness.

    The historical ``manders_*_costes`` columns mix Costes and run thresholds.
    Only the converged-only columns support a figure labelled Costes.
    """
    path = Path(path).resolve()
    if not path.is_file():
        raise ReportInputError(f'missing standard panel per-nucleus table: {path}')
    df = _groups(pd.read_csv(path))
    required = set(SIMPLE_METRICS) | {'image','nucleus_id','condition','group',
                                     'secondary_only','costes_converged'}
    if not required <= set(df):
        raise ReportInputError(f'missing simple coloc columns: {sorted(required-set(df))}')
    for column in ('secondary_only', 'costes_converged'):
        if not df[column].isin([True, False]).all():
            raise ReportInputError(f'missing or invalid boolean: {column}')
    if df[['image','nucleus_id','condition','group']].isna().any().any():
        raise ReportInputError('missing simple coloc nucleus identity')
    if df.duplicated(['condition','image','nucleus_id']).any():
        raise ReportInputError('duplicate simple coloc nucleus keys')
    for metric in SIMPLE_METRICS:
        df[metric] = pd.to_numeric(df[metric], errors='raise')
        lo,hi = (-1,1) if metric.startswith('pearson') else (-.5,.5) if metric.startswith('li_') else (0,1)
        if not df[metric].dropna().between(lo,hi).all():
            raise ReportInputError(f'invalid simple coloc range: {metric}')
        if 'costes_only' in metric and df.loc[~df.costes_converged,metric].notna().any():
            raise ReportInputError('Costes failure has a converged-only measurement')
    if 'well_id' in df and not df.well_id.equals(df.condition):
        raise ReportInputError('panel well_id differs from condition')
    df['well_id'] = df.condition
    df['source_file'],df['source_sha256'] = str(path),sha256(path)
    df['source_row'] = np.arange(len(df))+2
    return df


def simple_rollups(nuclei):
    """Equal nucleus weight inside FOV, equal defined FOV weight inside well."""
    bio = nuclei.loc[~nuclei.secondary_only.eq(True)]
    keys = ['group','well_id','image']
    fields = bio.groupby(keys,sort=True)[list(SIMPLE_METRICS)].mean().reset_index()
    counts = bio.groupby(keys,sort=True)[list(SIMPLE_METRICS)].count().reset_index()
    fields = fields.merge(counts.rename(columns={m:m+'_n_nuclei' for m in SIMPLE_METRICS}),on=keys,validate='one_to_one')
    wells = fields.groupby(keys[:2],sort=True)[list(SIMPLE_METRICS)].mean().reset_index()
    return fields,wells


def representative_nuclei(nuclei):
    """Closest to each arm's median r; ties by well, image, numeric nucleus ID."""
    bio = nuclei.loc[~nuclei.secondary_only.eq(True) & np.isfinite(nuclei.pearson_r_csp)].copy()
    rows=[]
    for _,sub in bio.groupby('group',sort=True):
        median=float(sub.pearson_r_csp.median())
        sub['distance_to_arm_median']=(sub.pearson_r_csp-median).abs()
        minimum=sub.distance_to_arm_median.min()
        tied=sub.loc[np.isclose(sub.distance_to_arm_median,minimum,rtol=0,atol=1e-14)]
        row=tied.sort_values(['well_id','image','nucleus_id']).iloc[0].copy()
        row['arm_median_pearson']=median
        rows.append(row)
    return pd.DataFrame(rows).reset_index(drop=True)


def simple_statistics(nuclei, reference):
    from .sensitivities import mixed_model
    from .stats import welch
    fields,wells=simple_rollups(nuclei)
    groups=sorted(wells.group.unique())
    if len(groups)!=2 or reference not in groups:
        raise ReportInputError('simple coloc requires two arms and an explicit reference')
    test=next(g for g in groups if g!=reference)
    rows=[]
    bio=nuclei.loc[~nuclei.secondary_only.eq(True)]
    for metric in SIMPLE_METRICS:
        a,b=[wells.loc[wells.group.eq(g),metric].dropna().to_numpy() for g in (test,reference)]
        row=dict(endpoint=metric,test_group=test,reference_group=reference,
                 threshold='Costes-converged nuclei only' if 'manders' in metric else 'threshold-free',
                 **welch(a,b),**mixed_model(bio[['group','well_id','image',metric]].rename(columns={metric:'value'}),test,reference))
        row['p_mixed']=row['sensitivity_mixed_p']
        rows.append(row)
    return fields,wells,pd.DataFrame(rows)


def render_simple_coloc(nuclei, fields, wells, contrasts, out_dir, rna, partner, reference, run_name=None):
    """Four persisted metrics with biological replicate dots and mixed-model bracket."""
    from types import SimpleNamespace
    from . import figures as fig
    order=[reference]+[g for g in wells.group.unique() if g!=reference]
    ctx=SimpleNamespace(group_order=order, colors=dict(zip(order,['#595959','#D67AE5' if partner=='RNASEH2B' else '#CC79A7'])),
                        technical_layer='none',reference=reference,footer=lambda x:x)
    fig.set_style()
    canvas,axes=fig.plt.subplots(1,4,figsize=(14,5.2))
    names=['Pearson r','Manders M1','Manders M2','Li ICQ']
    labels=[f'Pearson r, {rna} × {partner}\n(per nucleus)',
            f'Fraction of {rna} signal in\n{partner}-positive pixels (Costes)',
            f'Fraction of {partner} signal in\n{rna}-positive pixels (Costes)',
            f'Li ICQ, {rna} × {partner}\n(per nucleus)']
    footer=[]
    for ax,metric,name,label,letter in zip(axes,SIMPLE_METRICS,names,labels,'ABCD'):
        # A15 explicitly changes inference and keeps coefficients as fractions;
        # do not inherit the historical registry's descriptive/percent policies.
        display_endpoint='simple_coloc:'+metric
        w=wells[['group','well_id',metric]].rename(columns={metric:'well_mean_of_field_values'}).assign(endpoint=display_endpoint)
        f=fields[['group','well_id','image',metric,metric+'_n_nuclei']].rename(columns={metric:'field_value',metric+'_n_nuclei':'n_nuclei_nonmissing'}).assign(endpoint=display_endpoint,level='nucleus')
        display_contrasts=contrasts.loc[contrasts.endpoint.eq(metric)].assign(endpoint=display_endpoint)
        fig.draw_replicate_simple(ax,ctx,display_endpoint,w,f,nuclei,display_contrasts,label,metric)
        # Keep coefficient units (0..1), with inference outside the data axes.
        for line in ax.lines:
            if len(line.get_xdata())==2 and np.array_equal(line.get_xdata(),[0,1]):
                line.set_visible(False)
        for text in ax.texts:
            if 'p =' in text.get_text(): text.set_visible(False)
        row=contrasts.loc[contrasts.endpoint.eq(metric)].iloc[0]
        p=row.p_mixed
        headline=f'mixed p = {fig.fmt_p(p)}' if row.sensitivity_mixed_status=='ok' else 'mixed p unavailable'
        ax.plot([0,0,1,1],[1.02,1.05,1.05,1.02],transform=ax.get_xaxis_transform(),clip_on=False,color='black',lw=.8)
        ax.text(.5,1.06,headline,transform=ax.transAxes,ha='center',fontsize=10)
        vals=w.well_mean_of_field_values.dropna().to_numpy()
        extrema=list(vals)
        for _,sub in w.groupby('group'):
            v=sub.well_mean_of_field_values.dropna()
            if len(v)>1: extrema.extend([v.mean()-v.std(),v.mean()+v.std()])
        def setter(variant, ax=ax,metric=metric,extrema=extrema):
            bound=1 if metric.startswith('pearson') else .5 if metric.startswith('li_') else 1
            lo,hi=(-bound,bound) if 'manders' not in metric else (0,1)
            if variant=='focus' and extrema:
                pad=max((max(extrema)-min(extrema))*.15,.02)
                lo,hi=max(lo,min(extrema)-pad),min(hi,max(extrema)+pad)
                if lo<0 and 'manders' not in metric: lo,hi=-max(abs(lo),abs(hi)),max(abs(lo),abs(hi))
            ax.set_ylim(lo,hi)
            for patch in ax.patches:
                if patch.get_gid() and patch.get_gid().startswith('group-mean:'):
                    value=patch.get_y()+patch.get_height(); patch.set_y(0); patch.set_height(value)
        ax._replicate_simple_axis=setter
        ax.set_title(f'{letter}  {name}',loc='left',pad=48,fontsize=13,fontweight='bold')
        fig.no_box(canvas,ax)
        footer.append(f'{letter}: Welch (wells) p {fig.fmt_p(row.p_welch)}')
    canvas.subplots_adjust(left=.065,right=.985,bottom=.22,top=.73,wspace=.63)
    if run_name is None and 'source_file' in nuclei:
        run_name = Path(nuclei.source_file.iloc[0]).parent.parent.name
    foot = ' | '.join(footer) + '; bars = well mean +/- SD; dots = wells; Manders: Costes-converged only; axis: focus window'
    if run_name: foot += '; run ' + run_name
    artist = canvas.text(.02,.015,foot,fontsize=9)
    canvas.canvas.draw()
    width = artist.get_window_extent().width
    if width > canvas.bbox.width*.96: artist.set_fontsize(9*canvas.bbox.width*.95/width)
    records=[]
    fig.save(canvas,Path(out_dir),'FIG_SIMPLE_COLOC',records,f'{rna} × {partner}','Simple coloc metrics / Simple coloc wells')
    return records


def render_cytofluorogram(selected, pixels, out_dir, rna, partner):
    """Plot exact recorded-plane nuclear pixels; never regenerate thresholds.

    ``pixels`` maps (well_id, image, nucleus_id) to (rna1, partner) vectors.
    The caller must supply raw intensities inside the persisted label mask.
    """
    from . import figures as fig
    if len(selected)!=2:
        raise ReportInputError('cytofluorogram requires one representative per arm')
    vectors=[]
    for row in selected.itertuples():
        key=(row.well_id,row.image,row.nucleus_id)
        if key not in pixels: raise ReportInputError(f'missing nuclear pixels: {key}')
        x,y=[np.asarray(a,float).ravel() for a in pixels[key]]
        if len(x)!=len(y) or len(x)!=int(row.n_pix) or not np.isfinite([x,y]).all():
            raise ReportInputError(f'invalid nuclear pixel vectors: {key}')
        r=float(np.corrcoef(x,y)[0,1])
        if not np.isclose(r,row.pearson_r_csp,atol=1e-8,rtol=1e-8):
            raise ReportInputError(f'pixel Pearson differs from persisted panel: {key}')
        vectors.append((x,y))
    xmax=max(x.max() for x,y in vectors)*1.03
    ymax=max(y.max() for x,y in vectors)*1.03
    fig.set_style(); canvas,axes=fig.plt.subplots(1,2,figsize=(12.4,5.6),sharex=True,sharey=True)
    for ax,row,(x,y) in zip(axes,selected.itertuples(),vectors):
        ax.hexbin(x,y,gridsize=90,mincnt=1,bins='log',cmap='cividis',extent=(0,xmax,0,ymax),rasterized=True)
        if row.costes_converged:
            ax.axvline(row.costes_thr_rna1,ls='--',color='#D55E00',lw=1)
            ax.axhline(row.costes_thr_partner,ls='--',color='#D55E00',lw=1)
            threshold='Dashed lines: recorded Costes thresholds'
        else:
            threshold='Costes failed: thresholds unavailable'
        ax.set(xlim=(0,xmax),ylim=(0,ymax),xlabel=f'{rna} intensity (raw units)',ylabel=f'{partner} intensity (raw units)')
        ax.set_title(f'{row.group} · Pearson r = {row.pearson_r_csp:.3f}',fontsize=15)
        ax.text(.02,.98,threshold,transform=ax.transAxes,va='top',fontsize=9)
        fig.no_box(canvas,ax)
    canvas.subplots_adjust(left=.08,right=.98,bottom=.19,top=.87,wspace=.24)
    canvas.text(.5,.035,'Representative nucleus nearest arm median Pearson r · identical axes · colour shows log pixel count',ha='center',fontsize=11)
    records=[]; fig.save(canvas,Path(out_dir),'FIG_CYTOFLUOROGRAM',records,f'{rna} × {partner}','Simple coloc representatives')
    return records


def load_representative_pixels(selected, run_dir, image_paths):
    """Read explicitly named images and existing masks; no search or measurement run."""
    import importlib.util
    script = Path(__file__).resolve().parents[3] / 'scripts' / 'coloc_standard_panel.py'
    spec = importlib.util.spec_from_file_location('fishsuite_stored_coloc_panel', script)
    panel = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(panel)
    run_dir = Path(run_dir)
    loader = panel.Run(run_dir)
    # Override only path resolution: the original resolver recursively searches.
    # Plane and mask loading remain the panel script's own implementation.
    loader.image_path = lambda name: Path(image_paths[name])
    pixels, sources = {}, []
    for row in selected.itertuples():
        key = (row.well_id, row.image, row.nucleus_id)
        path = loader.image_path(row.image)
        mask_path = run_dir/'masks'/f'{row.stem}__nuclei_label_mask.tif'
        if not path.is_file() or not mask_path.is_file():
            raise ReportInputError(f'missing representative image or mask: {path}; {mask_path}')
        mask = loader.label_mask(row.stem) == row.nucleus_id
        if int(mask.sum()) != int(row.n_pix):
            raise ReportInputError('representative mask pixel count mismatch')
        _, rna, partner, _ = loader.planes(row.image, row.z_plane)
        if rna.shape != mask.shape or partner.shape != mask.shape:
            raise ReportInputError('representative plane/mask shape mismatch')
        pixels[key] = (np.asarray(rna[mask],float), np.asarray(partner[mask],float))
        sources.append(dict(image=row.image,nucleus_id=row.nucleus_id,well_id=row.well_id,
            raw_image=str(path.resolve()),raw_image_sha256=sha256(path),mask=str(mask_path.resolve()),
            mask_sha256=sha256(mask_path),z_plane=row.z_plane,n_pix=int(mask.sum())))
    return pixels,pd.DataFrame(sources)


@dataclass
class ExistingPanel:
    path: Path
    tables: dict
    sources: pd.DataFrame
    manifest: dict


def verify_pinned(path: Path, manifest: dict) -> str:
    path = Path(path).resolve()
    if not path.is_file():
        raise ReportInputError(f'missing source: {path}')
    records = manifest.get('named_files', []) + manifest.get('supplemental_files', [])
    matches = [r for r in records if Path(r['path']).resolve() == path]
    if len(matches) != 1 or not matches[0].get('sha256'):
        raise ReportInputError(f'missing unique frozen hash: {path}')
    digest = sha256(path)
    if digest.lower() != matches[0]['sha256'].lower():
        raise ReportInputError(f'stale source hash: {path}')
    return digest


def _groups(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    if 'line' in frame:
        if 'group' in frame and not frame['group'].fillna('').equals(frame['line'].fillna('')):
            raise ReportInputError('conflicting panel line/group columns')
        frame['group'] = frame['line']
    return frame


def load_existing_panel(path: Path, figure_index: Path, baseline_manifest: Path,
                        figure_index_sha256: str | None = None) -> ExistingPanel:
    path, figure_index = Path(path), Path(figure_index)
    if not path.is_file() or not figure_index.is_file():
        raise ReportInputError(f'missing persisted panel or FIGURE_INDEX: {path}; {figure_index}')
    if baseline_manifest is None or not Path(baseline_manifest).is_file():
        raise ReportInputError('missing baseline manifest for persisted panel')
    manifest = json.loads(Path(baseline_manifest).read_text(encoding='utf-8'))
    digest = verify_pinned(path, manifest)
    # A0 pins the measurement workbook but does not pin this prose-only index.
    # Record its current bytes; deck specs additionally pin the A3 observed hash.
    index_digest = sha256(figure_index)
    if figure_index_sha256 and index_digest != figure_index_sha256:
        raise ReportInputError(f'stale figure index: {figure_index}')
    book = load_workbook(path, read_only=True, data_only=True)
    tables, sources = {}, []
    required = ('README', 'per_nucleus', 'per_FOV', 'per_well', 'contrasts',
                'line_profiles', 'overlays_index')
    try:
        for name in required:
            if name not in book.sheetnames:
                raise ReportInputError(f'missing panel sheet: {name}')
            ws = book[name]
            values = list(ws.values)
            if not values or len(set(values[0])) != len(values[0]):
                raise ReportInputError(f'missing or duplicate panel headers: {name}')
            # Use pandas' native dtype inference, without modifying any source cells.
            tables[name] = _groups(pd.read_excel(path, sheet_name=name))
            for ri, row in enumerate(ws.iter_rows(min_row=2),2):
                for ci, cell in enumerate(row,1):
                    sources.append(dict(file=str(path.resolve()), sha256=digest, sheet=name,
                                        cell=f'{get_column_letter(ci)}{ri}', value=cell.value))
    finally:
        book.close()
    sources.append(dict(file=str(figure_index.resolve()), sha256=index_digest,
                        sheet='', cell='', value=figure_index.read_text(encoding='utf-8')))
    required_cols = set(ep.A3_PARTNER_ADDITIONS) | set(ep.A3_PANEL_ALIASES)
    for name in ('per_nucleus', 'per_FOV', 'per_well'):
        missing = required_cols - set(tables[name])
        if missing:
            raise ReportInputError(f'missing {name} endpoints: {sorted(missing)}')
    pn = tables['per_nucleus']
    # The v3 contrasts sheet appends a second pooled-pixel table after two blank
    # rows. Preserve the whole sheet, but its secondary header/arms are not tests.
    raw_contrasts=tables['contrasts']
    separators=np.flatnonzero(raw_contrasts['endpoint'].isna().to_numpy())
    stop=int(separators[0]) if len(separators) else len(raw_contrasts)
    tables['endpoint_contrasts']=raw_contrasts.iloc[:stop].copy()
    if pn.duplicated(['image', 'nucleus_id']).any():
        raise ReportInputError('duplicate panel nucleus keys')
    for c in ('manders_m1_costes_only', 'manders_m2_costes_only'):
        if pn.loc[~pn['costes_converged'].eq(True), c].notna().any():
            raise ReportInputError('Costes failure incorrectly recorded as converged-only value')
    return ExistingPanel(path.resolve(), tables, pd.DataFrame(sources), manifest)


def validate_pairing(panel_nuclei: pd.DataFrame, spots: pd.DataFrame) -> None:
    """Independent observed forward numerator/denominator from nuclear spot flags."""
    keys = ['image', 'nucleus_id']
    required = set(keys + ['channel', 'in_nucleus', 'paired_at_0p3um'])
    if not required <= set(spots):
        raise ReportInputError(f'missing pairing spot fields: {sorted(required-set(spots))}')
    nuclear = spots.loc[(spots.channel == 'rna1') & spots.in_nucleus.eq(1) & (spots.nucleus_id > 0)]
    if not nuclear['paired_at_0p3um'].isin([0, 1, False, True]).all():
        raise ReportInputError('missing or invalid nuclear pairing calls')
    expected = nuclear.groupby(keys)['paired_at_0p3um'].mean()
    actual = panel_nuclei.set_index(keys)['paired_frac_rna1_at_partner']
    wanted = expected.reindex(actual.index)
    if not np.allclose(actual, wanted, atol=1e-12, rtol=1e-10, equal_nan=True):
        raise ReportInputError('panel nuclear pairing differs from persisted nuclear spot flags')


def integrate_panel(data: dict, panel: ExistingPanel, endpoints: list) -> tuple:
    """Validate full retained cohort, aliases and anchors before adding columns."""
    nuc = data['nuclei']
    pn = panel.tables['per_nucleus']
    keys = ['image', 'nucleus_id']
    a, b = nuc.set_index(keys).sort_index(), pn.set_index(keys).sort_index()
    if not a.index.equals(b.index):
        raise ReportInputError('persisted panel cohort mismatch: nucleus roster')
    for c in ('condition', 'group', 'secondary_only'):
        mask = ~a.secondary_only.eq(True) if c == 'group' else pd.Series(True,index=a.index)
        if c in a and c in b and not a.loc[mask,c].fillna('').equals(b.loc[mask,c].fillna('')):
            raise ReportInputError(f'persisted panel cohort mismatch: {c}')
    for pc, nc in (('n_rna1_nuclear_puncta', 'nuclear_spot_count'),
                   ('n_partner_nuclear_puncta', 'n_spots_protein')):
        if not np.array_equal(a[nc], b[pc]):
            raise ReportInputError(f'persisted panel anchor count mismatch: {pc}')
    by_name = {e.name:e for e in endpoints}
    for alias, name in ep.A3_PANEL_ALIASES.items():
        if name not in by_name:
            raise ReportInputError(f'missing original alias endpoint: {name}')
        if not np.allclose(a[by_name[name].column], b[alias], atol=1e-12, rtol=1e-10, equal_nan=True):
            raise ReportInputError(f'persisted alias differs: {alias}')
    for record in panel.manifest['endpoint_registry']:
        old = by_name.get(record['name'])
        if old is None:
            raise ReportInputError(f'frozen registry member missing: {record["name"]}')
        for attr in ('family', 'descriptive_only', 'absolute_intensity', 'excluded_from_holm', 'usability_flag'):
            if getattr(old, attr) != record[attr]:
                raise ReportInputError(f'frozen registry membership changed: {old.name}/{attr}')
    added = ep.a3_endpoints(panel.tables['per_well'].columns)
    columns = [e.column for e in added if e.column in pn and e.column not in nuc]
    merged = nuc.merge(pn[keys+columns], on=keys, how='left', validate='one_to_one')
    for endpoint in added:
        if endpoint.column not in merged:
            raise ReportInputError(f'missing A endpoint: {endpoint.column}')
    data['nuclei'] = merged
    return endpoints + added, added


def render_existing(panel: ExistingPanel, ctx, contrasts: pd.DataFrame, out_dir: Path) -> list:
    """Re-render every persisted per-well metric using the common replicate-simple key."""
    from . import figures as fig
    out_dir = guard_output(out_dir)
    fig.set_style()
    manifest = []
    pw = panel.tables['per_well']
    pf = panel.tables['per_FOV']
    pn = panel.tables['per_nucleus']
    pw = pw.loc[~pw.secondary_only.eq(True)]
    pf = pf.loc[~pf.secondary_only.eq(True)]
    for metric in panel.tables['endpoint_contrasts']['endpoint']:
        if metric not in pw:
            raise ReportInputError(f'missing persisted per-well metric: {metric}')
        name = ep.A3_PANEL_ALIASES.get(metric, metric)
        w = pw[['group','well_id',metric]].rename(columns={metric:'well_mean_of_field_values'}).assign(endpoint=name)
        f = pf[['group','well_id','image',metric]].rename(columns={metric:'field_value'}).assign(endpoint=name,level='nucleus')
        finite = pn.loc[~pn.secondary_only.eq(True)].groupby('image')[metric].count()
        f['n_nuclei_nonmissing'] = f.image.map(finite)
        r = contrasts.loc[contrasts.endpoint == name]
        # Historical descriptive rows are not re-tested for a plot. Added tests use
        # the explicitly amended report contrasts; all other p cells stay missing.
        if metric not in ep.A3_PARTNER_ADDITIONS and metric not in ep.A3_PANEL_ALIASES:
            r = r.iloc[:0]
        desc = panel.tables['contrasts'].loc[lambda x: x.endpoint == metric].iloc[0]
        fig.superplot_standalone(ctx, name, str(desc['label']), str(desc['unit']),
                                w, f, pn, r, None, out_dir, metric, manifest)
        manifest[-1].update(endpoint=metric, source_file=str(panel.path),
                            source_sheet='per_well', sha256=sha256(out_dir / manifest[-1]['png']),
                            svg_sha256=sha256(out_dir / manifest[-1]['svg']),
                            source_cells=';'.join(f'per_well!{get_column_letter(pw.columns.get_loc(metric)+1)}{i+2}' for i in pw.index),
                            report_cells=';'.join(f'Coloc per well!{get_column_letter(pw.columns.get_loc(metric)+1)}{i+3}' for i in pw.index))
    (out_dir/'FIGURE_INDEX.md').write_text('# Persisted standard panel\n\n'
        'WT #595959; QKI-KO #D67AE5. Nuclear anchors only. No new nulls.\n\n' +
        '\n'.join(f'- {r["png"]}; full: {r.get("full_png", "")}: {r["endpoint"]}; source {panel.path} / per_well' for r in manifest), encoding='utf-8')
    return manifest
